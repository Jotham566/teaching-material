"""One-page roadmap (A4 landscape) for Enterprise Data Readiness for AI: editable PowerPoint and PDF.   python3 make_roadmap.py
Same system as the AI and Data Science Professional roadmap: type and hairlines, one ink, one accent used for workpiece sessions."""
import subprocess, sys, json
from pathlib import Path
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import deckkit as dk
ROOT = HERE.parents[2]
OUT = ROOT / '01 Course Overview'; OUT.mkdir(exist_ok=True)
C = json.load(open(HERE / 'curriculum.json'))

INK, BODY, SOFT, RULE, ACCENT, WHITE = '1F2328', '3B4046', '6A6F75', 'C9C5BC', 'B04A2B', 'FFFFFF'
SERIF, SANS = 'Cambria', 'Calibri'
T = dict(display=30.0, row=16.5, lead=14.0, body=12.5, small=11.0, dot=10.5)
PW, PH = 11.693, 8.268; MX = 0.6

prs = Presentation(); prs.slide_width = Inches(PW); prs.slide_height = Inches(PH)
s = prs.slides.add_slide(prs.slide_layouts[6]); sh = s.shapes
bg = s.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor.from_string(WHITE)

def text(x, y, w, h, t, size, face=SANS, color=BODY, bold=False, align='l', anchor='t', italic=False, name=None):
    tb, sz = dk.add_text(sh, x, y, w, h, t, size, bold, color, face, align, anchor, min_size=size, italic=italic, name=name)
    return tb
def line(x1, y1, x2, y2, color=RULE, w=0.75):
    c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(color); c.line.width = Pt(w)
    st = c._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: c._element.remove(st)
    return c
def dot(cx, cy, d, label, filled):
    o = sh.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    st = o._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: o._element.remove(st)
    o.fill.solid(); o.fill.fore_color.rgb = RGBColor.from_string(ACCENT if filled else WHITE)
    o.line.color.rgb = RGBColor.from_string(ACCENT if filled else INK); o.line.width = Pt(1.25)
    tf = o.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = str(label)
    r.font.name = SANS; r.font.size = Pt(T['dot']); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(WHITE if filled else INK)

# ---- header
text(MX, 0.42, 8.0, 0.62, C['title'], T['display'], SERIF, INK, True, name='Title')
text(MX, 1.06, 7.3, 0.62, 'Twelve sessions of two hours for leaders and their data teams. You leave with a scored baseline, a working search and a twelve-month plan.', T['lead'], SANS, BODY)
text(8.2, 0.5, PW - MX - 8.2, 1.15, [[('Bring', {'bold': True, 'color': INK})], [('A decision maker and someone from data or IT, from the same organisation. A laptop and 50 documents that are not confidential.', {})]], T['body'], SANS, BODY)
line(MX, 1.68, PW - MX, 1.68, INK, 1.25)

# ---- column heads
cols = [MX, 3.4, 7.05, 9.0]; widths = [2.7, 3.5, 1.8, PW - MX - 9.0]
for x, w, h in zip(cols, widths, ['Module and sessions', 'What you do', 'Tools', 'What you make']):
    text(x, 1.76, w, 0.26, h, T['small'], SANS, SOFT)
line(MX, 2.04, PW - MX, 2.04, RULE, 0.75)

CHECK = set(C['workpieces'])
y = 2.12; rh = 0.92
for i, m in enumerate(C['modules']):
    text(cols[0], y + 0.04, widths[0], 0.32, m['name'], T['row'], SERIF, INK, True, name=f'Module {i+1}')
    nums = m['sessions']; d = 0.3; step = 0.4; cy = y + 0.64
    line(cols[0] + d / 2, cy, cols[0] + d / 2 + step * (len(nums) - 1), cy, INK, 1.0)
    for k, n in enumerate(nums): dot(cols[0] + d / 2 + k * step, cy, d, n, n in CHECK)
    text(cols[1], y + 0.05, widths[1], rh - 0.1, m['what'], T['body'], SANS, BODY)
    text(cols[2], y + 0.05, widths[2], rh - 0.1, m['tools'], T['body'], SANS, BODY)
    text(cols[3], y + 0.05, widths[3], rh - 0.1, m['make'], T['body'], SANS, BODY)
    y += rh
    line(MX, y, PW - MX, y, INK if i == len(C['modules']) - 1 else RULE, 1.25 if i == len(C['modules']) - 1 else 0.5)

# ---- below the table
y += 0.1
dot(MX + 0.11, y + 0.13, 0.22, '', True)
text(MX + 0.32, y, 4.6, 0.28, 'Workpiece sessions: 2, 7 and 12', T['small'], SANS, BODY)
text(MX, y + 0.3, 5.0, 0.5, [[('Case. ', {'bold': True, 'color': INK}), ('We work on Kitara Group, a fictional bank and insurer, then on your own organisation.', {})]], T['body'], SANS, BODY)
text(6.0, y, PW - MX - 6.0, 0.5, [[('Every session. ', {'bold': True, 'color': INK}), ('Case 10 minutes, teach 30, live demo 25, break 5, workshop 40, decide 10.', {})]], T['body'], SANS, BODY)
text(6.0, y + 0.5, PW - MX - 6.0, 0.5, [[('Everything is free. ', {'bold': True, 'color': INK}), ('No paid tool, cloud account or graphics card. We read each tool’s licence in session 10.', {})]], T['body'], SANS, BODY)
text(MX, PH - 0.33, 8, 0.26, 'KLAi Professional Learning. Tutor: Jotham Martin Wambi.', T['small'], SANS, SOFT)

pp = OUT / 'Course Roadmap (one page).pptx'; prs.save(str(pp))
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(pp)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=200)
print('written', pp.name)
