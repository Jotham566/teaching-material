"""Rendering helpers for the Kitara corpus: text PDFs, .docx, .txt, scanned PDFs, OCR error rate."""
import re, os, glob, subprocess, tempfile
from pathlib import Path
from xml.sax.saxutils import escape
import numpy as np

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Flowable, KeepTogether, SimpleDocTemplate)
from reportlab.pdfgen import canvas as rl_canvas

BODY = ParagraphStyle("body", fontName="Helvetica", fontSize=11, leading=15.5, spaceAfter=9)
H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12.5, leading=16, spaceBefore=10, spaceAfter=4)
TITLE = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=16, leading=20, spaceAfter=14)
BUL = ParagraphStyle("bul", parent=BODY, leftIndent=14, bulletIndent=3, spaceAfter=4)


def parse_blocks(body):
    """Yield ('h', text), ('b', text) or ('p', [lines])."""
    for block in re.split(r"\n\s*\n", body.strip()):
        lines = [l.rstrip() for l in block.split("\n")]
        if len(lines) == 1 and lines[0].startswith("## "):
            yield ("h", lines[0][3:])
            continue
        cur = []
        for l in lines:
            if l.startswith("- "):
                if cur:
                    yield ("p", cur)
                    cur = []
                yield ("b", l[2:])
            elif l.startswith("## "):
                if cur:
                    yield ("p", cur)
                    cur = []
                yield ("h", l[3:])
            else:
                cur.append(l)
        if cur:
            yield ("p", cur)


class NumberedCanvas(rl_canvas.Canvas):
    def __init__(self, *a, **k):
        self._hdr = k.pop("hdr", None)
        self._ftr = k.pop("ftr", None)
        super().__init__(*a, **k)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        n = len(self._saved)
        for st in self._saved:
            self.__dict__.update(st)
            if self._hdr:
                self.setFont("Helvetica", 8)
                self.setFillColor(colors.grey)
                self.drawString(20 * mm, A4[1] - 12 * mm, self._hdr)
                self.setStrokeColor(colors.lightgrey)
                self.setLineWidth(0.4)
                self.line(20 * mm, A4[1] - 14 * mm, A4[0] - 20 * mm, A4[1] - 14 * mm)
            if self._ftr:
                self.setFont("Helvetica", 8)
                self.setFillColor(colors.grey)
                self.drawString(20 * mm, 10 * mm, self._ftr)
                self.drawRightString(A4[0] - 20 * mm, 10 * mm, f"Page {self._pageNumber} of {n}")
            super().showPage()
        super().save()


class ColumnwiseTable(Flowable):
    """A ruled table whose cell text is written column by column. It looks normal on screen and with a table
    reader that uses the ruling lines, but a plain text extractor returns column 1, then column 2, and so on."""

    def __init__(self, rows, widths, rowh=17):
        super().__init__()
        self.rows, self.widths, self.rowh = rows, widths, rowh
        self.width, self.height = sum(widths), rowh * len(rows)

    def wrap(self, aw, ah):
        return self.width, self.height

    def draw(self):
        c = self.canv
        c.setLineWidth(0.5)
        c.setStrokeColor(colors.black)
        y0 = self.height
        for i in range(len(self.rows) + 1):
            c.line(0, y0 - i * self.rowh, self.width, y0 - i * self.rowh)
        x = 0
        for w in self.widths + [0]:
            c.line(x, 0, x, self.height)
            x += w
        x = 0
        for j, w in enumerate(self.widths):
            for i, r in enumerate(self.rows):
                c.setFont("Helvetica-Bold" if i == 0 else "Helvetica", 9)
                c.drawString(x + 4, y0 - (i + 1) * self.rowh + 5, str(r[j]))
            x += w


def render_text_pdf(d, path):
    story = []
    is_letter = d["family"] in ("claims_letter",)
    if not is_letter:
        story.append(Paragraph(escape(d["title"]), TITLE))
    for kind, val in parse_blocks(d["body"]):
        if kind == "h":
            story.append(Paragraph(escape(val), H2))
        elif kind == "b":
            story.append(Paragraph(escape(val), BUL, bulletText="-"))
        else:
            story.append(Paragraph("<br/>".join(escape(l) for l in val), BODY))
    if d.get("table"):
        story.append(Spacer(1, 8))
        rows = d["table"]
        ncol = len(rows[0])
        widths = [62 * mm, 38 * mm, 70 * mm][:ncol] if ncol == 3 else [85 * mm, 85 * mm]
        story.append(ColumnwiseTable(rows, widths))
        story.append(Spacer(1, 8))
        story.append(Paragraph("Source: Kitara Group rate card. Check the intranet for the latest edition.", BODY))
    hdr = f"KITARA GROUP  |  {d['title']}" if d.get("hf") else None
    ftr = f"Classification: {d['sensitivity']}" if d.get("hf") else None

    def mk(*a, **k):
        return NumberedCanvas(*a, hdr=hdr, ftr=ftr, **k)

    top = 24 * mm if hdr else 18 * mm
    bot = 18 * mm
    if d.get("cols", 1) == 2:
        from reportlab.platypus.flowables import BalancedColumns
        title_part, rest = story[:1], story[1:]
        story = title_part + [BalancedColumns(rest, nCols=2, innerPadding=6 * mm, spaceBefore=2, spaceAfter=2)]
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=top, bottomMargin=bot,
                            title=d["title"], author="Kitara Group")
    doc.build(story, canvasmaker=mk)


def render_docx(d, path):
    import docx
    from docx.shared import Pt
    D = docx.Document()
    st = D.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    D.core_properties.title = d["title"]
    D.core_properties.author = "Kitara Group"
    import datetime as _dt
    D.core_properties.created = D.core_properties.modified = D.core_properties.last_printed = _dt.datetime(2026, 10, 29, 9, 0, 0)
    blocks = list(parse_blocks(d["body"]))
    first = True
    for kind, val in blocks:
        if kind == "h":
            D.add_heading(val, level=2)
        elif kind == "b":
            D.add_paragraph(val, style="List Bullet")
        else:
            txt = "\n".join(val)
            if first:
                D.add_heading(txt.split("\n")[0], level=1)
                if "\n" in txt:
                    D.add_paragraph("\n".join(txt.split("\n")[1:]))
            else:
                D.add_paragraph(txt)
        first = False
    D.save(str(path))


def render_txt(d, path):
    Path(path).write_text(f"KITARA GROUP\n{d['title']}\n\n" + d["body"].strip() + "\n", encoding="utf-8")


# ------------------------------------------------------------------ scans
def _font(size, serif=False):
    from PIL import ImageFont
    cands = (["/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"] if serif else
             ["/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf", "/usr/share/fonts/truetype/liberation2/LiberationMono-Regular.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", "/Library/Fonts/Courier New.ttf", "/System/Library/Fonts/Courier.ttc"])
    for c in cands:
        if os.path.exists(c):
            return ImageFont.truetype(c, size)
    return ImageFont.load_default()


HV = dict(scale=0.54, blur=1.5, noise=12, specks=0, fm=0.9, fa=12, streaks=0, blot=8)
NV = dict(blur=0.9, noise=14)


def render_scan(d, path, seed, serif=False):
    from PIL import Image, ImageDraw, ImageFilter
    g = np.random.default_rng(seed)
    heavy = d.get("meta_note") == "heavy"
    W, H = 1240, 1754
    img = Image.new("L", (W, H), 246)
    dr = ImageDraw.Draw(img)
    fnt = _font(20 if not heavy else 21, serif)
    y = 150
    x = 120
    for line in d["body"].strip().split("\n"):
        dr.text((x, y), line, font=fnt, fill=25)
        y += 33
        if y > H - 150:
            break
    if "Signed:" in d["body"]:  # a hand drawn signature scribble
        sy = 150 + 33 * (d["body"].strip().split("\n").index([l for l in d["body"].split("\n") if "Signed:" in l][0])) + 14
        pts = [(x + 120 + i * 9, sy + 10 * np.sin(i * 0.9) + g.normal(0, 3)) for i in range(18)]
        dr.line(pts, fill=40, width=3)
    arr = np.asarray(img, dtype=np.float32)
    if heavy:
        arr = arr * HV['fm'] + HV['fa']  # faded
    # paper shading
    yy, xx = np.mgrid[0:H, 0:W]
    arr = arr - (0.05 if not heavy else 0.12) * 60 * (xx / W) * (yy / H)
    img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
    if heavy:
        small = img.resize((int(W * HV['scale']), int(H * HV['scale'])), Image.BILINEAR)
        img = small.resize((W, H), Image.BILINEAR)
        img = img.filter(ImageFilter.GaussianBlur(HV['blur']))
    else:
        img = img.filter(ImageFilter.GaussianBlur(NV['blur']))
    ang = float(g.uniform(0.5, 1.5)) * (1 if g.random() < 0.5 else -1)
    if heavy:
        ang = float(g.uniform(1.3, 1.9)) * (1 if g.random() < 0.5 else -1)
    img = img.rotate(ang, resample=Image.BICUBIC, expand=False, fillcolor=240)
    arr = np.asarray(img, dtype=np.float32)
    arr = arr + g.normal(0, HV['noise'] if heavy else NV['noise'], arr.shape)
    if heavy:  # washed-out patches (stains, fold wear) that hide part of some lines
        for _ in range(HV['blot']):
            r0, c0 = int(g.integers(140, 1100)), int(g.integers(60, 800))
            hh, ww = int(g.integers(14, 34)), int(g.integers(160, 420))
            arr[r0:r0 + hh, c0:c0 + ww] = np.maximum(arr[r0:r0 + hh, c0:c0 + ww], g.uniform(150, 215))
    if heavy:  # dark specks and streaks
        for _ in range(HV['specks']):
            r, c = g.integers(0, H), g.integers(0, W)
            arr[r:r + int(g.integers(1, 4)), c:c + int(g.integers(1, 4))] = g.integers(0, 90)
        for _ in range(HV['streaks']):
            c = int(g.integers(0, W))
            arr[:, c:c + 2] -= g.uniform(25, 55)
    img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
    import datetime as _dt
    _d = _dt.datetime(2026, 10, 29, 9, 0, 0)
    img.save(str(path), "PDF", resolution=150.0, title=d["title"], author="Kitara Group", creationDate="D:20261029090000Z", modDate="D:20261029090000Z")


def ocr_pdf(path, dpi=150):
    import pytesseract
    from PIL import Image
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["pdftoppm", "-r", str(dpi), "-png", str(path), os.path.join(td, "p")], check=True)
        return "\n".join(pytesseract.image_to_string(Image.open(p)) for p in sorted(glob.glob(os.path.join(td, "p*.png"))))


def edit_distance(a, b):
    a = np.frombuffer(a.encode("utf-32-le"), dtype=np.uint32)
    b = np.frombuffer(b.encode("utf-32-le"), dtype=np.uint32)
    n, m = len(a), len(b)
    prev = np.arange(m + 1, dtype=np.int64)
    idx = np.arange(m + 1, dtype=np.int64)
    for i in range(1, n + 1):
        cost = (b != a[i - 1]).astype(np.int64)
        tmp = np.empty(m + 1, dtype=np.int64)
        tmp[0] = i
        tmp[1:] = np.minimum(prev[1:] + 1, prev[:-1] + cost)
        prev = np.minimum.accumulate(tmp - idx) + idx
    return int(prev[m])


def cer(reference, hypothesis):
    ref = " ".join(reference.split())
    hyp = " ".join(hypothesis.split())
    return edit_distance(ref, hyp) / max(1, len(ref))
