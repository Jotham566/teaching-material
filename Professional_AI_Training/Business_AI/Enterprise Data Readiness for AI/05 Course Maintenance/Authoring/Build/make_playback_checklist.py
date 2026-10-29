"""Write 04 Facilitator Materials/Playback and Rehearsal Checklist.txt from the built decks and docs_data.json.  python3 make_playback_checklist.py"""
import json
from pathlib import Path
from pptx import Presentation
HERE = Path(__file__).resolve().parent; ROOT = HERE.parents[2]
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
data = json.loads((HERE / 'docs_data.json').read_text(encoding='utf-8'))
out = ["PLAYBACK AND REHEARSAL CHECKLIST", "",
 "Part 1, in PowerPoint on the machine you will teach from. Open each deck, choose Slide Show, and play every animated slide from the first click to the last. Confirm the parts appear in the order listed and nothing overlaps. Editing view shows the finished slide, so only Slide Show proves the build. Tick each line.",
 "Part 2, read the notes aloud once per session, timing yourself. The estimate is the speaking time of the notes alone at 130 words a minute, before pauses. The first paragraph is read before any click, then one paragraph per click, then any wrap-up.",
 "Part 3, run every cell that needs a heavy library once in Colab (sentence-transformers, chromadb, pgvector with psycopg2, pandera, great_expectations, presidio, dagster, duckdb with pyarrow, docling). run_all_labs.py in this folder runs everything else. Open the workbook named for each session and complete one workshop task with it.", ""]
for f in sorted((ROOT / '02 Class Slides').glob('[0-2][0-9] - *.pptx')):
    n = int(f.name[:2]); s = data['sessions'][n - 1]; prs = Presentation(str(f)); lines = []
    for i, sl in enumerate(prs.slides, 1):
        t = sl._element.find(P + 'timing')
        if t is None: continue
        ms = [x for x in t.iter(P + 'cTn') if x.get('nodeType') == 'mainSeq'][0]; k = len(ms.find(P + 'childTnLst'))
        title = sl.shapes.title.text_frame.text if sl.shapes.title is not None else f'slide {i}'
        lines.append(f"  [ ] Slide {i}: {title}, {k} click{'s' if k != 1 else ''}")
    heavy = [f"  [ ] Run in Colab once: slide {d['no']} ({d['title']}) needs {' and '.join(d['libs'])}" for d in s['demo'] if d['libs']] + [f"  [ ] Run in Colab once: notebook {h['where']} needs {' and '.join(h['libs'])}" for h in s['heavy_cells']]
    if n == 5: heavy.append("  [ ] Run the OCR cells once in Colab (the setup installs Tesseract and poppler with apt-get)")
    wb = s['workbooks'] + [f"page {p} in the Participant Workbook" for p in s['pages']]
    out += [f"SESSION {n}: {s['title']}", f"  Notes read aloud: about {s['minutes_spoken']} minutes.", "  [ ] Deck opens in PowerPoint and Slide Show runs from slide 1", *lines,
            *heavy, "  [ ] Notes read aloud once, wording changed where it does not sound like you",
            f"  [ ] Open and try: {'; '.join(wb) if wb else 'no template file (Task 1 describes the sheet participants make)'}", f"  [ ] Solution notebook runs in Colab: {s['solution_notebook']}", ""]
out += ["TOOLS THAT NEED YOUR ACCOUNT OR A DOWNLOAD", "  [ ] Colab opens and runs a cell with a Google account", "  [ ] data.zip hosted, DATA_ZIP_URL given to participants (or they upload it)", "  [ ] sentence-transformers model downloads in Colab (Session 6)", "  [ ] Model cards and licence files read for every tool marked check in the Setup Guide", ""]
(ROOT / '04 Facilitator Materials' / 'Playback and Rehearsal Checklist.txt').write_text("\n".join(out), encoding='utf-8'); print('written', len(out), 'lines')
