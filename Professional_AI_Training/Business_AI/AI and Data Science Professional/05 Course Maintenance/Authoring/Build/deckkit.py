"""Slide engine for AI and Data Science Professional (second design, October 2026).

Design: type and hairlines. White slides, one ink colour, one accent (clay) used for what matters, almost no fills.
No cards, no numbered circles, no kicker labels, no banner bars. Layouts differ by slide type so neighbours do not repeat.
Every text block is measured with metric-compatible fonts (Carlito for Calibri, Caladea for Cambria, Liberation Mono for
Courier New) and a block that cannot fit at its minimum size raises. Animations are native PowerPoint click effects.

Slide types: concept, flow, cycle, compare, table, code, figure, stat, steps, trap, check, define, matrix, bar, erd, summary.
"""
import re
from pathlib import Path
from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

# ------------------------------------------------------------------ palette and type (shared with the roadmap)
INK, BODY_C, SOFT, RULE, ACCENT, WHITE = '1F2328', '3B4046', '6A6F75', 'C9C5BC', 'B04A2B', 'FFFFFF'
DARK_BG, DARK_TEXT, DARK_SOFT, DARK_RULE, DARK_ACCENT = '1F2328', 'F3F0EA', 'B8B2A7', '4A4F55', 'E3A78F'
WELL = 'F5F3EF'                  # the one light fill: output panels
MOSS = '5E7A57'
W, H = 13.333, 7.5
MX = 0.75
Y0, Y1 = 1.75, 6.8
HEAD, BODY, MONO = 'Cambria', 'Calibri', 'Courier New'
FONTDIR = Path('/usr/share/fonts/truetype')
_FP = {('Calibri', False): 'crosextra/Carlito-Regular.ttf', ('Calibri', True): 'crosextra/Carlito-Bold.ttf', ('Cambria', False): 'crosextra/Caladea-Regular.ttf',
       ('Cambria', True): 'crosextra/Caladea-Bold.ttf', ('Courier New', False): 'liberation2/LiberationMono-Regular.ttf', ('Courier New', True): 'liberation2/LiberationMono-Bold.ttf'}
_cache = {}
def _font(face, bold, size):
    k = (face, bold, round(size * 4))
    if k not in _cache: _cache[k] = ImageFont.truetype(str(FONTDIR / _FP[(face, bold)]), max(4, int(round(size * 4))))
    return _cache[k]
def text_width_pt(s, face, bold, size): return _font(face, bold, size).getlength(s) / 4.0

def wrap_lines(text, face, bold, size, width_pt):
    out = []
    for para in text.split('\n'):
        if para == '': out.append(''); continue
        cur = ''
        for w in para.split(' '):
            t = (cur + ' ' + w) if cur else w
            if text_width_pt(t, face, bold, size) <= width_pt or not cur: cur = t
            else: out.append(cur); cur = w
        out.append(cur)
    return out

LS = 1.2
def block_height_in(text, face, bold, size, width_in, space_after_pt=0, mono=False):
    lines = wrap_lines(text, face, bold, size, width_in * 72)
    return (len(lines) * size * (1.17 if mono else LS) + space_after_pt) / 72.0 * 1.04 + 0.02

class FitError(Exception): pass
def fit_size(text, width_in, height_in, start, minimum, face=BODY, bold=False, mono=False):
    s = start
    longest_word = max((w for w in re.split(r'\s+', text) if w), key=len, default='')
    while s >= minimum:
        if (mono or text_width_pt(longest_word, face, bold, s) <= width_in * 72) and block_height_in(text, face, bold, s, width_in, mono=mono) <= height_in: return s
        s -= 0.5
    raise FitError(f"text does not fit at {minimum}pt in {width_in:.1f}x{height_in:.1f}in: {text[:70]!r}")

def rgb(h): return RGBColor.from_string(h)
def _set_font(run, face, size, bold=False, color=INK, italic=False):
    f = run.font; f.name = face; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = rgb(color)

def add_text(shapes, x, y, w, h, text, size=18, bold=False, color=BODY_C, face=BODY, align='l', anchor='t', margin=0.0, italic=False, name=None, min_size=None, space_after=0):
    """Measured text box. `text`: str, list of str paragraphs, list of run tuples for one paragraph, or list of run lists."""
    tb = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(text, list) and text and isinstance(text[0], tuple): paras = [text]
    elif isinstance(text, list): paras = text
    else: paras = [text]
    plain = '\n'.join(p if isinstance(p, str) else ''.join(r[0] for r in p) for p in paras)
    sz = fit_size(plain, w - 2 * margin, h, size, min_size if min_size else size, face, bold)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        if space_after: para.space_after = Pt(space_after)
        for t, o in ([(p, {})] if isinstance(p, str) else p):
            r = para.add_run(); r.text = t
            _set_font(r, o.get('face', face), o.get('size', sz), o.get('bold', bold), o.get('color', color), o.get('italic', italic))
    return tb, sz

def rule(shapes, x1, y1, x2, y2, color=RULE, w=0.75):
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    st = c._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: c._element.remove(st)
    return c

def box(shapes, x, y, w, h, fill=None, line=RULE, lw=0.75, shape=MSO_SHAPE.RECTANGLE):
    s = shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    st = s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: s._element.remove(st)
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    return s

def node(shapes, cx, cy, d=0.24, filled=False, color=INK):
    s = box(shapes, cx - d / 2, cy - d / 2, d, d, fill=(color if filled else WHITE), line=color, lw=1.5, shape=MSO_SHAPE.OVAL); return s

# ------------------------------------------------------------------ animation
P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
def animate(slide, groups):
    if not groups: return 0
    n = [2]
    def nid(): n[0] += 1; return n[0]
    clicks = ''
    for g in groups:
        inner = ''
        for k, sid in enumerate(g):
            node_t = 'clickEffect' if k == 0 else 'withEffect'
            inner += (f'<p:par><p:cTn id="{nid()}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node_t}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                      f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                      f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{nid()}" dur="500"/><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')
        clicks += (f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                   f'{inner}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
           f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    el = etree.fromstring(xml); sld = slide._element
    old = sld.find(f'{{{P_NS}}}timing')
    if old is not None: sld.remove(old)
    ext = sld.find(f'{{{P_NS}}}extLst')
    (ext.addprevious(el) if ext is not None else sld.append(el))
    return len(groups)

# ------------------------------------------------------------------ code colouring (four colours only)
PY_KW = set('and as assert break class continue def del elif else except False finally for from global if import in is lambda None nonlocal not or pass raise return True try while with yield'.split())
SQL_KW = set('SELECT FROM WHERE GROUP BY ORDER HAVING JOIN LEFT RIGHT INNER OUTER ON AS AND OR NOT IN IS NULL LIMIT DISTINCT COUNT SUM AVG MIN MAX CASE WHEN THEN ELSE END DESC ASC UNION ALL CREATE TABLE INSERT INTO VALUES LIKE BETWEEN'.split())
COL = dict(kw=INK, str=ACCENT, num=BODY_C, com=SOFT, plain=INK)
def tokenize(line, lang):
    toks = []; i = 0
    while i < len(line):
        c = line[i]
        if (lang == 'python' and c == '#') or (lang == 'sql' and line.startswith('--', i)): toks.append((line[i:], 'com')); break
        if c in '"\'':
            j = line.find(c, i + 1); j = len(line) - 1 if j < 0 else j; toks.append((line[i:j + 1], 'str')); i = j + 1; continue
        m = re.match(r'[A-Za-z_][A-Za-z_0-9]*', line[i:])
        if m:
            w = m.group(0); kind = 'plain'
            if lang == 'python' and w in PY_KW: kind = 'kw'
            elif lang == 'sql' and w.upper() in SQL_KW: kind = 'kw'
            toks.append((w, kind)); i += len(w); continue
        m = re.match(r'\d+(\.\d+)?', line[i:])
        if m: toks.append((m.group(0), 'num')); i += len(m.group(0)); continue
        toks.append((c, 'plain')); i += 1
    return toks

# ------------------------------------------------------------------ the deck
class Deck:
    def __init__(self, sess, vocab=None, fig_dir=None, hero=None):
        self.sess = sess; self.vocab = vocab or []; self.fig_dir = Path(fig_dir) if fig_dir else None
        self.prs = Presentation(); self.prs.slide_width = Inches(W); self.prs.slide_height = Inches(H)
        self.n = 0; self.clicks = 0; self.animated = 0

    def _slide(self, dark=False, title=None, notes='', kicker=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[5]); self.n += 1
        bg = s.background.fill; bg.solid(); bg.fore_color.rgb = rgb(DARK_BG if dark else WHITE)
        t = s.shapes.title
        if title is None: t._element.getparent().remove(t._element)
        else:
            t.left, t.top, t.width, t.height = Inches(MX), Inches(0.5), Inches(W - 2 * MX), Inches(1.05)
            tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP; tf.margin_left = tf.margin_right = 0; tf.margin_top = Inches(0.02)
            sz = fit_size(title, W - 2 * MX, 1.0, 32, 27, HEAD, True)
            tf.text = ''; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = title; _set_font(r, HEAD, sz, True, DARK_TEXT if dark else INK)
        foot = f"Session {self.sess['no']}   {self.sess['title']}"
        add_text(s.shapes, MX, 7.02, 9.5, 0.28, foot, 11, False, DARK_SOFT if dark else SOFT, name='Footer')
        add_text(s.shapes, W - MX - 1.0, 7.02, 1.0, 0.28, str(self.n), 11, False, DARK_SOFT if dark else SOFT, align='r', name='Page')
        if notes: s.notes_slide.notes_text_frame.text = notes.strip()
        return s

    def title_slide(self):
        s = self._slide(dark=True, notes=self.sess.get('notes_title', '')); sh = s.shapes
        add_text(sh, 7.6, 0.55, 5.2, 4.4, f"{self.sess['no']:02d}", 230, True, '33383E', HEAD, align='r', anchor='t', name='Session number')
        add_text(sh, MX, 1.05, 8.0, 0.45, f"Session {self.sess['no']} of 24", 18, False, DARK_ACCENT, HEAD, italic=True)
        add_text(sh, MX, 1.7, 8.2, 2.5, self.sess['title'], 46, True, DARK_TEXT, HEAD, min_size=32)
        add_text(sh, MX, 4.55, 7.4, 1.4, self.sess['question'], 22, False, DARK_SOFT, min_size=17)
        add_text(sh, MX, 6.45, 9, 0.4, 'Jotham Martin Wambi, KLAi Professional Learning', 14, False, DARK_SOFT)
        return s

    def outcomes_slide(self):
        s = self._slide(title='By the end of today', notes=self.sess.get('notes_outcomes', '')); sh = s.shapes
        y = Y0 + 0.1; cw = 5.9
        rule(sh, MX, y - 0.05, MX + cw, y - 0.05, INK, 1.25)
        for o in self.sess['outcomes']:
            _, sz = add_text(sh, MX, y + 0.12, cw, 1.1, o, 24, False, INK, HEAD, min_size=19)
            hh = block_height_in(o, HEAD, False, sz, cw) + 0.3; y += hh; rule(sh, MX, y - 0.06, MX + cw, y - 0.06, RULE, 0.5)
        x0 = 7.6; wr = W - MX - x0
        add_text(sh, x0, Y0 - 0.3, wr, 0.3, 'How the two hours run', 14, False, SOFT)
        rule(sh, x0, Y0 + 0.05, x0 + wr, Y0 + 0.05, INK, 1.25); yy = Y0 + 0.12; rh = 0.66
        for name, mins in self.sess['shape']:
            lab = 'lab' in name.lower()
            add_text(sh, x0, yy + 0.1, wr - 1.3, rh - 0.1, name, 19, lab, INK if lab else BODY_C, min_size=15)
            add_text(sh, x0 + wr - 1.2, yy + 0.1, 1.2, rh - 0.1, f'{mins} min', 19, lab, ACCENT if lab else BODY_C, align='r')
            yy += rh; rule(sh, x0, yy, x0 + wr, yy, RULE, 0.5)
        return s

    def words_slide(self):
        if not self.vocab: return None
        s = self._slide(title='Vocabulary', notes=self.sess.get('notes_words', ''))
        rows = [['Word', 'Plain meaning', 'In a spreadsheet']] + [[t, d, (x or 'No equivalent. A new idea.')] for t, d, x in self.vocab]
        self._table(s, rows, MX, Y0 - 0.05, W - 2 * MX, Y1 - Y0 + 0.05, colw=[2.4, 6.0, 3.6], size=21, min_size=14, firstbold=True)
        return s

    # ---- booktabs table: no fills, header rule, hairline row rules
    def _table(self, s, rows, x, y, w, h, colw=None, size=20, min_size=14, header=True, hl_rows=(), firstbold=False):
        nr, nc = len(rows), len(rows[0])
        if not colw:
            minw = []; wt = []
            for ci in range(nc):
                words = [q for r in rows for q in str(r[ci]).split()]
                minw.append(max(text_width_pt(q, BODY, True, size) for q in words) / 72 + 0.36); wt.append(max(6.0, sum(len(str(r[ci])) for r in rows) / nr))
            spare = w - sum(minw); colw = [m + (spare * v / sum(wt) if spare > 0 else 0) for m, v in zip(minw, wt)]
        sc = w / sum(colw); colw = [c * sc for c in colw]
        for pad in (0.2, 0.12):
            sz = size
            while True:
                rh = [max(block_height_in(str(c), BODY, (header and ri == 0) or (firstbold and ci == 0), sz, colw[ci] - 0.3) for ci, c in enumerate(r)) + pad for ri, r in enumerate(rows)]
                if sum(rh) <= h or sz <= min_size: break
                sz -= 0.5
            if sum(rh) <= h + 0.05: break
        if sum(rh) > h + 0.05: raise FitError(f"table does not fit at {min_size}pt ({sum(rh):.2f}in > {h:.2f}in): {rows[0]}")
        extra = max(0.0, min(h - sum(rh), 0.25 * nr) / nr); rh = [r + extra for r in rh]
        gs = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(sum(rh))); gs.name = 'Table'
        tbl = gs.table; tblPr = tbl._tbl.tblPr; tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
        sid = tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
        if sid is not None: sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'      # No Style, No Grid
        for ci, cw in enumerate(colw): tbl.columns[ci].width = Inches(cw)
        for ri, r in enumerate(rows):
            tbl.rows[ri].height = Inches(rh[ri])
            for ci, c in enumerate(r):
                cell = tbl.cell(ri, ci); cell.margin_left = Inches(0.0 if ci == 0 else 0.15); cell.margin_right = Inches(0.1); cell.margin_top = cell.margin_bottom = Inches(0.07); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.background()
                tf = cell.text_frame; tf.word_wrap = True; tf.text = ''
                hd = header and ri == 0; run = tf.paragraphs[0].add_run(); run.text = str(c)
                _set_font(run, BODY, sz, bold=hd or (firstbold and ci == 0) or ri in hl_rows, color=INK if (hd or ri in hl_rows or (firstbold and ci == 0)) else BODY_C)
                self._cell_rule(cell, INK if hd else RULE, 15875 if hd else 6350)
        return gs, sz

    @staticmethod
    def _cell_rule(cell, color, wemu):
        A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ('lnL', 'lnR', 'lnT'):
            ln = etree.SubElement(tcPr, A + tag); ln.set('w', '0'); etree.SubElement(ln, A + 'noFill')
        ln = etree.SubElement(tcPr, A + 'lnB'); ln.set('w', str(wemu)); sf = etree.SubElement(ln, A + 'solidFill'); c = etree.SubElement(sf, A + 'srgbClr'); c.set('val', color)
        for f in tcPr.findall(A + 'noFill'): tcPr.remove(f); tcPr.append(f)

    def _group(self, s, name):
        g = s.shapes.add_group_shape(); g.name = name; return g

    def _picture(self, s, name, x, y, w, h, alt=None, align='c'):
        path = self.fig_dir / f"{name}.png"
        if not path.exists(): raise FileNotFoundError(f"figure {name!r} not found in {self.fig_dir}")
        from PIL import Image
        iw, ih = Image.open(path).size; ar = iw / ih
        pw, ph = (h * ar, h) if w / h > ar else (w, w / ar)
        px = x if align == 'l' else x + (w - pw) / 2
        p = s.shapes.add_picture(str(path), Inches(px), Inches(y if align == 'l' else y + (h - ph) / 2), Inches(pw), Inches(ph)); p.name = f'figure-{name}'
        p._element.nvPicPr.cNvPr.set('descr', alt or name.replace('_', ' '))
        return p, pw, ph

    def add(self, spec):
        fn = getattr(self, f"_t_{spec['type']}", None)
        if fn is None: raise ValueError(f"unknown slide type {spec['type']!r}")
        s = fn(spec); spec['_slide'] = s; return s

    # ---- a list of observations: plain paragraphs between hairlines (used beside figures, code and charts)
    def _notes_list(self, s, items, x, y, w, h, size=21, groups=None, label=None):
        sh = s.shapes; n = len(items)
        if label:
            add_text(sh, x, y - 0.32, w, 0.3, label, 14, False, SOFT);
        rule(sh, x, y, x + w, y, INK, 1.25); gap = 0.0; ph = min(1.5, h / n); top = y
        for i, tx in enumerate(items):
            g = self._group(s, f'note-{i+1}')
            add_text(g.shapes, x, top + 0.08, 0.35, ph - 0.1, str(i + 1), size, True, ACCENT, HEAD)
            add_text(g.shapes, x + 0.4, top + 0.08, w - 0.4, ph - 0.12, tx, size, False, INK, min_size=15)
            top += ph; rule(g.shapes, x, top, x + w, top, RULE, 0.5)
            if groups is not None: groups.append([g.shape_id])
        return top

    # ---- slide types
    def _t_concept(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        pts = sp.get('points', []); fig = sp.get('fig'); lead = sp.get('lead'); wide = False
        if fig:
            from PIL import Image as _I
            nm = fig['name'] if isinstance(fig, dict) else fig; _iw, _ih = _I.open(self.fig_dir / f'{nm}.png').size; wide = _iw / _ih > 1.8
        if fig and wide:
            y = Y0
            if lead: _, lsz = add_text(sh, MX, y, W - 2 * MX, 1.0, lead, 24, False, INK, HEAD, min_size=19); y += block_height_in(lead, HEAD, False, lsz, W - 2 * MX) + 0.1
            fh = min(2.5, (W - 2 * MX) / (_iw / _ih)); self._picture(s, nm, MX, y, W - 2 * MX, fh, sp.get('alt'), 'l'); y += fh + 0.2
            n = len(pts); gap = 0.4; cw = (W - 2 * MX - gap * (n - 1)) / max(n, 1)
            for i, p in enumerate(pts):
                x = MX + i * (cw + gap); rule(sh, x, y, x + cw, y, INK, 1.25)
                add_text(sh, x, y + 0.1, cw, Y1 - y - 0.1, [[(p['h'] + '. ', {'bold': True, 'color': INK}), (p['t'], {})]], 20, False, BODY_C, min_size=15)
        elif fig:
            wl = 5.8; y = Y0
            if lead: _, lsz = add_text(sh, MX, y, wl, 1.3, lead, 23, False, INK, HEAD, min_size=18); y += block_height_in(lead, HEAD, False, lsz, wl) + 0.2
            rule(sh, MX, y, MX + wl, y, INK, 1.25); n = len(pts); ph = (Y1 - y) / max(n, 1); ph = min(ph, 1.3)
            for p in pts:
                add_text(sh, MX, y + 0.08, wl, ph - 0.1, [[(p['h'] + '. ', {'bold': True, 'color': INK}), (p['t'], {})]], 20, False, BODY_C, anchor='t', min_size=15)
                y += ph; rule(sh, MX, y, MX + wl, y, RULE, 0.5)
            self._picture(s, nm, 6.95, Y0, W - MX - 6.95, Y1 - Y0, sp.get('alt'))
        else:
            y = Y0
            if lead: _, lsz = add_text(sh, MX, y, 10.5, 1.2, lead, 26, False, INK, HEAD, min_size=20); y += block_height_in(lead, HEAD, False, lsz, 10.5) + 0.35
            n = len(pts); rule(sh, MX, y, W - MX, y, INK, 1.25); ph = min(1.25, (Y1 - y) / max(n, 1))
            for p in pts:
                add_text(sh, MX, y + 0.12, 2.9, ph - 0.12, p['h'], 24, True, INK, HEAD, min_size=18)
                add_text(sh, MX + 3.2, y + 0.14, W - 2 * MX - 3.2, ph - 0.14, p['t'], 22, False, BODY_C, min_size=16)
                y += ph; rule(sh, MX, y, W - MX, y, RULE, 0.5)
        return s

    def _t_flow(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        st = sp['steps']; n = len(st); span = W - 2 * MX; colw = span / n; ly = Y0 + 0.75
        rule(sh, MX + 0.12, ly, W - MX - 0.12 - (colw - 0.5) * 0, ly, INK, 1.5)
        groups = []
        for i, p in enumerate(st):
            x = MX + i * colw; g = self._group(s, f'step-{i+1}'); hot = p.get('hot')
            node(g.shapes, x + 0.12, ly, 0.26, filled=bool(hot), color=ACCENT if hot else INK)
            _, hs = add_text(g.shapes, x, ly + 0.35, colw - 0.25, 0.95, p['h'], 24, True, ACCENT if hot else INK, HEAD, min_size=18)
            hh = block_height_in(p['h'], HEAD, True, hs, colw - 0.25) + 0.12
            add_text(g.shapes, x, ly + 0.35 + hh, colw - 0.25, 2.8 - hh, p['t'], 22, False, BODY_C, min_size=14)
            groups.append([g.shape_id])
        if sp.get('caption'):
            rule(sh, MX, 5.6, W - MX, 5.6, INK, 1.25)
            add_text(sh, MX, 5.75, W - 2 * MX, 1.0, sp['caption'], 22, False, INK, HEAD, min_size=17)
        if sp.get('animate', True): self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_cycle(self, sp):
        import math
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        st = sp['steps']; n = len(st); cx, cy = 3.9, (Y0 + Y1) / 2 + 0.05; rx, ry = 2.55, 1.85
        e = box(sh, cx - rx, cy - ry, 2 * rx, 2 * ry, fill=None, line=INK, lw=1.25, shape=MSO_SHAPE.OVAL)
        add_text(sh, cx - 1.1, cy - 0.4, 2.2, 0.8, sp.get('center', ''), 22, False, SOFT, HEAD, align='c', anchor='m', italic=True, min_size=16)
        groups = []; right = []
        for i in range(n):
            a = -math.pi / 2 + 2 * math.pi * i / n; px, py = cx + rx * math.cos(a), cy + ry * math.sin(a)
            g = self._group(s, f'step-{i+1}'); node(g.shapes, px, py, 0.26, filled=True, color=INK)
            lw_ = 2.2; side = 1 if math.cos(a) > 0.2 else (-1 if math.cos(a) < -0.2 else 0)
            lx = px + 0.25 if side == 1 else (px - 0.25 - lw_ if side == -1 else px - lw_ / 2); ly = py - 0.62 if side == 0 and math.sin(a) < 0 else (py + 0.2 if side == 0 else py - 0.2)
            add_text(g.shapes, lx, ly, lw_, 0.42, st[i]['h'], 22, True, INK, HEAD, align='r' if side == -1 else ('l' if side == 1 else 'c'), min_size=16)
            groups.append([g.shape_id])
        x0 = 7.9; wr = W - MX - x0; ph = (Y1 - Y0 - 0.1) / n
        rule(sh, x0, Y0 + 0.02, x0 + wr, Y0 + 0.02, INK, 1.25)
        for i, p in enumerate(st):
            g2 = self._group(s, f'note-{i+1}')
            add_text(g2.shapes, x0, Y0 + 0.08 + i * ph, wr, ph - 0.1, [[(p['h'] + '. ', {'bold': True, 'color': INK}), (p['t'], {})]], 19, False, BODY_C, min_size=14)
            rule(g2.shapes, x0, Y0 + 0.02 + (i + 1) * ph, x0 + wr, Y0 + 0.02 + (i + 1) * ph, RULE, 0.5)
            groups[i].append(g2.shape_id)
        if sp.get('animate', True): self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_compare(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        L, R = sp['left'], sp['right']; gap = 0.7; cw = (W - 2 * MX - gap) / 2; groups = []; bh = 3.7 if sp.get('verdict') else 4.9
        rule(sh, MX + cw + gap / 2, Y0 + 0.05, MX + cw + gap / 2, Y0 + bh, RULE, 0.75)
        for k, side in enumerate((L, R)):
            x = MX + k * (cw + gap); g = self._group(s, f'side-{k+1}')
            add_text(g.shapes, x, Y0, cw, 0.6, side['h'], 28, True, INK if k == 0 else ACCENT, HEAD, min_size=20)
            rule(g.shapes, x, Y0 + 0.62, x + cw, Y0 + 0.62, INK, 1.25)
            body = side['points']; ph = (bh - 0.7) / max(len(body), 1)
            for i, pt in enumerate(body):
                add_text(g.shapes, x, Y0 + 0.7 + i * ph + 0.04, cw, ph - 0.06, pt, 22, False, BODY_C, anchor='m', min_size=16)
                if i < len(body) - 1: rule(g.shapes, x, Y0 + 0.7 + (i + 1) * ph, x + cw, Y0 + 0.7 + (i + 1) * ph, RULE, 0.5)
            groups.append([g.shape_id])
        if sp.get('verdict'):
            v = self._group(s, 'verdict'); rule(v.shapes, MX, Y0 + bh + 0.3, W - MX, Y0 + bh + 0.3, INK, 1.25)
            add_text(v.shapes, MX, Y0 + bh + 0.42, W - 2 * MX, 0.95, sp['verdict'], 24, False, INK, HEAD, min_size=18); groups.append([v.shape_id])
        if sp.get('animate', True) and len(groups) > 1: self.clicks += animate(s, groups[1:]); self.animated += 1
        return s

    def _t_table(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', ''))
        rows = [sp['header']] + sp['rows']; h = Y1 - Y0 - (1.15 if sp.get('caption') else 0.1)
        self._table(s, rows, MX, Y0, W - 2 * MX, h, colw=sp.get('colw'), size=sp.get('size', 22), min_size=14, hl_rows=sp.get('hl', ()))
        if sp.get('caption'):
            rule(s.shapes, MX, Y1 - 0.85, W - MX, Y1 - 0.85, INK, 1.25); add_text(s.shapes, MX, Y1 - 0.75, W - 2 * MX, 0.75, sp['caption'], 21, False, INK, HEAD, min_size=15)
        return s

    def _t_code(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        lang = sp.get('lang', 'python'); code = sp['code'].rstrip('\n'); lines = code.split('\n'); out = sp.get('output'); says = sp.get('says', [])
        longest = max(len(l) for l in lines); avail = Y1 - Y0; ol_ = out.rstrip('\n').split('\n') if out else []
        def code_h(c): return len(lines) * c * 1.2 / 72 + 0.34
        def out_h_for(o=14): return (len(ol_) * o * 1.2 / 72 + 0.34) if out else 0
        layout = None
        if says:
            cw = 7.6; csz = 21.0
            while csz >= 14:
                if longest * 0.6 * csz / 72 + 0.5 <= cw and (not out or max(len(l) for l in ol_) * 0.6 * 14 / 72 + 0.4 <= cw) and code_h(csz) + (out_h_for() + 0.35 if out else 0) <= avail: layout = 'side'; break
                csz -= 0.5
        if layout is None:
            cw = W - 2 * MX; csz = 21.0; says_h = 1.05 if says else 0
            while csz >= 14:
                if longest * 0.6 * csz / 72 + 0.5 <= cw and code_h(csz) + (out_h_for() + 0.35 if out else 0) + (says_h + 0.2 if says else 0) <= avail: layout = 'stack'; break
                csz -= 0.5
        if layout is None and says:
            cw = 7.4; csz = 21.0; wr_ = W - MX - (MX + cw + 0.45)
            ow = (max(len(l) for l in ol_) * 0.6 * 13 / 72 + 0.4) if out else 0
            while csz >= 14:
                if longest * 0.6 * csz / 72 + 0.5 <= cw and code_h(csz) <= avail and ow <= wr_ and ((len(ol_) * 13 * 1.2 / 72 + 0.64) if out else 0) + 0.3 + len(says) * 0.72 <= avail: layout = 'split'; break
                csz -= 0.5
        if layout is None: raise FitError(f"code does not fit at 14pt ({longest} chars wide, {len(lines)} lines, output {len(ol_)} lines): {lines[0][:50]}")
        ch = code_h(csz); box(sh, MX, Y0, cw, ch, fill=None, line=INK, lw=1.0)
        tb = sh.add_textbox(Inches(MX + 0.22), Inches(Y0 + 0.18), Inches(cw - 0.3), Inches(ch - 0.3)); tb.name = 'Code'
        tf = tb.text_frame; tf.word_wrap = False; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if not line: r = p.add_run(); r.text = ' '; _set_font(r, MONO, csz, False, INK)
            for tok, kind in tokenize(line, lang): r = p.add_run(); r.text = tok; _set_font(r, MONO, csz, kind == 'kw', COL.get(kind, INK), italic=(kind == 'com'))
        out_group = None; y_next = Y0 + ch + 0.12
        if out:
            osz = 13 if layout == 'split' else 14; ox = (MX + cw + 0.45) if layout == 'split' else MX; ow_ = (W - MX - ox) if layout == 'split' else cw; oy = Y0 if layout == 'split' else y_next
            oh = (len(ol_) * osz * 1.2 / 72 + 0.64) if layout == 'split' else out_h_for(); og = self._group(s, 'output')
            add_text(og.shapes, ox, oy, 2.0, 0.26, 'Output', 13, False, SOFT)
            box(og.shapes, ox, oy + 0.28, ow_, oh - 0.28, fill=WELL, line=None)
            tb2 = og.shapes.add_textbox(Inches(ox + 0.2), Inches(oy + 0.28 + 0.1), Inches(ow_ - 0.3), Inches(oh - 0.4)); tf2 = tb2.text_frame; tf2.word_wrap = False
            tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0
            if max(len(l) for l in ol_) * 0.6 * osz / 72 + 0.4 > ow_: raise FitError(f"output line too wide for the slide: {max(ol_, key=len)[:60]!r}")
            for i, line in enumerate(ol_):
                p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph(); r = p.add_run(); r.text = line if line else ' '; _set_font(r, MONO, osz, False, INK)
            out_group = og.shape_id
            if layout != 'split': y_next += oh + 0.2
            else: split_y = oy + oh + 0.35
        groups = []
        if says and layout == 'side':
            x0 = MX + cw + 0.45; wr = W - MX - x0; n = len(says); ph = min(1.4, (Y1 - Y0 - 0.1) / n)
            groups_ = []; self._notes_list(s, says, x0, Y0 + 0.02, wr, ph * n, size=20, groups=groups_); groups = groups_
            if out_group: groups[-1].append(out_group)
        elif says and layout == 'split':
            x0 = MX + cw + 0.45; wr = W - MX - x0; yy = split_y if out else Y0 + 0.02; n = len(says); ph_ = min(0.95, (Y1 - yy) / n)
            groups_ = []; self._notes_list(s, says, x0, yy, wr, ph_ * n, size=18, groups=groups_); groups = groups_
            if out_group: groups[-1].append(out_group)
        elif says:
            n = len(says); gap = 0.45; wc = (W - 2 * MX - gap * (n - 1)) / n; hh = min(Y1 - y_next, 1.5)
            for i, tx in enumerate(says):
                g = self._group(s, f'says-{i+1}'); x = MX + i * (wc + gap); rule(g.shapes, x, y_next, x + wc, y_next, INK, 1.25)
                add_text(g.shapes, x, y_next + 0.08, 0.32, hh - 0.1, str(i + 1), 19, True, ACCENT, HEAD); add_text(g.shapes, x + 0.38, y_next + 0.08, wc - 0.38, hh - 0.1, tx, 19, False, INK, min_size=14)
                groups.append([g.shape_id])
            if out_group: groups[-1].append(out_group)
        elif out_group: groups.append([out_group])
        if sp.get('animate', True) and groups: self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_figure(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        look = sp.get('look', []); wf = 7.7 if look else W - 2 * MX; hf = Y1 - Y0 - (0.95 if sp.get('caption') else 0)
        p, pw, ph = self._picture(s, sp['fig'], MX, Y0, wf, hf, sp.get('alt'), 'l')
        groups = []
        if look:
            x0 = MX + wf + 0.5; wr = W - MX - x0; self._notes_list(s, look, x0, Y0 + 0.02, wr, min(4.6, 1.5 * len(look)), size=21, groups=groups)
        if sp.get('caption'):
            rule(sh, MX, Y1 - 0.8, W - MX, Y1 - 0.8, RULE, 0.75); add_text(sh, MX, Y1 - 0.72, W - 2 * MX, 0.72, sp['caption'], 20, False, INK, HEAD, italic=True, min_size=15)
        if sp.get('animate', True) and groups: self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_stat(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        st = sp['stats']; n = len(st); y = Y0; ph = min(1.15, (3.9) / n); rule(sh, MX, y, W - MX, y, INK, 1.25)
        for p in st:
            nsz = 48
            while nsz > 24 and text_width_pt(p['n'], HEAD, True, nsz) > 4.4 * 72: nsz -= 2
            add_text(sh, MX, y + 0.05, 4.6, ph - 0.05, p['n'], nsz, True, INK, HEAD, anchor='m', min_size=nsz)
            add_text(sh, MX + 5.0, y + 0.05, W - 2 * MX - 5.0, ph - 0.05, p['label'], 22, False, BODY_C, anchor='m', min_size=16)
            y += ph; rule(sh, MX, y, W - MX, y, RULE, 0.5)
        if sp.get('caption'): add_text(sh, MX, y + 0.3, W - 2 * MX, 1.0, sp['caption'], 23, False, INK, HEAD, min_size=17)
        return s

    def _t_steps(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        st = sp['steps']; n = len(st); wl = 7.9; ph = min(0.8, (Y1 - Y0) / n); y = Y0; rule(sh, MX, y, MX + wl, y, INK, 1.25)
        for i, tx in enumerate(st):
            add_text(sh, MX, y + 0.1, 0.5, ph - 0.1, str(i + 1), 24, True, ACCENT, HEAD, min_size=24)
            add_text(sh, MX + 0.6, y + 0.1, wl - 0.6, ph - 0.1, tx, 22, False, INK, min_size=15)
            y += ph; rule(sh, MX, y, MX + wl, y, RULE, 0.5)
        x0 = MX + wl + 0.6; wr = W - MX - x0
        rule(sh, x0, Y0, x0 + wr, Y0, INK, 1.25)
        add_text(sh, x0, Y0 + 0.12, wr, 0.3, 'Tool and time', 14, False, SOFT)
        add_text(sh, x0, Y0 + 0.45, wr, 1.2, sp.get('tool', 'Google Colab') + '\n' + sp.get('time', '40 minutes'), 21, True, INK, HEAD, min_size=16)
        rule(sh, x0, Y0 + 1.75, x0 + wr, Y0 + 1.75, RULE, 0.5)
        add_text(sh, x0, Y0 + 1.87, wr, 0.3, 'You are done when', 14, False, SOFT)
        add_text(sh, x0, Y0 + 2.2, wr, Y1 - Y0 - 2.2, sp['done'], 21, False, INK, min_size=15)
        return s

    def _t_trap(self, sp):
        s = self._slide(dark=True, title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        blocks = [('What it looks like', sp['looks'], DARK_ACCENT), ('Why it goes wrong', sp['why'], DARK_ACCENT), ('The fix', sp['fix'], DARK_TEXT)]
        gap = 0.55; cw = (W - 2 * MX - gap * 2) / 3; groups = []
        for i, (h, tx, col) in enumerate(blocks):
            x = MX + i * (cw + gap); g = self._group(s, f'trap-{i+1}')
            rule(g.shapes, x, Y0 + 0.1, x + cw, Y0 + 0.1, DARK_SOFT, 1.0)
            add_text(g.shapes, x, Y0 + 0.25, cw, 0.4, h, 18, True, col, HEAD)
            add_text(g.shapes, x, Y0 + 0.8, cw, Y1 - Y0 - 0.9, tx, 24, False, DARK_TEXT, min_size=16)
            groups.append([g.shape_id])
        if sp.get('animate', True): self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_check(self, sp):
        s = self._slide(title=sp.get('title', 'Check your understanding'), notes=sp.get('notes', '')); sh = s.shapes
        qs = sp['questions']; n = len(qs); ph = (Y1 - Y0) / n; y = Y0; groups = []
        for i, q in enumerate(qs):
            rule(sh, MX, y, W - MX, y, INK if i == 0 else RULE, 1.25 if i == 0 else 0.5)
            qh = ph * 0.46
            add_text(sh, MX, y + 0.1, 0.5, qh, str(i + 1), 26, True, ACCENT, HEAD, min_size=26)
            add_text(sh, MX + 0.6, y + 0.1, W - 2 * MX - 0.6, qh, q['q'], 24, True, INK, HEAD, min_size=16)
            g = self._group(s, f'answer-{i+1}')
            add_text(g.shapes, MX + 0.6, y + 0.1 + qh, W - 2 * MX - 0.6, ph - qh - 0.2, q['a'], 22, False, BODY_C, min_size=15)
            groups.append([g.shape_id]); y += ph
        rule(sh, MX, y, W - MX, y, RULE, 0.5)
        if sp.get('animate', True): self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_define(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        _, dsz = add_text(sh, MX, Y0, 10.8, 2.0, sp['definition'], 32, False, INK, HEAD, min_size=22); hh = block_height_in(sp['definition'], HEAD, False, dsz, 10.8) + 0.4
        y = Y0 + hh; gap = 0.8; cw = (W - 2 * MX - gap) / 2
        rule(sh, MX + cw + gap / 2, y, MX + cw + gap / 2, Y1, RULE, 0.75)
        for k, (lab, tx) in enumerate((('In our data', sp['example']), ('In a spreadsheet', sp.get('excel', 'No spreadsheet equivalent. This is a new idea.')))):
            x = MX + k * (cw + gap); rule(sh, x, y, x + cw, y, INK, 1.25)
            add_text(sh, x, y + 0.1, cw, 0.3, lab, 14, False, SOFT); add_text(sh, x, y + 0.45, cw, Y1 - y - 0.5, tx, 22, False, BODY_C, min_size=16)
        return s

    def _t_matrix(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        cells = sp['cells']; x0, y0, cw, chh = 3.0, Y0 + 0.95, 3.1, 1.75
        add_text(sh, x0, Y0 + 0.05, 2 * cw, 0.3, 'What the model said', 14, False, SOFT, align='c')
        add_text(sh, x0, Y0 + 0.4, cw, 0.4, sp.get('pos', 'Default'), 19, True, INK, HEAD, align='c'); add_text(sh, x0 + cw, Y0 + 0.4, cw, 0.4, sp.get('neg', 'No default'), 19, True, INK, HEAD, align='c')
        add_text(sh, MX, y0 + 0.2, 2.1, 1.35, 'What really happened:\n' + sp.get('pos', 'Default'), 16, False, SOFT, align='r', anchor='m', min_size=13)
        add_text(sh, MX, y0 + chh + 0.2, 2.1, 1.35, 'What really happened:\n' + sp.get('neg', 'No default'), 16, False, SOFT, align='r', anchor='m', min_size=13)
        pos = {'tp': (0, 0), 'fn': (0, 1), 'fp': (1, 0), 'tn': (1, 1)}; groups = []
        for k, (r, c) in pos.items():
            lab, val = cells[k]; err = k in ('fn', 'fp'); g = self._group(s, f'cell-{k}'); x = x0 + c * cw; y = y0 + r * chh
            box(g.shapes, x, y, cw, chh, fill=None, line=ACCENT if err else INK, lw=1.5 if err else 1.0)
            add_text(g.shapes, x + 0.15, y + 0.15, cw - 0.3, 0.85, str(val), 44, True, ACCENT if err else INK, HEAD, align='c', anchor='m', min_size=30)
            add_text(g.shapes, x + 0.15, y + 1.05, cw - 0.3, chh - 1.1, lab, 19, False, BODY_C, align='c', min_size=14); groups.append([g.shape_id])
        x1 = x0 + 2 * cw + 0.5; wr = W - MX - x1
        if sp.get('side'): rule(sh, x1, y0, x1 + wr, y0, INK, 1.25); add_text(sh, x1, y0 + 0.12, wr, 2 * chh - 0.1, sp['side'], 22, False, INK, HEAD, min_size=16)
        if sp.get('animate', True): self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_bar(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        cd = CategoryChartData(); cd.categories = sp['categories']
        for name, vals in sp['series']: cd.add_series(name, vals)
        wf = 8.0 if sp.get('look') else W - 2 * MX
        gf = sh.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED if not sp.get('horizontal') else XL_CHART_TYPE.BAR_CLUSTERED, Inches(MX), Inches(Y0), Inches(wf), Inches(Y1 - Y0 - (0.95 if sp.get('caption') else 0)), cd)
        gf.name = 'Chart'; ch = gf.chart; ch.has_legend = len(sp['series']) > 1
        if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(16)
        ch.font.size = Pt(16); ch.font.name = BODY
        pl = ch.plots[0]; pl.gap_width = 70; pl.has_data_labels = True; dl = pl.data_labels; dl.font.size = Pt(16); dl.font.bold = True; dl.number_format = sp.get('fmt', '#,##0'); dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
        for i, ser in enumerate(pl.series): ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb([INK, ACCENT, SOFT, MOSS][i % 4])
        va = ch.value_axis; va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb(RULE); va.major_gridlines.format.line.width = Pt(0.5); va.format.line.fill.background(); va.tick_labels.font.size = Pt(14)
        if sp.get('vmax'): va.maximum_scale = sp['vmax']
        va.minimum_scale = 0; ch.category_axis.tick_labels.font.size = Pt(15)
        groups = []
        if sp.get('look'):
            x0 = MX + wf + 0.5; wr = W - MX - x0; self._notes_list(s, sp['look'], x0, Y0 + 0.02, wr, min(4.4, 1.5 * len(sp['look'])), size=21, groups=groups)
        if sp.get('caption'):
            rule(sh, MX, Y1 - 0.8, W - MX, Y1 - 0.8, RULE, 0.75); add_text(sh, MX, Y1 - 0.72, W - 2 * MX, 0.72, sp['caption'], 20, False, INK, HEAD, italic=True, min_size=15)
        if sp.get('animate', True) and groups: self.clicks += animate(s, groups); self.animated += 1
        return s

    def _t_erd(self, sp):
        s = self._slide(title=sp['title'], notes=sp.get('notes', '')); sh = s.shapes
        tabs = sp['tables']; pos = sp['layout']; cw, rh = 3.45, 1.95; xs = [MX, MX + 4.2, MX + 8.4]; ys = [Y0 + 0.05, Y0 + 2.4]; ids = {}
        for t, (c, r) in zip(tabs, pos):
            x, y = xs[c], ys[r]; box(sh, x, y, cw, rh, fill=None, line=INK, lw=1.0)
            add_text(sh, x + 0.15, y + 0.06, cw - 0.3, 0.4, t['name'], 20, True, INK, HEAD); rule(sh, x, y + 0.5, x + cw, y + 0.5, INK, 0.75)
            cols_txt = [[(('' if not c_.startswith('*') else ''), {}), (c_.lstrip('*'), {'bold': c_.startswith('*'), 'color': INK if c_.startswith('*') else BODY_C})] for c_ in t['cols']]
            add_text(sh, x + 0.15, y + 0.58, cw - 0.3, rh - 0.65, cols_txt, 18, False, BODY_C, min_size=13); ids[t['name']] = (x, y)
        for a, b, label in sp.get('links', []):
            (x1, y1), (x2, y2) = ids[a], ids[b]
            if abs(y1 - y2) < 0.1: rule(sh, x1 + cw if x1 < x2 else x1, y1 + 1.0, x2 if x1 < x2 else x2 + cw, y2 + 1.0, ACCENT, 1.75)
            else: rule(sh, x1 + cw / 2, y1 + rh if y1 < y2 else y1, x2 + cw / 2, y2 if y1 < y2 else y2 + rh, ACCENT, 1.75)
        if sp.get('caption'):
            rule(sh, MX, Y1 - 0.55, W - MX, Y1 - 0.55, RULE, 0.75); add_text(sh, MX, Y1 - 0.5, W - 2 * MX, 0.5, sp['caption'], 18, False, INK, HEAD, italic=True, min_size=14)
        return s

    def _t_summary(self, sp):
        s = self._slide(title=sp.get('title', 'Summary'), notes=sp.get('notes', '')); sh = s.shapes
        pts = sp['points']; n = len(pts); ph = min(1.35, (Y1 - Y0 - 1.0) / n); y = Y0; rule(sh, MX, y, W - MX, y, INK, 1.25)
        for tx in pts:
            add_text(sh, MX, y + 0.1, 10.8, ph - 0.1, tx, 26, False, INK, HEAD, anchor='m', min_size=17); y += ph; rule(sh, MX, y, W - MX, y, RULE, 0.5)
        if sp.get('next'): add_text(sh, MX, y + 0.35, W - 2 * MX, 0.6, [[('Next session. ', {'bold': True, 'color': ACCENT}), (sp['next'], {})]], 22, False, INK, min_size=16)
        return s

    def save(self, path): self.prs.save(str(path))
