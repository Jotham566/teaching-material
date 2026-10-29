"""Write 04 Facilitator Materials/Playback and Rehearsal Checklist.txt from the built decks.  python3 make_playback_checklist.py"""
import json
from pathlib import Path
from pptx import Presentation
HERE = Path(__file__).resolve().parent; ROOT = HERE.parents[2]
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
data = json.loads((HERE / 'docs_data.json').read_text(encoding='utf-8'))
out = ["PLAYBACK AND REHEARSAL CHECKLIST", "",
 "Part 1, in PowerPoint on the machine you will teach from. Open each deck, choose Slide Show, and play every animated slide from the first click to the last. Confirm the parts appear in the order listed and nothing overlaps. Editing view shows the finished slide, so only Slide Show proves the build. Tick each line.",
 "Part 2, read the notes aloud once per session, timing yourself. The estimate is the speaking time of the notes alone at 130 words a minute, before pauses. The notes are written to be read as they stand: the first paragraph before any click, then one paragraph per click, then any wrap-up.",
 "Part 3, run every lab solution in Colab (run_all_labs.py in this folder does this on a machine with the full library set) and tick it here. Sessions 1 to 4 also have Sheets tasks: work them once in Google Sheets and check yourself against Sheets Lab Answers.docx.", ""]
for f in sorted((ROOT / '02 Class Slides').glob('[0-2][0-9] - *.pptx')):
    n = int(f.name[:2]); s = data['sessions'][n - 1]; prs = Presentation(str(f)); lines = []
    for i, sl in enumerate(prs.slides, 1):
        t = sl._element.find(P + 'timing')
        if t is None: continue
        ms = [x for x in t.iter(P + 'cTn') if x.get('nodeType') == 'mainSeq'][0]; k = len(ms.find(P + 'childTnLst'))
        title = sl.shapes.title.text_frame.text if sl.shapes.title is not None else f'slide {i}'
        lines.append(f"  [ ] Slide {i}: {title}, {k} click{'s' if k != 1 else ''}")
    out += [f"SESSION {n}: {s['title']}", f"  Notes read aloud: about {s['minutes_spoken']} minutes.", "  [ ] Deck opens in PowerPoint and Slide Show runs from slide 1", *lines,
            "  [ ] Notes read aloud once, wording changed where it does not sound like you", *(["  [ ] Lab solution notebook runs top to bottom in Colab"] if s['code'] or n == 1 else []), *(["  [ ] Sheets tasks done once in Google Sheets with the workbook (answers are in Sheets Lab Answers.docx)"] if n <= 4 else []), *(["  [ ] Rubric and presentation tasks read through once (Sheets Lab Sheets.docx, Session 24)"] if n == 24 else []), ""]
out += ["TOOLS THAT NEED YOUR ACCOUNT OR A DOWNLOAD", "  [ ] Colab opens and runs a cell with a Google account", "  [ ] data.zip hosted, DATA_ZIP_URL set (or students will upload it)", "  [ ] sentence-transformers model downloads in Colab (Session 21 and 22)", "  [ ] Small language model licence read on its model card, and the model runs in Colab (Session 22, optional)", "  [ ] GitHub repository created and pushed; Streamlit Community Cloud app deploys (Session 23)", "  [ ] TensorFlow Playground opens (Session 20)", ""]
(ROOT / '04 Facilitator Materials' / 'Playback and Rehearsal Checklist.txt').write_text("\n".join(out), encoding='utf-8'); print('written', len(out), 'lines')
