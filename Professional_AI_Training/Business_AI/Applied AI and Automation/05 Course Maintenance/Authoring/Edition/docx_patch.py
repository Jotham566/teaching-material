"""Paragraph-text patch between the builder's Word documents and the teaching edition.

  python3 docx_patch.py make  BASE_FOLDER EDITED_FOLDER PATCH_FILE
  python3 docx_patch.py apply BASE_FOLDER PATCH_FILE OUT_FOLDER

Folders hold the four documents by their edition names (Course Overview.docx, Learner
Workbook.docx, Facilitator Guide.docx, Law Firm Implementation Playbook.docx). The patch
lists every body paragraph and table-cell paragraph whose text differs, by position.
Applying writes the new text into the first run and empties the other runs, which keeps
the paragraph style.
"""
import sys, json
from pathlib import Path
from docx import Document

NAMES = ['Course Overview.docx', 'Learner Workbook.docx', 'Facilitator Guide.docx', 'Law Firm Implementation Playbook.docx']

def paras(doc):
    for i, p in enumerate(doc.paragraphs): yield f'p{i}', p
    for ti, t in enumerate(doc.tables):
        for ri, row in enumerate(t.rows):
            seen = set()
            for ci, cell in enumerate(row.cells):
                if id(cell._tc) in seen: continue
                seen.add(id(cell._tc))
                for pi, p in enumerate(cell.paragraphs): yield f't{ti}.{ri}.{ci}.{pi}', p

def set_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]: r.text = ''
    else: p.add_run(text)

def make(base, edited, out):
    patch = {}
    for n in NAMES:
        a = dict((k, p.text) for k, p in paras(Document(str(Path(base) / n))))
        b = dict((k, p.text) for k, p in paras(Document(str(Path(edited) / n))))
        assert list(a) == list(b), n
        patch[n] = {k: b[k] for k in a if a[k] != b[k]}
    Path(out).write_text(json.dumps(patch, indent=1, ensure_ascii=False), encoding='utf-8')
    print({n: len(v) for n, v in patch.items()})

def apply(base, patch_file, out):
    patch = json.loads(Path(patch_file).read_text(encoding='utf-8')); Path(out).mkdir(parents=True, exist_ok=True)
    for n in NAMES:
        d = Document(str(Path(base) / n))
        for k, p in paras(d):
            if k in patch[n]: set_text(p, patch[n][k])
        d.save(str(Path(out) / n))
    print('documents written to', out)

if __name__ == '__main__':
    (make if sys.argv[1] == 'make' else apply)(*sys.argv[2:5])
