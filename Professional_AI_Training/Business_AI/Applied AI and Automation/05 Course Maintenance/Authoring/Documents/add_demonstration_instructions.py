"""Append the maintained session demonstration instructions to the patched core Guide.
Input must be the core guide emitted by docx_patch.py, never the already assembled Guide.
"""
from pathlib import Path
import json,sys
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
P=ROOT/'03 Learner Materials/Practice Files'
TR=ROOT/'04 Facilitator Materials/Teaching Resources'

def mark(p,name,i):
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(i));a.set(qn('w:name'),name)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(i));p._p.insert(0,a);p._p.append(b)
def link(p,label,anchor=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor:h.set(qn('w:anchor'),anchor)
    else:h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'145F7A');pr.append(co);r.append(pr)
    t=OxmlElement('w:t');t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=label;r.append(t);h.append(r);p._p.append(h)
def styles(d):
    for name,base,size,bold in [('Demo Body','Normal',11,False),('Demo Small','Normal',9,False),('Demo Code','Normal',9,False),('Demo Heading','Heading 2',12,True),('Demo Chapter','Heading 1',16,True)]:
        if name not in d.styles:d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s=d.styles[name];s.base_style=d.styles[base];s.font.name='Calibri' if name!='Demo Code' else 'Courier New';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.line_spacing=1.0;s.paragraph_format.space_after=Pt(2)
        if name in ('Demo Heading','Demo Chapter'):s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(5)
        else:s.paragraph_format.keep_with_next=False
        s.paragraph_format.widow_control=True
    d.styles['Demo Small'].paragraph_format.line_spacing=1.0
    d.styles['Demo Small'].paragraph_format.space_after=Pt(2)
    d.styles['Demo Chapter'].paragraph_format.page_break_before=True

def append(path):
    d=Document(path)
    assert not any(p.text=='Demonstration instructions' for p in d.paragraphs),'Use the core Guide as input.'
    data=json.loads((HERE/'demonstrations.json').read_text());styles(d)
    def para(t='',style='Demo Body'):return d.add_paragraph(t,style)
    def heading(t):return para(t,'Demo Heading')
    # Link each existing session's summary directly to its detailed route without adding another summary document.
    found=0
    for p in list(d.paragraphs):
        if p.text=='Demonstration and guided application':
            found+=1;link(p,'  Full steps',anchor=f'demo{found:02d}')
    assert found==12
    link(d.paragraphs[2],'  Open demonstration index',anchor='demo_index')
    h=para('Demonstration instructions','Demo Chapter');mark(h,'demo_index',900)
    para('These instructions turn each session’s demonstration into an executable sequence. I operate the tools and complete the working records with the learner during class. The earlier session tables remain the 120-minute plan; the instructions below explain how to carry out its worked and guided blocks.')
    para('I use the session links below or Full steps in a session’s demonstration heading. I keep the current Guide, the matching deck and a working record open. Prompt files contain only text to copy; their names match the prompt IDs printed here.')
    for x in data:
        p=para();link(p,f"Session {x['session']:02d}  {x['title']}",anchor=f"demo{x['session']:02d}")
    heading('Build on existing AI use')
    para('The learner already uses Claude, ChatGPT and similar assistants. I begin each demonstration with its professional decision and ask how they currently approach it. The step-by-step tool instructions below are my private preparation reference; I do not narrate routine typing, account access or uploads unless help is requested.')
    para('I ask for a short explanation or prediction before teaching a mechanism. If the reasoning is sound, I connect it briefly to the diagram and use the remaining block for the harder case. Familiar terminology alone does not establish understanding. We still explain unfamiliar AI and ML mechanisms in sequence, with concrete legal and tax examples and no assumed mathematics or coding.')
    para('I pre-open sources and run checks before class. One purposeful live run plus inspected evidence is usually sufficient; saved repetitive or slow runs stay clearly labelled. Each session ends with a decision or reusable working record for the firm. The deeper challenge uses the existing guided block and adds no homework or extra core time.')
    heading('Folders and working copies')
    para('Practice means 03 Learner Materials/Practice Files inside the course folder. Teaching means 04 Facilitator Materials/Teaching Resources. Each file link below opens the existing local source. I copy blank records to my session output folder before typing, so the course templates stay reusable.')
    para('Course folder: '+str(ROOT),'Demo Small')
    para('Prompt files: '+str(TR/'Demo Materials'),'Demo Small')
    para('The launcher saves each run to Downloads/KLAi Demo Runs with a date/time and demo name. The printed OUTPUT line identifies the exact folder. I copy the prompt, raw reply and learner observations there. I do not save new run outputs over the checked historical evidence or into the course templates.')
    heading('Prepare the teaching Mac')
    para('1. In Finder, open Applications then Utilities then Terminal. Copy the following cd command, including its quotes, and press Return. This changes the working folder; it does not run a model.')
    para('cd "'+str(P)+'"','Demo Code')
    para('2. Run the preflight command. It prints the installed runtimes, Swift OCR runner, Claude CLI and source locations. A missing optional tool selects the saved-output route for its session; I do not install accounts or packages during class.')
    para('./facilitator_demo.command preflight','Demo Code')
    para('3. To list all routes, run the following command. In each session, I replace the demo name with the exact command printed in that section. I wait for the shell prompt to return before reading the output.')
    para('./facilitator_demo.command --list','Demo Code')
    para('4. I open the folder printed after OUTPUT in Finder using Go > Go to Folder, paste the path and select Return. I open the named text or JSON result in a text editor. A Python traceback, STOP message or held result is a finding to inspect, not a successful demonstration.')
    para('5. If a command is still waiting after two minutes, I press Control-C once, retain any output already saved, and use the named saved route. Large OCR regeneration and account setup happen before class. I do not fabricate a completed result.')
    heading('Use the chat and notebook routes')
    para('Claude web: I open https://claude.ai in my existing teaching account, choose a new chat, select + then Add files or photos, choose the listed local files and select Open. I confirm their names appear before pasting the prompt and submitting. Dragging the files into the chat is an equivalent upload route. I avoid a project containing extra evidence when a source-scope comparison needs a fresh context. CSV/TXT/JSON are sufficient here; an Excel add-in or XLSX feature is unnecessary.')
    para('Google notebook: I open the relevant existing notebook below. In Sources I select the checkbox for the intended file and deselect other sources. If the source is absent, I choose Add source > Upload files, select the exact local file and wait for processing. I paste the prompt in the chat box. On the completed answer I select a numbered citation and read the supporting passage in context. A citation proves neither the inference nor current applicability.')
    p=para();link(p,'Fictional legal and tax retrieval notebook',url='https://notebook.google.com/notebook/2fa0c6ce-9d08-436a-801a-a905c7067692')
    p=para();link(p,'Public law research notebook',url='https://notebook.google.com/notebook/467e5157-8f0a-4f85-b22f-c92b5a7f174c')
    para('Local workflow page: in Edge I press Command-O, choose Practice/04 Automation and Tax/workflow_demo.html and open it. workflow_engine.js must remain beside the HTML. The page states that it is a local deterministic simulation. Reload restores the visible default form; Reset exercise resets engine state only. Its mock acknowledgement is not an external filing or payment.')
    heading('Describe and save what actually happened')
    para('Before a run I state the purpose and ask for a prediction. After it I inspect the evidence with the learner, then record the conclusion and its limits. I label each example live model output, saved model output with date, scripted host test, deterministic simulation or constructed exercise. If a model gives a correct answer I retain the success; I do not keep altering the prompt to force a failure. A revision after feedback is saved separately from the first attempt.')
    para('Before sharing my screen I close unrelated mail, client files and browser tabs. All required sources here are fictional or public. No demonstration requires signing up, buying a plan, changing security settings, connecting employer systems, publishing a Zap or sending a real communication.')
    heading('Interface references checked on 28 October 2026')
    for label,url in [('Claude file upload instructions','https://support.claude.com/en/articles/8241126-upload-files-to-claude'),('Google notebook chat and source citations','https://support.google.com/gemininotebook/answer/16179559'),('AI by Zapier configuration and Preview','https://help.zapier.com/hc/en-us/articles/8496342944013-Use-AI-by-Zapier-to-analyze-and-return-data')]:
        p=para(style='Demo Small');link(p,label,url=url)
    for x in data:
        n=x['session'];h=para(f"Session {n:02d} demonstration instructions",'Demo Chapter');mark(h,f'demo{n:02d}',900+n);link(h,'  Index',anchor='demo_index')
        para(x['title'],'Demo Heading');para('Purpose. '+x['purpose']);para('Primary tools. '+x['tool'])
        heading('Begin with the decision')
        para(x['executive_task']);para('What I listen for. '+x['executive_check'])
        heading('Fit the demonstration into the class')
        plans=json.loads((HERE/'session_delivery.json').read_text());blocks=plans[n-1]['blocks']
        para('Worked demonstration: slides '+blocks[2]['slides']+' ('+str(blocks[2]['end_minute']-blocks[2]['start_minute'])+' minutes). Guided work: slides '+blocks[4]['slides']+' ('+str(blocks[4]['end_minute']-blocks[4]['start_minute'])+' minutes). Return to the firm decision on slides '+blocks[5]['slides']+'. The challenge fits within the guided block of the existing 120-minute plan.')
        heading('Open these files')
        for rel in x['files']:
            fp=ROOT/rel;assert fp.exists(),rel
            label=('Practice/' + rel.split('Practice Files/',1)[1]) if 'Practice Files/' in rel else ('Teaching/'+rel.split('Teaching Resources/',1)[1])
            p=para(style='Demo Small');link(p,label,url=fp.as_uri())
        heading('Before the learner arrives')
        for t in x['prepare']:para(t)
        if x['commands']:
            para('Before class, run these from the prepared Terminal folder. Each prints its output location. Use the saved result to avoid waiting in class; identify model runs and scripted tests separately.')
            outputs={
                's02_vectors':'Open Embedding Selftest.txt. Inspect shape, norms, repeatability and the actual cosine pairs.',
                's05_ocr':'Open Original Source.txt, Synthetic Scan.txt and Degraded Scan - moderate.txt beside their images. Apple Vision OCR.json records engine, configuration and source hashes.',
                's06_baseline':'Open s06_baseline.txt. K3 and K8 are expected historical failures; this is not the repaired route.',
                's06_reviewed':'Open Retrieval Tests.json. Read passed/total and the detail for a selected check; the recorded target is 19/19.',
                's06_query':'Open s06_query.txt as JSON text. Inspect status, selected_source, hits and mandatory_context.',
                's06_missing':'Open s06_missing.txt. Expect no_source, empty hits and no substitute context.',
                's09_host':'Open s09_host.txt. Seven scripted checks should pass; no model was called.',
                's09_protocol':'Open s09_protocol.txt. Three scripted format/grounding checks should pass; no model was called.',
                's09_agent':'Open Agent Trace.json. Read actual_model_output, host_decision, actual_tool_result and the final result.',
                's09_injection':'Open Agent Trace.json. Inspect actual behaviour and the host boundary; do not infer an unrequested forbidden action was denied.',
                's09_unknown':'Open Agent Trace.json. Expect a missing-source request for CONTRACT-99 without substituting CONTRACT-01.',
                's09_onecall':'Open Agent Trace.json and compare lookup_calls with max_lookup_calls. A model using only one call does not exercise a second-call refusal.',
                's10_reconcile':'Open Result.json. Inspect small and larger results, source identities, periods and the conditional balances.',
                's10_controls':'Open s10_controls.txt. All twelve scripted checks should show pass: true; this is a local engine test, not agent behaviour.',
                's11_ocr_ladder':'Open OCR Quality Summary.txt and Quality Results.csv. The live Apple Vision sample is three selected passages under four conditions; source images and hashes are in Apple Vision OCR.json.',
                's12_costs':'Open Result.json. Inspect values and units; compare the main and stress cases without converting capacity into automatic cash savings.'}
            for cmd in x['commands']:
                para('./facilitator_demo.command '+cmd,'Demo Code')
                para(outputs[cmd],'Demo Small')
        heading('Run the demonstration')
        for i,t in enumerate(x['steps'],1):para(f'{i}. {t}')
        for item in x['prompts']:
            heading('Copy prompt '+item['id'])
            p=para(style='Demo Small');link(p,'Open '+item['id']+'.txt',url=(TR/'Demo Materials'/(item['id']+'.txt')).as_uri());p.paragraph_format.keep_with_next=True
            p=para(item['text']);p.paragraph_format.keep_together=True
        if x['optional']:
            heading('Optional builder or research extension')
            for t in x['optional']:para(t)
        review_heading=heading('Check the result before accepting it')
        if n==10:review_heading.paragraph_format.page_break_before=True
        for i,t in enumerate(x['check'],1):para(f'{i}. {t}')
        heading('Ask and explain')
        for q,a in x['discuss']:
            para('I ask: '+q);para('I listen for: '+a)
        heading('If the demonstration does not go to plan')
        for problem,action in x['rescue']:
            p=para();p.add_run(problem+'. ').bold=True;p.add_run(action)
        # Precise saved route under every session, independent of account availability.
        saved={1:'Observed Runs/Corrections 2026-10-28/Contract Retests/R1 Revision Response.txt',2:'Observed Runs/Mac Rehearsal 2026-10-27/06 Neural Embedding',3:'Observed Runs/Corrections 2026-10-28/Contract Retests',4:'Observed Runs/Corrections 2026-10-28/Contract Notebook/Revision After Feedback - Raw UI.txt',5:'Observed Runs/Observed OCR Output.txt',6:'Observed Runs/Corrections 2026-10-28/Retrieval Retest.json',7:'Observed Runs/Corrections 2026-10-28/Research Retest/Notebook Retest - Raw UI.txt',8:'Checked Contract Answers.txt',9:'Observed Runs/Corrections 2026-10-28/Agent Retests - Structured/Normal.json',10:'Observed Runs/Observed Labelled Reconciliation.txt',11:'Observed Runs/Observed OCR Output.txt',12:'../Law Firm Implementation Playbook.docx'}[n]
        sp=(TR/saved).resolve()
        if n==11:sp=P/'02 Data and Retrieval/OCR Experiment/summary.txt';saved='Practice/02 Data and Retrieval/OCR Experiment/summary.txt'
        if n==12:sp=P/'05 Firm Adoption/Cost and Effort Example.csv';saved='Practice/05 Firm Adoption/Cost and Effort Example.csv'
        assert sp.exists(),sp
        p=para('Saved/reference route: ','Demo Small');link(p,saved,url=sp.as_uri())
        para('I identify the recorded producer/date or label the reference authored, mark the live step unexecuted, and inspect its evidence with the learner.')
        heading('Save and close')
        para(x['record']);para(x['transition'])
    d.save(path)
    print('Added twelve linked demonstration sections to',path)
if __name__=='__main__':append(Path(sys.argv[1]))
