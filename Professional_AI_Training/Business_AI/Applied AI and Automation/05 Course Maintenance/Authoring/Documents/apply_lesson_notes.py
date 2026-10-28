"""Apply the read-aloud speaker notes in lesson_notes.json to the class decks.

Run on any machine with python-pptx:
  python3 apply_lesson_notes.py "<folder with the 12 deck files>"

Each slide's note is exactly the script in lesson_notes.json, nothing added. Sources, click counts and the
paragraph convention for animated slides are kept out of the notes; make_presenter_backup.py writes them to
04 Facilitator Materials/Teaching Resources/Presenter Notes Backup.txt.
Run it after any rebuild of the decks, because a rebuild regenerates the notes from the builder's own text.
"""
import sys, json
from pathlib import Path
from pptx import Presentation

here = Path(__file__).resolve().parent
notes = json.loads((here / 'lesson_notes.json').read_text(encoding='utf-8'))
folder = Path(sys.argv[1])
total = 0
for f in sorted(folder.glob('[0-1][0-9] - *.pptx')):
    n = f.name[:2]
    prs = Presentation(str(f))
    sess = notes[n]
    assert len(sess) == len(prs.slides), (f.name, len(sess), len(prs.slides))
    for i, slide in enumerate(prs.slides, 1):
        slide.notes_slide.notes_text_frame.text = sess[str(i)].strip()
        total += 1
    prs.save(str(f))
    print(f.name, len(prs.slides))
print('slides updated:', total)
