"""Build aligned course documents from legal_tax_course.py."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from course_paths import ROOT,MAINTENANCE,SOURCE
from legal_tax_course import COURSE,SOURCES,TIMING,CASE,CLAUSE
from definitions import D
B=MAINTENANCE/'Authoring/.build/legal_tax_rebuild';OUT=B/'Documents';OUT.mkdir(exist_ok=True)
INDEX=json.loads((B/'Practice Index.json').read_text())
def doc(title,subtitle):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.62);sec.left_margin=sec.right_margin=Inches(.7)
 for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.05;d.styles['Normal'].paragraph_format.space_after=Pt(5)
 for n,size in [('Title',23),('Heading 1',18),('Heading 2',12),('Heading 3',11)]:
  d.styles[n].font.size=Pt(size);d.styles[n].font.bold=True;d.styles[n].paragraph_format.space_before=Pt(8);d.styles[n].paragraph_format.space_after=Pt(5)
 d.styles['Subtitle'].font.size=Pt(11)
 d.add_paragraph(title,'Title');d.add_paragraph(subtitle,'Subtitle')
 d.core_properties.author='Jotham Martin Wambi';d.core_properties.title=title
 f=sec.footer.paragraphs[0];f.text='KLAi Professional Learning   ';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 f.runs[0].font.size=Pt(9)
 return d
def head(d,t):d.add_paragraph(t,'Heading 2')
def page(d,t):d.add_page_break();d.add_paragraph(t,'Heading 1')
def para(d,t):return d.add_paragraph(t)
def bullet(d,t):d.add_paragraph(t,'List Bullet')
def tab(d,headers,rows,widths=None,size=10):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Inches(w)
 vals=[headers]+rows
 for i,row in enumerate(vals):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,v in enumerate(row):
   c=cells[j];c.text=str(v);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if widths:c.width=Inches(widths[j])
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1C3144' if i==0 else ('F2F5F7' if i%2 else 'FFFFFF'));pr.append(sh)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar);b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.02
    for rr in p.runs:rr.font.name='Arial';rr.font.size=Pt(size);rr.font.bold=i==0;rr.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
  trpr=t.rows[i]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 d.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
def save(d,name):
 for el in (d.element,d.styles.element):
  for x in list(el.iter(qn('w:pBdr'))):x.getparent().remove(x)
 p=OUT/name;d.save(p);print(p)
def locate(name):return INDEX.get(name,name)

def ranges(nums):
 if not nums:return '—'
 groups=[];start=prev=nums[0]
 for n in nums[1:]:
  if n==prev+1:prev=n
  else:groups.append(str(start) if start==prev else f'{start}–{prev}');start=prev=n
 groups.append(str(start) if start==prev else f'{start}–{prev}')
 return ', '.join(groups)
g=doc('Facilitator Guide','Applied AI and Automation for legal and tax leadership')
para(g,'I will teach through one fictional customer-side matter. We will build understanding, predict what a tool might do, inspect actual results and make a proposed firm decision. All required work happens in class. This guide is my teaching plan; the slide notes provide the explanation, question, expected response and help for a partly correct answer.')
head(g,'The matter');para(g,CASE)
head(g,'The teaching agreement');para(g,'I will distinguish the base reorganisation, Variant A share sale and Variant B assignment. The firm acts for Victoria, the customer. Unknown ownership facts and the missing signed schedule remain unknown. Public judgments concern different parties and must not be presented as deciding our fictional matter.')
head(g,'Before we start');para(g,'Use the entry diagnostic before showing the explanations. Record reasons and misconceptions. Confirm the actual teaching account, test the selected demonstrations and save the observed outputs. Read Teaching Resources/Demonstration Preflight.txt. Prepared examples, local control simulations and real AI runs carry different labels.')
head(g,'Pacing and optional content');para(g,'Each session has its own 120-minute plan, including a break and ten-minute buffer. Optional slides appear after the essential closing check and are marked in their footer. They are not required within two hours. Ask the written questions at the main teaching landmarks; do not turn every support slide into a separate discussion. If we need more time, agree it and revisit the point in class. No homework is required.')
head(g,'Resources');para(g,'Learner Workbook contains working prompts, not the answer key. Teaching Resources/Expected Findings.txt gives authored model results and transfer answers. Actual saved tool outputs are separately labelled. Use the progress record to distinguish independent understanding from an explanation that needed support.')
for c in COURSE:
 n=c['session'];page(g,f"Session {n:02} {c['title']}");para(g,c['aim']);head(g,'Result');para(g,c['result'])
 head(g,'Recall before explaining');para(g,c['slides'][1]['question']);para(g,'Expected: '+c['slides'][1]['answer'])
 tab(g,['Block','Min','Slides','What I will do'],[[b['block'],b['minutes'],ranges(b['slides']) or '—',('Ask the recall question before the answer.' if i==0 else 'Explain the essential concepts; pause at the landmarks below.' if i==1 else 'Show the first example and retain its actual output.' if i==2 else 'Break.' if i==3 else 'Let the learner direct the changed example and verify evidence.' if i==4 else 'Record the firm decision and ask the transfer questions.' if i==5 else 'Recover setup time or revisit a misconception.') ] for i,b in enumerate(c['run_sheet'])],[1.35,.45,1.0,4.1],size=9.5)
 head(g,'Essential teaching landmarks')
 inner=c['slides'][2:len(c['core_slides'])-4]
 for j in range(0,len(inner),4):
  group=inner[j:j+4];nums=[c['slides'].index(x)+1 for x in group];para(g,'Slides '+ranges(nums)+': '+'; '.join(x['title'] for x in group)+'.')
 if c['optional_slides']:para(g,'Optional after the closing check: slides '+ranges(c['optional_slides'])+'. Use only for a useful question or agreed extra time.')
 page(g,f'Session {n:02} Application and answers')
 head(g,'Demonstration and guided application')
 for j,t in enumerate(c['demo'],1):para(g,str(j)+'. '+t)
 head(g,'A good result contains');para(g,c['expected'])
 head(g,'Transfer questions before showing the key')
 for q in c['understanding_check']['questions']:bullet(g,q)
 para(g,'Expected: '+c['understanding_check']['expected'])
 head(g,'Misconceptions to watch for')
 for m in c['misconceptions']:para(g,m['wrong']+': '+m['response'])
 head(g,'Decision and recovery');para(g,c['decision']);para(g,c['risk']);para(g,'If the connected tool is unavailable, use the labelled example, mark that step unexecuted and preserve time for the explanation. The preflight file names the account-dependent tasks.')
 head(g,'Open before class');para(g,'; '.join(locate(f) for f in c['files']))
save(g,'Facilitator Guide.docx')

w=doc('Learner Workbook','Working notes for our sessions')
para(w,'We will use this workbook together during class. We will explain ideas, inspect sources and tool outputs, and record decisions that may help your firm. No assignments are required between sessions. The working prompts are for discussion; We will discuss each case before comparing our conclusions.')
head(w,'The recurring matter');para(w,CASE)
head(w,'What we are working towards');para(w,'A clearer understanding of AI and machine learning, practical habits for choosing and checking tools, and a proposal for a bounded organisational pilot. The proposal will name the problem, information, owners, quality tests, costs, risks and next decision.')
head(w,'Our working record');para(w,'Record an observation, the source or output supporting it, and the decision it changes. Proposed firm actions remain subject to the firm’s authority. Fictional documents, constructed examples, local simulations and observed AI runs are labelled separately.')
tab(w,['No.','Question'],[[c['session'],c['aim']] for c in COURSE],[.65,6.25],size=10)
for c in COURSE:
 page(w,f"Session {c['session']:02} {c['title']}");para(w,c['aim'])
 head(w,'What we will do together')
 for t in c['demo']:bullet(w,t)
 head(w,'Working record')
 tab(w,['Observation or question','Source or result to inspect','Implication or next action'],[['','',''],['','',''],['','','']],[2.3,2.3,2.3],size=10)
 head(w,'Explain it to a fellow partner')
 for q in c['understanding_check']['questions']:bullet(w,q)
 head(w,'Decision record');tab(w,['Proposed change','Owner to confirm','Evidence or condition'],[['','','']],[3.0,1.6,2.3],size=10)
 para(w,'Record what remains uncertain and what we will revisit during class. A proposed action here is not firm approval.')
plain={
 'AI':'Systems infer outputs from inputs; the field includes several different methods.',
 'Data analysis':'Examine data to answer a question, such as why invoices and receipts differ.',
 'Data science':'Use computing, statistics and subject knowledge to study data and evaluate useful results.',
 'Machine learning':'Learn a mapping or pattern from examples rather than manually writing every rule.',
 'Deep learning':'Machine learning using multiple neural-network layers to learn representations.',
 'Generative AI':'Produce content, such as a draft, from learned patterns and current context.',
 'Model':'A learned computational component; the product also includes files, tools and controls.',
 'Training':'Adjust numerical parameters using examples and an objective.',
 'Inference':'Use the model on an input; this need not update its parameters.',
 'Foundation model':'A broadly trained model that can be adapted or used for different tasks.',
 'Token':'A model input/output unit, which may be a word, subword or punctuation.',
 'Embedding':'A numerical representation used to compare items; similarity is not legal equivalence.',
 'Context window':'The amount of information the model can process in a request, subject to product rules.',
 'Attention':'Input-dependent weighting that combines information from relevant positions.',
 'Hallucination':'Generated content that is false, unsupported or inconsistent with the evidence.',
 'RAG':'Retrieve source material and supply it to a model when it generates an answer.',
 'Fine-tuning':'Further training of a model, distinct from simply supplying a document as context.',
 'OCR':'Recover text from a page image; verify words, layout and relationships.',
 'Chunking':'Split material into useful pieces while preserving source identity and necessary context.',
 'Index':'A structure that supports finding relevant records efficiently.',
 'NER':'Identify spans such as organisation names and dates; roles need further analysis.',
 'Automation':'Software follows defined work steps, with or without AI.',
 'Agent':'A system in which a model can select permitted actions based on observations.',
 'API':'A defined way for software to request an operation and receive a response.',
 'MCP':'A protocol for connecting clients to tools and context; it does not grant business permission.',
 'Human review':'A defined person checks evidence and can stop or change the dependent action.',
 'Evaluation':'Compare observed results with agreed criteria on representative cases.',
 'JSON':'Text using named fields and values; valid structure does not establish true content.',
 'CAG':'A specialised approach that preloads a manageable collection and reuses supported cached computation.',
 'AI-native organisation':'Our working term for deliberately designing suitable work with AI, maintained information and clear responsibilities.'}
items=list(plain.items())
for i in range(0,len(items),8):
 page(w,'Reference concepts '+str(i//8+1))
 for term,t in items[i:i+8]:head(w,term);para(w,t)
 para(w,'These concise explanations are reference aids. We will use the fuller definitions, sources and examples in class. Source links are in Public Sources.txt.')
save(w,'Learner Workbook.docx')

p=doc('Law Firm Implementation Playbook','Working decisions and a completed teaching example')
para(p,'We will turn the course into a proposal for one bounded firm pilot. This document includes a worked fictional example so that the required decisions are concrete. We will adapt the reasoning to the firm during class. Completing it does not approve production use.')
head(p,'Candidate scope');para(p,'Internal precedent search and partner-review drafting are the worked example. Exclude external delivery, cross-matter access, filing and payment. A tax evidence workflow is an alternative, not an additional requirement.')
head(p,'Our proposal');tab(p,['Decision','Working note'],[['Problem and present process',''],['Selected workflow and exclusions',''],['Sponsor and practice owner',''],['Sources and information owner','']],[2.0,4.9])
head(p,'Evidence before a purchase');para(p,'Require a representative trial, actual data terms, permissions, review effort, total costs and an exit route. Use Vendor Due Diligence.csv. A vendor demo and a classroom success do not establish firm-wide reliability.')
page(p,'Information and workflow design')
tab(p,['Design area','Decision to make'],[['Source readiness','Systems, modalities, formats, authority and missing evidence'],['Access','Who may search which matters and for what purpose'],['Preparation','Extraction checks, provenance, chunks and update policy'],['Workflow','Inputs, allowed tools, review and excluded actions'],['Operations','Logs, monitoring, support, recovery and changes']],[1.65,5.25])
head(p,'Our source and review record');tab(p,['Source or stage','Owner to confirm','Decision and evidence'],[['','',''],['','',''],['','','']],[2.1,1.8,3.0])
head(p,'A consequential action');para(p,'The reviewer sees the exact version and supporting evidence. Rejection stops the dependent action; changed fields require renewed review. The execution host enforces permitted operations, recipients and budgets.')
page(p,'Evaluation and acceptance')
para(p,'First measure the current manual process. Include ordinary and difficult cases, not just the demonstration clause. Separate source selection, answer fidelity, access, tool execution and professional usefulness. Define critical failures before collecting results.')
tab(p,['Measure','Evidence'],[['Quality','Correct source, complete meaning, uncertainty and useful result'],['Critical errors','Restricted disclosure, invented material terms or unapproved actions'],['Total effort','Preparation, checking, corrections, escalation and support'],['Value','Turnaround, capacity or better client decisions'],['Change','Performance after source, model, access or workflow updates']],[1.5,5.4])
head(p,'Our acceptance decision');tab(p,['Threshold or stop condition','Who decides','Evidence needed'],[['','',''],['','','']],[3.0,1.5,2.4])
para(p,'A classroom 0–2 rubric supports discussion. It is not a statistically validated production threshold. Set representative test coverage and acceptance conditions with the firm.')
page(p,'Budget and benefit assessment')
para(p,'Use actual quotations and observed review time before committing funds. The figures below are fictional assumptions for learning, not market prices.')
tab(p,['Assumption','Teaching example'],[['Monthly licence, usage and maintenance','UGX 120,000 + 30,000 + 90,000 = 240,000'],['Initial setup','UGX 600,000 once'],['First three months','UGX 600,000 + 3 × 240,000 = 1,320,000'],['Main capacity case','40 tasks × (30 − 20) minutes = 400 minutes released/month'],['Stress case','40 tasks × (30 − 33) = −120 minutes/month']],[2.7,4.2])
para(p,'Released time is not automatically cash saving. Specify how the firm will use capacity, whether demand exists, what review adds and which costs are actually avoidable. Keep currencies separate unless an exchange rate is explicitly sourced.')
head(p,'Our cost and benefit record');tab(p,['Item','Evidence or assumption','Amount or time'],[['Licence and usage','',''],['Setup and support','',''],['Review and correction','',''],['Benefit that can be tested','','']],[1.7,3.5,1.7])
page(p,'Risk ownership and use rules')
para(p,'The following is a proposed teaching register. Replace role labels with agreed owners. Privacy, professional confidentiality, quality and action authority need separate assessment. Use Uganda Legal Frame.txt for the sourced legal questions.')
tab(p,['Risk','Control and test','Owner role'],[['Cross-matter disclosure','Filter before retrieval; test forbidden queries','Information / technical'],['Lost exception','Source-linked review; changed-condition test','Practice reviewer'],['Action after rejection','Version-bound gate; decline/stale-approval test','Workflow owner']],[1.8,3.6,1.5],size=10)
head(p,'Proposed use rule');para(p,'Use only the approved product/account and source collection. Keep matter permissions. A qualified reviewer checks material claims before external use. No new recipient, filing, payment or broader access is authorised by a chat instruction. Report errors through the agreed route. Stop, preserve evidence and assess duties after a material incident. Retest before authorised restart.')
head(p,'Our decisions');tab(p,['Approved scope and account','Incident contact','Review date'],[['','','']],[3.3,1.8,1.8])
page(p,'Worked pilot proposal')
para(p,'Fictional example: internal precedent search and review drafting for M-01. The problem is time spent finding current clauses and correcting incomplete summaries. The proposal excludes external action and other matters.')
head(p,'Approach and evidence');para(p,'Compare one available configured product with manual review using a current permitted collection. Start with ten representative ordinary and difficult tasks, recording review time and critical errors. This is an initial test set, not proof of broad reliability. Expand coverage before a wider decision.')
head(p,'Proposed acceptance conditions');para(p,'No restricted disclosure or unapproved action. No accepted draft with an omitted operative exception or invented material term. Meet the quality threshold agreed before testing and improve total review effort against the measured baseline. Critical failures stop expansion.')
head(p,'Responsibilities');para(p,'Partner sponsor: scope and investment. Practice reviewer: meaning and release. Information owner: sources, permissions and retention. Technical/workflow owner: enforced controls, support and changes. Names require the firm’s agreement.')
head(p,'Resources and decision');para(p,'Use the budget assessment with actual quotations. Stop or revise if correction effort erases the benefit. Retain a manual fallback. The sponsor decides whether evidence supports a bounded pilot, revision or stopping; this example does not grant authority.')
page(p,'People and staged implementation')
para(p,'Explain the proposed change to affected staff, train reviewers and maintain professional skills. Identify how an approved route will be easier to use than informal workarounds. Include feedback and correction ownership.')
tab(p,['Period after approval','Proposed owner','Decision evidence'],[['Day 0–30','Practice and information owners','Scope, sources, permissions and baseline: ready to shadow?'],['Day 31–60','Workflow owner and reviewers','Quality, effort and failure tests: ready for bounded use?'],['Day 61–90','Sponsor and risk owner','Value and remaining exposure: expand, revise or stop?']],[1.1,2.1,3.7],size=10)
head(p,'Dates and people to agree');tab(p,['Milestone','Actual owner','Actual date'],[['Readiness review','',''],['Shadow review','',''],['Pilot decision','','']],[2.6,2.6,1.7])
para(p,'These are relative planning stages, not a promise to deploy in 90 days. No follow-up appointment has been booked. Agree actual dates around the firm’s capacity and dependencies.')
page(p,'Partner handover and optional investment transfer')
para(p,'The handover should contain the problem, selected workflow, sources, results, limitations, costs, owners, controls and next decision. Another partner should be able to understand why the proposal is justified without attending every lesson.')
head(p,'Handover checklist')
for t in ['A bounded problem and the manual baseline','Current source and permission record','Actual test outputs and acceptance decision','Total cost and review-effort assumptions','Named operational and professional owners','Stop, incident and change arrangements','A dated next decision approved by the firm']:bullet(p,t)
head(p,'Applying the ideas to an investment');para(p,'Ask what user problem is evidenced, what models and data rights the business depends on, how performance was evaluated, what human-review costs remain and how it handles changes or provider exit. Request evidence that could disprove the investment story. Use Investment AI Due Diligence.txt as an optional discussion aid.')
save(p,'Law Firm Implementation Playbook.docx')

o=doc('Applied AI and Automation','Executive learning through legal and tax work')
para(o,'12 sessions of 2 hours • 24 hours total • flexible dates and frequency')
para(o,'We will build a practical understanding of AI and machine learning, use and evaluate tools, and develop a proposal for a first organisational pilot. One fictional customer-side legal and tax matter connects the sessions. All required activities happen together in class; no homework is required.')
tab(o,['No.','Focus'],[[c['session'],c['title']] for c in COURSE],[.6,6.3],size=10)
head(o,'What you will be able to do');para(o,'Explain core mechanisms and limitations; prepare and assess information; check research and contract outputs; distinguish workflows and agents; identify risks; and sponsor a bounded pilot using evidence, owners, costs and acceptance conditions.')
head(o,'How we will work');para(o,'Short explanations and editable illustrations lead into guided demonstrations and changed-fact discussions. An entry/exit comparison and in-class explanations help us adjust the pace. Optional technical detail remains available after each session’s essential slides.')
head(o,'Before we begin');para(o,'No coding prerequisite. Bring a laptop and access to one agreed AI assistant. I will confirm the available account and any builder access before class. We will use fictional and public material; production use and real client data require separate firm decisions.')
save(o,'Course Overview.docx')
