from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/client'; OUT.mkdir(parents=True,exist_ok=True)

def new(title,subtitle):
    d=Document(); sec=d.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    sec.header_distance=Inches(.2); sec.footer_distance=Inches(.25)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        st=d.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(6)
    d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.line_spacing=1.08
    d.styles['Title'].font.size=Pt(25); d.styles['Title'].font.bold=True
    d.styles['Heading 1'].font.size=Pt(17); d.styles['Heading 1'].font.bold=True
    d.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    d.styles['Heading 2'].font.size=Pt(12); d.styles['Heading 2'].font.bold=True
    d.styles['Heading 2'].paragraph_format.space_before=Pt(8)
    d.styles['Subtitle'].font.size=Pt(11)
    d.core_properties.author='Jotham Martin Wambi'; d.core_properties.title=title
    d.add_paragraph(title,'Title'); d.add_paragraph(subtitle,'Subtitle')
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run('KLAi  |  '); r.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
    return d

def save(d,path):
    for root in (d.element,d.styles.element):
        for border in list(root.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    d.save(path)

def p(d,t,boldlead=None):
    a=d.add_paragraph()
    if boldlead and t.startswith(boldlead):
        a.add_run(boldlead).bold=True; a.add_run(t[len(boldlead):])
    else:a.add_run(t)
    return a

def h(d,t): return d.add_paragraph(t,'Heading 1')
def h2(d,t): return d.add_paragraph(t,'Heading 2')
def bullet(d,t):
    a=d.add_paragraph(t,style='List Bullet'); a.paragraph_format.space_after=Pt(4); return a

def page(d,t):
    a=h(d,t); a.paragraph_format.page_break_before=True

def table(d,heads,rows,widths):
    t=d.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,x in enumerate(heads):t.rows[0].cells[i].text=x
    for row in rows:
        c=t.add_row().cells
        for i,x in enumerate(row):c[i].text=str(x)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        avoid=OxmlElement('w:cantSplit');trPr.append(avoid)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'17384B' if ri==0 else ('F1F5F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for a in c.paragraphs:
                a.paragraph_format.space_after=Pt(3);a.paragraph_format.line_spacing=1.02
                for r in a.runs:
                    r.font.name='Arial';r.font.size=Pt(10.5)
                    if ri==0:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def link(d,label,url,size=9.5):
    a=d.add_paragraph();a.paragraph_format.space_after=Pt(4)
    hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),d.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));pr.append(sz)
    col=OxmlElement('w:color');col.set(qn('w:val'),'175B85');pr.append(col)
    under=OxmlElement('w:u');under.set(qn('w:val'),'single');pr.append(under)
    r.append(pr);text=OxmlElement('w:t');text.text=label;r.append(text);hyp.append(r);a._p.append(hyp)

D=new('Practical AI and Automation for Finance','Study roadmap  |  KLAi Professional Learning')
p(D,'This course is for experienced finance professionals who want to spend less time on repetitive work. You will use AI to review complex workbooks, prepare reports and presentations, and build an automation for a task you currently do manually.')
p(D,'Proposed format: Ten 90-minute sessions, two or three times per week. No prior coding is required. We will agree the schedule and project after the short needs assessment.',boldlead='Proposed format: ')
h(D,'Course outline')
table(D,['Sessions and focus','Practical work and outcome'],[
('1  Effective and safe AI use','Compare Claude, ChatGPT and Codex on finance tasks. Learn how to give clear instructions, spot unreliable answers and protect confidential information. Save useful prompts and a data handling checklist.'),
('2 to 3  AI with Excel','Ask Claude to review complex formulas, find inconsistencies across sheets and investigate unusual movements. Check its findings against the source records before accepting changes. Produce a workbook review with cell references and unresolved questions.'),
('4  AI assisted finance reporting','Use AI to investigate variances, question scenario assumptions and draft management commentary. Check the figures and remove explanations the data does not support. Produce a report you can defend in a management meeting.'),
('5  Presentation creation and quality','Use Claude or Codex to create an editable finance presentation. Replace generic wording, repetitive layouts and decorative visuals with clear findings and useful charts. Produce a five-slide briefing and a checklist for reviewing future decks.'),
('6  Excel and connected workflow automation','Use AI to write and debug a recurring Excel task. Practise connecting apps with Zapier or Power Automate, including filters, approval steps and failed runs. Build a small automation you can use as a starting point for the capstone.'),
('7  Power BI reporting','Use AI to help create a finance dashboard in Power BI, including DAX measures and management KPIs. Check that the calculations and filters give the right results. Produce an interactive report for reviewing financial performance.'),
('8  Claude and Codex computer use','Practise having Claude or Codex carry out a task in a browser or desktop app, such as updating a workbook or revising slides. Learn when to intervene and which actions need your approval.'),
('9 to 10  Workflow automation capstone','Choose a repetitive manual task and build an automation for it. Test a new batch of data and common failures, then compare the time taken with your current method. Deliver the automation and instructions for running and maintaining it.')],[1.85,5.15])

page(D,'Projects and working practices')
h2(D,'What you should leave with')
for x in [
'A workbook review that identifies issues and shows where each finding came from.',
'An editable finance presentation that sounds like you and explains the figures clearly.',
'A finance dashboard and prompts you can reuse for analysis and reporting.',
'A working automation for a repetitive task, with instructions for handling errors.'
]:bullet(D,x)
h2(D,'Your automation project')
p(D,'Start with a task that takes up time every week or month. Examples include consolidating exports, preparing reconciliation exceptions or collecting invoice details for review. We will map the current process and choose the software that fits it. You will build the automation with guidance and test it on a new batch, including missing information and duplicate records. By the end, you should be able to run it yourself, check the result and recover from a failed run. We will measure the time saved, including review and corrections.')
h2(D,'Presentations that sound and look like your work')
p(D,'Give the AI a clear audience, the decision to be made and an example of the writing or design you want. Review the outline before generating the deck. Edit vague claims and stock phrases into specific findings you would say yourself. Check chart scales, text size and layout, and remove visuals that add no information. Read the slides aloud and review them at presentation size before sharing.')
h2(D,'Handling workplace information')
p(D,'We will begin with sample data. The course covers what can be shared with an AI tool, how to remove identifying details and how to check storage and sharing settings. You will also practise spotting instructions hidden in documents and limiting what an AI tool can access. Reports and financial actions remain subject to your usual review and approval process.')
h2(D,'Before the first session')
p(D,'Complete the online needs assessment with your preferred class times, AI tools, Excel setup and the task you would most like to automate. We will agree the project and confirm any software setup with you before class.')
link(D,'Complete the needs assessment  |  forms.gle/FGdVz8Eswpc1nvRD8','https://forms.gle/FGdVz8Eswpc1nvRD8',11)
qr=D.add_paragraph()
qr.paragraph_format.space_before=Pt(3)
qr.add_run().add_picture(str(ROOT/'assets/finance_needs_assessment_qr.png'),width=Inches(1.15))
p(D,'Scan to complete the needs assessment.')
save(D,OUT/'Practical_AI_for_Finance_Roadmap.docx')
print(OUT/'Practical_AI_for_Finance_Roadmap.docx')
