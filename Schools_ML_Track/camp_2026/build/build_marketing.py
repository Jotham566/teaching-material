from common import ROOT,ROADMAP
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, Color, white
from reportlab.lib.pagesizes import A4
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from xml.sax.saxutils import escape

OUT=ROOT/'marketing'; OUT.mkdir(exist_ok=True)
import os
# Arial on the Mac. Liberation Sans has the same metrics, so layout is unchanged where Arial is not installed.
_FONT_SETS=[('/System/Library/Fonts/Supplemental/','Arial.ttf','Arial Bold.ttf'),('/usr/share/fonts/truetype/liberation2/','LiberationSans-Regular.ttf','LiberationSans-Bold.ttf'),('/usr/share/fonts/truetype/liberation/','LiberationSans-Regular.ttf','LiberationSans-Bold.ttf')]
fonts,_reg,_bold=next(x for x in _FONT_SETS if os.path.exists(x[0]+x[1]))
for name,file in [('Camp',_reg),('CampBold',_bold)]:pdfmetrics.registerFont(TTFont(name,fonts+file))
pdfmetrics.registerFontFamily('Camp',normal='Camp',bold='CampBold',italic='Camp',boldItalic='CampBold')
NAVY=HexColor('#142B42');INK=HexColor('#183044');TEAL=HexColor('#007E83');PURPLE=HexColor('#6751AA');PALE=HexColor('#EDF3F6')
W,H=A4
BIO='Jotham Martin Wambi • Main tutor'
BIO_BODY='Data scientist and ML/AI engineer; AI & ML lecturer at IUEA. Master’s, Yamaguchi University, Japan; incoming PhD student in Applied Informatics at Osaka Metropolitan University, Japan (April 2027).'
CONTACT='+256 750 990 718 | jothammartin.wambi@gmail.com'
LINKEDIN='linkedin.com/in/jothammartinwambi'
SCHEDULE='Three core classes weekly. ML adds one foundation lab weekly in weeks 1–4. Other days: project practice.'


def para(c,text,x,y,w,size=11,color=INK,bold=False,leading=None):
 p=Paragraph(text,ParagraphStyle('x',fontName='CampBold' if bold else 'Camp',fontSize=size,leading=leading or size*1.35,textColor=color))
 _,h=p.wrap(w,1000);p.drawOn(c,x,y-h);return y-h

def line(c,text,x,y,size=11,color=INK,bold=False):
 c.setFillColor(color);c.setFont('CampBold' if bold else 'Camp',size);c.drawString(x,y,text)

def art(c,track,x,y,w,h,accent):
 c.setFillColor(HexColor('#203F59'));c.roundRect(x,y,w,h,12,fill=1,stroke=0)
 if track=='Data Analysis':
  line(c,'MAKE THE DATA CLEAR',x+13,y+h-22,9,white,True)
  vals=[.40,.80,.62]
  for i,v in enumerate(vals):
   c.setFillColor([HexColor('#7DCBCD'),HexColor('#F7C66C'),HexColor('#A3B1EB')][i]);c.roundRect(x+18+i*43,y+37,25,v*58,3,fill=1,stroke=0)
  for i,t in enumerate(['Book','Pen','Juice']):line(c,t,x+15+i*43,y+22,8,white)
  line(c,'Illustrative units sold',x+13,y+7,8,HexColor('#C9DCE9'))
 else:
  line(c,'LEARN FROM EXAMPLES',x+13,y+h-22,9,white,True)
  c.setStrokeColor(HexColor('#6FBABD'));c.setLineWidth(1)
  c.line(x+18,y+27,x+w-16,y+27);c.line(x+18,y+27,x+18,y+h-38)
  pts=[(28,36),(47,44),(62,58),(82,65),(100,78),(122,86)]
  for dx,dy in pts:c.setFillColor(HexColor('#F7C66C'));c.circle(x+dx,y+dy,3,fill=1,stroke=0)
  c.setStrokeColor(HexColor('#91E0D3'));c.setLineWidth(2);c.line(x+24,y+34,x+w-24,y+h-39)
  line(c,'A simple prediction pattern',x+13,y+10,8,HexColor('#C9DCE9'))

WA_URL='https://wa.me/256750990718'
RELEVANCE='Build a practical foundation for university study in computing, engineering, business, economics, statistics and data science.'
TOOLS={'Data Analysis':'Python • Jupyter • NumPy • pandas • matplotlib • seaborn', 'Machine Learning':'Python • Jupyter • NumPy • pandas • scikit-learn'}
OUTCOMES={
 'Data Analysis':['Analyse data with Python and pandas','Clean messy records and keep an audit trail','Visualise patterns and explain the evidence','Explore data with NumPy and seaborn','Merge tables and check every total','Present your individual analysis project'],
 'Machine Learning':['Write basic Python programs','Load, clean and explore datasets','Build simple prediction models','Test on data the model has not seen','Explain errors and model limitations','Present your individual ML project']}
PHASES={
 'Data Analysis':[('1–2','Write and run your first Python programs'),('3–4','Clean data, compare groups and explain charts'),('5–6','Explore patterns and merge tables carefully'),('7–8','Build and present your individual project')],
 'Machine Learning':[('1–2','Write Python and practise the foundations'),('3–4','Explore data and understand the key maths'),('5–6','Build models and test whether they work'),('7–8','Build and present your individual project')]}
weeks=[('1–2','Python + guided foundation labs','Python fundamentals in Jupyter'),('3–4','Data handling + maths/statistics labs','pandas, cleaning and charts'),('5–6','Regression, classification and evaluation','NumPy, seaborn, merges and result checks'),('7–8','Individual model, testing, AI literacy & showcase','Individual analysis, AI literacy & showcase')]

def make_qr():
 from reportlab.graphics.barcode.qr import QrCodeWidget
 from PIL import Image,ImageDraw
 code=QrCodeWidget(WA_URL);code.qr.make();modules=code.qr.modules
 border=4;scale=12;n=len(modules)
 im=Image.new('RGB',((n+2*border)*scale,)*2,'white');draw=ImageDraw.Draw(im)
 for row,values in enumerate(modules):
  for col,filled in enumerate(values):
   if filled:draw.rectangle(((col+border)*scale,(row+border)*scale,(col+border+1)*scale-1,(row+border+1)*scale-1),fill='black')
 im.save(OUT/'WhatsApp_QR.png')

def cta(c,y):
 # A scan starts an enquiry. It does not promise a confirmed booking.
 c.drawImage(str(OUT/'WhatsApp_QR.png'),W-112,y-72,72,72)
 c.linkURL(WA_URL,(W-112,y-72,W-40,y),relative=0)
 y=para(c,'<b>Reserve your place</b>',36,y,W-164,15,TEAL)
 y=para(c,'<link href="'+WA_URL+'">WhatsApp/Call: +256 750 990 718</link><br/>jothammartin.wambi@gmail.com',36,y-4,W-164,10)
 y=para(c,'Scan to WhatsApp. Confirm dates and booking terms before payment.',36,y-5,W-164,9)
 return y

def fees(c,y):
 c.setFillColor(NAVY);c.roundRect(36,y-63,W-72,63,8,fill=1,stroke=0)
 line(c,'UGX 1,000,000',49,y-23,22,white,True)
 line(c,'Entire programme, per learner',49,y-40,10,white)
 line(c,'Same fee for groups of 1 to 5 learners',49,y-54,9,white)
 line(c,'Venue',303,y-19,9,HexColor('#B8DFDE'),True)
 line(c,'The National ICT Innovation Hub, Nakawa',303,y-34,11,white,True)
 line(c,'Dates and times confirmed with each group.',303,y-48,9,white)
 return y-63

def flyer(track):
 accent=TEAL if track=='Data Analysis' else PURPLE
 fname=OUT/(track.replace(' ','_')+'_Flyer.pdf');c=canvas.Canvas(str(fname),pagesize=A4)
 c.setTitle(('Python for Data Analysis' if track=='Data Analysis' else track)+' holiday camp');c.setAuthor('Jotham Martin Wambi')
 c.setFillColor(NAVY);c.rect(0,H-182,W,182,fill=1,stroke=0)
 line(c,'HOLIDAY DATA AND AI CAMP',36,H-28,10,HexColor('#B8DFDE'),True)
 line(c,'DECEMBER 2026 INTAKE',36,H-45,9,white)
 title=['Python for','Data Analysis'] if track=='Data Analysis' else track.split(' ')
 line(c,title[0],36,H-84,32,white,True);line(c,' '.join(title[1:]),36,H-120,32,white,True)
 tagline='Turn raw data into clear findings.' if track=='Data Analysis' else 'Build prediction models. Learn how to test them.'
 para(c,tagline,36,H-134,290,13,white)
 art(c,track,W-205,H-169,168,128,accent)
 c.setFillColor(accent);c.rect(0,H-208,W,26,fill=1,stroke=0)
 line(c,'8 WEEKS • ONLY 15 PLACES • MAX. 5 LEARNERS PER GROUP',36,H-199,10.5,white,True)
 y=para(c,'<b>For S.4–S.6 students and vacists.</b>',36,H-222,W-72,11)
 y=para(c,'Own laptop required. No coding experience needed. Choose one track.',36,y-4,W-72,10.5)
 schedule='24 classes: three weekly. Other days: project practice.' if track=='Data Analysis' else '24 classes + 4 labs: three classes weekly; an extra lab in weeks 1–4. Other days: practice.'
 y=para(c,schedule,36,y-4,W-72,10)
 y=para(c,'<b>TOOLS YOU WILL USE</b>  '+TOOLS[track],36,y-9,W-72,10,accent)
 y-=19;line(c,'YOUR EIGHT WEEK ROADMAP',36,y,10,accent,True);y-=10
 for i,(week,desc) in enumerate(PHASES[track]):
  c.setFillColor(PALE);c.roundRect(36,y-24,W-72,24,4,fill=1,stroke=0)
  line(c,'WEEKS '+week,46,y-16,9,accent,True);line(c,desc,131,y-16,10.2,INK,i==3)
  y-=28
 y-=10;line(c,'BY THE END, YOU WILL BE ABLE TO',36,y,10,accent,True);y-=9
 for row in range(3):
  for col in range(2):
   para(c,'• '+OUTCOMES[track][row+3*col],36+col*268,y,253,10)
  y-=16
 y-=4
 result='Analyse sample sales data to identify the best-selling products and explain the evidence. Leave with a Python notebook, charts, a cleaning log and a report.' if track=='Data Analysis' else 'Build and test a rent-estimation model using sample housing data. Leave with a Python notebook, baseline comparison and test results.'
 y=para(c,'<b>Your individual final project</b><br/>'+result,36,y,W-72,10.5)
 y=para(c,'Projects use synthetic sample data. Demonstrate your work to parents/guardians and earn a completion certificate by meeting the course requirements.',36,y-5,W-72,9.5)
 y=para(c,'<b>AI literacy:</b> Learn when AI coding assistants can help—and how to check their answers.',36,y-4,W-72,9)
 y=para(c,'<b>'+BIO+'</b><br/>'+BIO_BODY,36,y-9,W-72,10)
 y=para(c,'<link href="https://www.linkedin.com/in/jothammartinwambi/">'+LINKEDIN+'</link>',36,y-3,W-72,9,accent)
 y=fees(c,y-10)
 y=cta(c,y-10)
 if y<22:raise RuntimeError(('flyer overflow',track,y))
 c.showPage();c.save();print(fname.name,'bottom',round(y,1))

def onepager_pdf():
 c=canvas.Canvas(str(OUT/'Parent_Student_One_Pager.pdf'),pagesize=A4);c.setTitle('Holiday Data and AI Camp overview')
 line(c,'DECEMBER 2026 • ONLY 15 PLACES',36,H-29,9,TEAL,True)
 line(c,'Holiday Data and AI Camp',36,H-62,27,INK,True)
 y=para(c,'Build practical data and AI foundations before university.',36,H-78,W-72,13,bold=True)
 y=para(c,'For S.4–S.6 students and vacists. No coding experience needed. Own laptop required. Maximum five learners per group; choose one track.',36,y-6,W-72,10.5)
 y-=10
 for x,name,accent,body in [(36,'MACHINE LEARNING',PURPLE,'Build and test prediction models.<br/><b>Tools:</b> Python, Jupyter, NumPy, pandas, scikit-learn.<br/><b>Your project:</b> model, baseline comparison and test results.'),(306,'PYTHON FOR DATA ANALYSIS',TEAL,'Turn raw data into clear findings.<br/><b>Tools:</b> Python, Jupyter, NumPy, pandas, matplotlib, seaborn.<br/><b>Your project:</b> analysis notebook, charts, cleaning log and report.')]:
  c.setFillColor(PALE);c.roundRect(x,y-109,253,109,8,fill=1,stroke=0);line(c,name,x+13,y-22,11,accent,True);para(c,body,x+13,y-33,225,10.5)
 y-=122;line(c,'WHAT STUDENTS LEAVE WITH',36,y,11,TEAL,True);y-=12
 y=para(c,'<b>Practical skills:</b> write Python, work with data and check results.<br/><b>An individual project:</b> something each learner builds, tests and explains.<br/><b>Presentation experience:</b> a final demonstration to parents and guardians.<br/><b>A completion certificate:</b> awarded after meeting assessment and project requirements.<br/><b>AI literacy:</b> understand how modern AI coding assistants can help—and why you must be able to check, test and explain their answers.',36,y,W-72,10.5,leading=15)
 y=para(c,RELEVANCE,36,y-6,W-72,10.5)
 y-=15;line(c,'HOW THE EIGHT WEEKS WORK',36,y,10,TEAL,True);y-=12
 widths=[48,237,238];xs=[36,84,321];rowh=29
 c.setFillColor(NAVY);c.rect(36,y-25,W-72,25,fill=1,stroke=0)
 for x,t in zip(xs,['Weeks','Machine Learning','Python for Data Analysis']):line(c,t,x+6,y-16,9,white,True)
 y-=25
 for idx,row in enumerate(weeks):
  c.setFillColor(PALE if idx%2==0 else white);c.rect(36,y-rowh,W-72,rowh,fill=1,stroke=0)
  for x,w,t in zip(xs,widths,row):para(c,escape(t),x+6,y-4,w-14,9.5,leading=11.5)
  y-=rowh
 y=para(c,SCHEDULE+' ML: 24 classes + 4 labs. DA: 24 classes.',36,y-6,W-72,10)
 y=para(c,'<b>'+BIO+'</b><br/>'+BIO_BODY,36,y-7,W-72,10.5)
 y=para(c,'<link href="https://www.linkedin.com/in/jothammartinwambi/">'+LINKEDIN+'</link>',36,y-3,W-72,9,TEAL)
 y=para(c,'<b>Programme fee per learner:</b> UGX 1,000,000 for any group of one to five learners. <b>Venue:</b> The National ICT Innovation Hub, Nakawa.',36,y-7,W-72,10.5)
 y=para(c,'Projects use synthetic sample data. This is a beginner programme, not a professional qualification.',36,y-6,W-72,9.5)
 y=para(c,'<b>Foundations first:</b> AI coding assistants are introduced only after learners build the core skills themselves.',36,y-4,W-72,9.5)
 y=cta(c,y-6)
 if y<22:raise RuntimeError(('one pager overflow',y))
 c.showPage();c.save();print('one pager bottom',round(y,1))

def onepager_docx():
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69);sec.top_margin=sec.bottom_margin=Inches(.48);sec.left_margin=sec.right_margin=Inches(.55)
 normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(6)
 for name in ['Title','Subtitle','Heading 1','Heading 2']:
  d.styles[name].font.name='Arial';d.styles[name].font.color.rgb=RGBColor(0,0,0)
 d.styles['Title'].font.size=Pt(23);d.styles['Title'].paragraph_format.space_after=Pt(7)
 d.add_paragraph('Holiday Data and AI Camp',style='Title')
 p=d.add_paragraph('December 2026 intake for S.4–S.6 students and vacists.');p.runs[0].bold=True
 d.add_paragraph('Eight weeks. '+SCHEDULE+' Choose one track. Each student must have their own laptop and charger.')
 d.add_paragraph('Machine Learning: build and test prediction models using Python, Jupyter, NumPy, pandas and scikit-learn. Python for Data Analysis: turn data into findings using Python, Jupyter, NumPy, pandas, matplotlib and seaborn. No coding experience needed.')
 p=d.add_paragraph();p.add_run('What students leave with  ').bold=True;p.add_run('Practical Python and data skills; an individual project they build, test and explain; presentation experience at the parent/guardian showcase; and a completion certificate after meeting assessment and project requirements.');p.add_run('\nAI literacy: ').bold=True;p.add_run('understand how modern AI coding assistants can help—and why you must be able to check, test and explain their answers.')
 d.add_paragraph(RELEVANCE)
 t=d.add_table(rows=1,cols=3);t.autofit=False
 for c,w in zip(t.columns,[.6,3.26,3.26]):c.width=Inches(w)
 for c,txt in zip(t.rows[0].cells,['Weeks','Machine Learning','Python for Data Analysis']):c.text=txt
 for row in weeks:
  for c,txt in zip(t.add_row().cells,row):c.text=txt
 for ri,row in enumerate(t.rows):
  for c in row.cells:
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'142B42' if ri==0 else ('EDF3F6' if ri%2 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2)
    for r in p.runs:r.font.size=Pt(9);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '183044')
 rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
 for head,txt in [('Personal attention','Only 15 places across both tracks; maximum five learners per group. Worked examples, guided practice and individual feedback.'),('Individual project','Estimate rent from sample housing data with a tested model, or analyse sample sales data and explain product performance. Projects use synthetic sample data. Each learner presents their own work.'),('Fee and venue','Entire programme: UGX 1,000,000 per learner, the same for any group of one to five. Venue: The National ICT Innovation Hub, Nakawa.'),('Completion certificate','Awarded for meeting the published course requirements, including assessed work and an individual demonstration. This beginner programme is not a professional qualification.'),('Main tutor',BIO.replace(' • Main tutor','')+'. '+BIO_BODY),('Reserve your place','WhatsApp/Call: +256 750 990 718. Email: jothammartin.wambi@gmail.com. '+LINKEDIN+'. Scan to WhatsApp; confirm dates, venue and booking terms before payment.')]:
  p=d.add_paragraph();p.paragraph_format.space_before=Pt(5);p.add_run(head+'  ').bold=True;p.add_run(txt)
  if head=='Completion certificate':
   p=d.add_paragraph();p.add_run('Foundations first: ').bold=True;p.add_run('AI coding assistants are introduced only after learners build the core skills themselves.')
   for run in p.runs:run.font.size=Pt(9)
 d.add_picture(str(OUT/'WhatsApp_QR.png'),width=Inches(.85))
 for root in [d.styles.element,d.element]:
  for b in list(root.iter(qn('w:pBdr'))):b.getparent().remove(b)
 d.save(OUT/'Parent_Student_One_Pager.docx')

if __name__=='__main__':
 make_qr()
 for name in ROADMAP:flyer(name)
 onepager_pdf();onepager_docx()
