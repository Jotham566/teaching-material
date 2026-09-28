import sys,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from pptx import Presentation
import pptxkit as kit
class LecturerSlideNumbers(unittest.TestCase):
 def setUp(self):kit.LECTURER.clear();kit.notes_context('QA deck')
 def test_slides_without_notes_do_not_shift_the_guide(self):
  p=Presentation();p.slides.add_slide(p.slide_layouts[6]);second=p.slides.add_slide(p.slide_layouts[6]);kit.notes(second,'Student explanation. TO TEACH: Lecturer directions.')
  self.assertEqual(kit.LECTURER[-1]['slide'],2)
  self.assertNotIn('Lecturer directions',second.notes_slide.notes_text_frame.text)
 def test_repeated_and_out_of_order_notes_use_actual_position(self):
  p=Presentation();first=p.slides.add_slide(p.slide_layouts[6]);second=p.slides.add_slide(p.slide_layouts[6])
  kit.notes(second,'TO TEACH: Second slide.');kit.notes(first,'TO TEACH: First slide.');kit.notes(second,'TO TEACH: More on the second slide.')
  self.assertEqual([r['slide'] for r in kit.LECTURER],[2,1,2])
if __name__=='__main__':unittest.main()
