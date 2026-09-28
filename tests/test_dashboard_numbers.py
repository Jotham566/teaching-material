import json,sqlite3,sys,tempfile,unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];G=ROOT/'IUEA_Teaching_Plan/grader';sys.path.insert(0,str(G))
import dashboard_numbers as dn
class DashboardNumbers(unittest.TestCase):
 def text(self,expected):return '\n'.join(f'S15CHECK {k} revenue={v["revenue"]:.2f} orders={v["orders"]} aov={v["aov"]:.2f}' for k,v in expected.items())
 def test_independent_sql_reconciliation(self):
  values=dn.load_reference()
  with sqlite3.connect(ROOT/'DataAnalysis_Track/data/riverside.db') as db:
   for key,(year,region) in dn.CONTEXTS.items():
    args=[];where="o.status='Delivered'"
    if year:where+=' AND substr(o.order_date,1,4)=?';args.append(year)
    if region:where+=' AND r.region=?';args.append(region)
    amount,count=db.execute('SELECT SUM(i.quantity*i.unit_price_ugx*(1-i.discount)), COUNT(DISTINCT o.order_id) FROM orders o JOIN order_items i ON i.order_id=o.order_id JOIN reps r ON r.rep_id=o.rep_id WHERE '+where,args).fetchone()
    self.assertAlmostEqual(values[key]['revenue'],amount,places=2)
    self.assertEqual(values[key]['orders'],count)
    self.assertAlmostEqual(values[key]['aov'],amount/count,places=5)
 def test_wrong_denominator_duplicate_and_missing_context(self):
  expected=dn.load_reference();text=self.text(expected)
  self.assertEqual(dn.assess_text(text,expected)['score'],20)
  with self.assertRaises(ValueError):dn.assess_text(text+'\n'+text.splitlines()[0],expected)
  with self.assertRaises(ValueError):dn.assess_text('\n'.join(text.splitlines()[1:]),expected)
  wrong=json.loads(json.dumps(expected));wrong['ALL']['aov']/=2
  result=dn.assess_text(self.text(wrong),expected)
  self.assertEqual(result['score'],16);self.assertEqual(result['status'],'COMPONENT_ONLY')
 def test_unreadable_and_executable_appendix_rejected(self):
  expected=dn.load_reference()
  for text in ['',self.text(expected).replace('orders=','orders=__import__("os")')]:
   with self.assertRaises(ValueError):dn.assess_text(text,expected)
 def test_docx_extraction_in_isolation(self):
  from docx import Document
  from artifact_grade import extract
  expected=dn.load_reference()
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'evidence.docx';doc=Document();doc.add_heading('Synthetic dashboard QA',0)
   for line in self.text(expected).splitlines():doc.add_paragraph(line)
   doc.save(p);self.assertEqual(dn.assess_text(extract(p),expected)['score'],20)
if __name__=='__main__':unittest.main()
