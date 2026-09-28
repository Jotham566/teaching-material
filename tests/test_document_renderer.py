import sys,unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'IUEA_Teaching_Plan/grader'))
from render_submission import bounded_output
class RendererBoundaryTests(unittest.TestCase):
 def test_valid_stream(self):
  self.assertEqual(bounded_output([sys.executable,'-c','print("rendered")']),'rendered\n')
 def test_output_limit(self):
  with self.assertRaises(ValueError):bounded_output([sys.executable,'-c','print("x"*100000)'],limit=1024)
 def test_timeout(self):
  with self.assertRaises(TimeoutError):bounded_output([sys.executable,'-c','import time;time.sleep(5)'],timeout=.1)
 def test_failure(self):
  with self.assertRaises(RuntimeError):bounded_output([sys.executable,'-c','raise SystemExit(1)'])

class AccessibleDocumentTests(unittest.TestCase):
 def test_slide_order_and_narration_notes(self):
  import tempfile
  from pptx import Presentation
  from render_submission import extract_document
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'slides.pptx';prs=Presentation()
   for title,note in [('Original first','First narration'),('Now first','Second narration')]:
    slide=prs.slides.add_slide(prs.slide_layouts[1]);slide.shapes.title.text=title;slide.notes_slide.notes_text_frame.text=note
   ids=prs.slides._sldIdLst;ids.insert(0,ids[-1]);prs.save(p)
   text=extract_document(p)
   self.assertLess(text.index('Now first'),text.index('Original first'))
   self.assertLess(text.index('Second narration'),text.index('First narration'))
 def test_expanded_archive_limit(self):
  import tempfile,zipfile
  from render_submission import extract_document
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'large.docx'
   with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('word/document.xml',b'x'*(81*1024*1024))
   with self.assertRaises(RuntimeError):extract_document(p)
