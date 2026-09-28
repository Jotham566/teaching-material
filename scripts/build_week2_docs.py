"""Rebuild the concise Colab-first student guide and lecturer runsheet."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
ROOT=Path(__file__).resolve().parents[1]


def document(title):
    d=Document()
    for section in d.sections:
        section.page_width=Cm(21);section.page_height=Cm(29.7)
        section.top_margin=section.bottom_margin=Cm(1.6)
        section.left_margin=section.right_margin=Cm(2)
    for style in ('Normal','Title','Heading 1','Heading 2'):
        d.styles[style].font.name='Calibri'
        d.styles[style].font.color.rgb=RGBColor(0,0,0)
    d.styles['Title'].font.size=Pt(24)
    d.styles['Heading 1'].font.size=Pt(13)
    d.styles['Normal'].paragraph_format.line_spacing=1.0
    d.styles['Normal'].font.size=Pt(11)
    d.styles['Normal'].paragraph_format.space_after=Pt(4)
    d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    d.styles['Heading 1'].paragraph_format.space_after=Pt(4)
    for node in list(d.styles.element.iter()):
        if node.tag.endswith('}pBdr'):
            node.getparent().remove(node)
    d.add_paragraph(title,'Title')
    return d


def student():
    # The student how-to is the single source for the printable guide and ZIPs.
    import re
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    d=document('Student setup and assignment guide')
    source=(ROOT/'IUEA_Teaching_Plan/setup/STUDENT_HOW_TO.md').read_text()
    for i,line in enumerate(source.splitlines()):
        if i==0:continue
        if not line.strip():continue
        if line=='---PAGE---':d.add_page_break();continue
        if line.startswith('# '):d.add_paragraph(line[2:],'Title');continue
        if line.startswith('## '):d.add_paragraph(line[3:],'Heading 1');continue
        p=d.add_paragraph(style='List Bullet' if line.startswith('- ') else 'Normal')
        if line.startswith('- '):line=line[2:]
        for j,part in enumerate(re.split(r'(\*\*.*?\*\*)',line)):
            if part.startswith('**'):
                p.add_run(part[2:-2]).bold=True
                continue
            for fragment in re.split(r'(https://[^\s]+)',part):
                if not fragment.startswith('https://'):p.add_run(fragment);continue
                url=fragment.rstrip('.,')
                h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
                r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');pr.append(color);r.append(pr)
                t=OxmlElement('w:t');t.text=url;r.append(t);h.append(r);p._p.append(h)
                if fragment!=url:p.add_run(fragment[len(url):])
    d.save(ROOT/'IUEA_Teaching_Plan/IUEA_ML_Student_Setup_Guide.docx')
    return


def lecturer():
    d=document('Week 2 lecturer runsheet')
    d.add_paragraph('15 and 18 September 2026 | Three hours per cohort | Kampala time')
    d.add_paragraph('Teach notebooks and version control using Colab. The operational priority is that each student runs and saves their own work. Use the small supplied kit and pre-provisioned lab fallback.')
    d.add_paragraph('Before students arrive','Heading 1')
    for text in ['Test Google and GitHub sign-in from a student account. Verify the assignment form and correct Week 2 quiz.',
                 'Share the cohort student ZIP before class; students download and extract it, without doing the exercises yet. Keep a USB copy and rehearse the lab fallback.',
                 'Check projector readability, room power and the CSV upload step. Brief floor helpers to resolve access issues while teaching continues.',
                 'Preserve Week 1 submissions. A2 is W02 in the gradebook; the second quiz is Q2.']:
        d.add_paragraph(text,'List Bullet')
    d.add_paragraph('Teaching timetable','Heading 1')
    t=d.add_table(rows=1,cols=2);t.style='Light Shading Accent 1'
    t.rows[0].cells[0].text='Elapsed time';t.rows[0].cells[1].text='Teaching activity'
    for time,activity in [('00–30','Recall and four-tool map'),('30–100','Start proof_of_life; model repair 01, let students try 02–03, guide 04; return to proof_of_life and fresh run'),('100–110','Break'),('110–160','Git explanation and repository practical'),('160–170','Repair recap and individual help; verify main notebook and own README'),('170–175','Five-question Week 2 quiz'),('175–180','A2 submission instructions and exit check')]:
        cells=t.add_row().cells;cells[0].text=time;cells[1].text=activity
    d.add_paragraph('Opening forms','Heading 1')
    d.add_paragraph('Use IUEA_ML_Quick_Class_Checklist.docx for clickable owner controls. Assignment uploads open and close on their published schedule; run syncSubmissions in the assignment project if a due opening has not appeared. In the separate quiz project, select open_W2_CML3112, open_W2_CML3113 or open_W2_CML3111 and Run once when introducing the quiz. Re-running resets its 24-hour window. CML3112 Week 1 catch-up quiz closes 16 September at 06:31:20 Kampala; its upload stays open until 22 September at 11:00; do not reopen them.')
    d.add_paragraph('Quiz and deadline','Heading 1')
    d.add_paragraph('Allow sign-in before the five-minute question period. This is a low-stakes quiz with the existing 24-hour response window. The classroom countdown does not enforce a per-student timer. Check automatic closure after the window; failed closures must remain queued for retry.')
    d.add_paragraph('Fallback and exit check','Heading 1')
    d.add_paragraph('If Colab fails, move students to prepared lab computers with the same ZIP. Record an outage rather than silently penalizing delayed uploads. Before leaving, each student shows a fresh-run notebook and keeps a downloaded copy. A2 closes at the start of that cohort’s Week 3 lecture.')
    d.save(ROOT/'IUEA_Teaching_Plan/IUEA_ML_Week2_Lecturer_Runsheet.docx')

if __name__=='__main__':student();lecturer()
