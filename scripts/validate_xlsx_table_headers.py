"""Read-only Excel table/header consistency check after workbook export."""
import sys
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


def validate(filename):
    workbook = load_workbook(filename)
    try:
        count = 0
        for sheet in workbook:
            for table in sheet.tables.values():
                if table.headerRowCount == 0:
                    continue
                first_col, first_row, last_col, _ = range_boundaries(table.ref)
                visible = [str(sheet.cell(first_row, col).value or '') for col in range(first_col, last_col + 1)]
                stored = [column.name for column in table.tableColumns]
                if visible != stored:
                    raise ValueError(f'{sheet.title}: table {table.name} headings differ from visible cells')
                count += 1
        return count
    finally:
        workbook.close()


if __name__ == '__main__':
    print(f'Validated {validate(sys.argv[1])} Excel table headers.')
