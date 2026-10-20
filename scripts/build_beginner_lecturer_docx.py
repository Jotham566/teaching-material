"""Five-page quick guide to the revised beginner teaching route."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from beginner_pathway import WEEKS, VIDEOS
ROOT=Path(__file__).resolve().parents[1]
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.7);s.left_margin=s.right_margin=Cm(2)
for name in ('Normal','Title','Heading 1','Heading 2'):
 d.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 d.styles[name].paragraph_format.left_indent=Cm(0);d.styles[name].paragraph_format.right_indent=Cm(0);d.styles[name].paragraph_format.first_line_indent=Cm(0)
 d.styles[name].paragraph_format.line_spacing=1.1
 d.styles[name].font.name='Liberation Sans';d.styles[name].font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(7)
d.styles['Title'].font.size=Pt(24);d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 2'].font.size=Pt(12)
for element in list(d.styles.element.iter()):
 if element.tag==qn('w:pBdr'):element.getparent().remove(element)
def p(t):d.add_paragraph(t)
def h(t):d.add_paragraph(t,'Heading 1')
def page(title):
 para=d.add_paragraph(title,'Title');para.paragraph_format.page_break_before=True
 para.paragraph_format.space_after=Pt(18)
 d.add_paragraph('')
def week(w):
 x=WEEKS[w];h(f'Week {w}  {x["title"]}');p('Teach: '+x['core']);p('Check: '+x['gate']);p('Defer: '+x['stretch'])
d.add_paragraph('Teaching complete beginners','Title')
p('Lecturer quick guide | Revised 15 October 2026')
p('From Week 6 onward, continue using pandas for table work. The current slides and matching notebooks explain the new operations before examples. Check Teaching_Progress.md for cohort pacing.')
h('Open one place before class')
p('In IUEA_Teaching_Plan/decks, open the PowerPoint for your cohort and week. Use Slide Show for teaching and Presenter View for the explanation prompts in the speaker notes. The filenames are unchanged.')
p('Teach from the updated PowerPoint. Explain its pictures, annotated code and worked answers directly. Students repeat the examples individually after class. No live partner activity is required. Read the live assignment form for the current deadline; the separately authorised A3 reopening supersedes the older date printed in its slides.')
h('What to share')
p('Share the current cohort student ZIP. From Week 3 where relevant, it includes guided_practice.ipynb for the examples shown in the PowerPoint. It follows the revised slide order, with matching slide numbers. The issued Wxx_student.ipynb is the assignment notebook. Its opening checklist and task prompts contain the required work; separate instruction files are reference copies. For Week 2, proof_of_life.ipynb remains the main notebook and the four repairs stay practice.')
h('Repeat this teaching routine')
for t in ['1. Explain the physical question without coding words. State every quantity and unit.', '2. Define numbers, values and variables before assignment syntax; define every new word before use. Say what each symbol in the code does.', '3. Work one tiny example. Ask everyone to predict a changed value before running it.', '4. Demonstrate an analogous example and explain how each changed input affects the output.', '5. Set individual revision using the worked answers. Repeat a smaller example if learners cannot explain it.', '6. Explain the exact assignment output names, required files, checks and deadline. Students save their own work.']:
 p(t)
h('Support without extra marking')
p('Use the demonstrated outputs and visible practice checks. Hint in stages: restate the question, trace a smaller example, supply the first step, then model an analogous solution. Keep the existing automatic assessment workflow; do not add manually marked worksheets.')
page('Week 2 reference and Week 3 teaching')
week(2)
p('Week 2 reference: A1, Quiz 1, A2 and Quiz 2 are separate tasks. Check the current form deadlines before referring students back to them. Each quiz lasts 24 hours from activation.')
p('Use the private Week2 Repair Walkthrough for the exact corrections. Add the GitHub concept clips before saving a repository: 00:00–02:10 and 03:35–06:08 of the CodeWithChris video you shared. Desktop installation is optional; the course browser workflow remains the default.')
week(3)
p('After values, variables, types and operators, use each cohort’s simple formula: floor area for CML3112, pump flow for CML3113, DC electrical power for CML3111. Explain all quantities rather than assuming discipline knowledge. Then introduce reusable functions using double_value, return and a guarded division.')
p('A3 has five tasks; A4 has four named outputs. Each requires one completed notebook, with no separate written report. Point to the task, marks and answer cell. Students who already started may keep their existing work and use refreshed instructions as a reference.')
h('Three-hour teaching rhythm')
p('00–15 recall/access; 15–30 plain explanation; 30–50 worked example and optional short clip; 50–80 guided attempt; 80–90 break; 90–125 supported practice; 125–140 assignment bridge; 140–165 individual attempt/recovery; 165–170 quiz; 170–180 save and submission instructions.')
p('Quizzes close 24 hours after activation, whether opened by a scheduled control or manually. Released assignment uploads follow their configured opening and closing times. Use the Quick Class Checklist for controls and links. Assignment forms now accept any Student Registration Number format; the field is still required.')
page('Tables, cleaning and Test 1')
for w in (4,5,6):week(w)
h('Week 7 preparation')
p('Rehearse the offline environment, public interfaces and saving both notebook and typed Part B before the test. Keep the official assessment timetable. Confirm lab readiness and the separate grading release checks.')
p('Before Week 6, rehearse reading a CSV, selecting columns, filtering and grouping with pandas. Use the issued loading scaffold and the slide-by-slide notes. Let the cohort’s actual progress determine how much recap is needed.')
page('The modelling sequence')
for w in (8,9,10,11,12):
 x=WEEKS[w];h(f'Week {w}  {x["title"]}');p(x['core']);p('Before independent work: '+x['gate'].replace('return to the paper partitions','revisit the illustrated train/validation/test split'))
h('Correct these ideas explicitly')
p('MAE is an average error, not a ± guarantee. Teach the required regression tree in Week 8 before it is assessed. Keep repeated equipment groups together and fit preprocessing only on training data. Select models on validation data, reserving the final test. Lowering a classification threshold does not guarantee precision falls at every step.')
page('Projects and revision')
for w in (13,14,15):
 x=WEEKS[w];h(f'Week {w}  {x["title"]}');p(x['core'].replace('one useful peer review','an independent run check'));p(x['bridge'])
h('Short video use')
p('The PowerPoints contain clickable choices for Git, Python arithmetic, boxplots, regression, confusion matrices and cross-validation. Preview before class, state one listening question, play about four minutes at most, pause and do the related exercise. Optional entries still marked preview-required have not had a complete audiovisual review.')
p('Your GitHub link originally starts at 02:47, during installation. Your Colab link starts at 17:14 of a 22:23 tutorial. Share the un-timestamped links for optional revision. GPUs, TPUs and extra installations in tutorials are not Week 2 requirements.')
h('Before each class')
p('Open the actual cohort PowerPoint and run its practice notebook. Read the assignment brief and quiz privately. Rehearse any syntax used by the quiz, including the examples in the slides and speaker notes. Check the student ZIP and submission link. Preview the chosen clip, or use the local example if video access fails.')
p('Beginner_Lecturer_Guide.md gives weekly guidance. Curriculum_Benchmark_and_Teaching_Map.md lists the research and activities. Major projects, tests and demos still require their separate grading-readiness checks before release.')
# Explicit run fonts avoid theme-font substitution in document viewers.
for paragraph in d.paragraphs:
 for run in paragraph.runs:
  run.font.name='Liberation Sans'
  run.font.size=Pt(24 if paragraph.style.name=='Title' else 15 if paragraph.style.name=='Heading 1' else 10.5)
  run.font.bold=paragraph.style.name=='Heading 1'
  fonts=run._r.get_or_add_rPr().rFonts
  for attr in list(fonts.attrib):
   if attr.endswith(('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme')):del fonts.attrib[attr]
d.save(ROOT/'IUEA_Teaching_Plan/IUEA_ML_Beginner_Teaching_Quick_Guide.docx')
