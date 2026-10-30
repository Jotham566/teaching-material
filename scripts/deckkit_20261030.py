"""Slide-building helpers shared by the Week 8b to 12 patch. Wraps the helpers in
patch_week8_20261030 and adds idea/statement slides, deck re-ordering and numbering."""
import re, zipfile
from pathlib import Path
from pptx.util import Pt
from pptx.opc.packuri import PackURI
import patch_week8_20261030 as p8
from patch_week8_20261030 import (shape_rect, textbox, header, footer, set_notes, new_slide, move_slide, find_slide,
                                  delete_slide, figure_slide, table_slide, code_slide, rgb, slide_title_texts,
                                  INK, ORANGE, TEAL, MUTED, PALE1, PALE2)


def idea_slide(prs, kicker, title, left, right, notes, label, left_fill=PALE1, right_fill=PALE2, size=22):
    """Two panels like the existing 'one idea at a time' slides. left/right: list of (text, bold) or str."""
    s = new_slide(prs); header(s, kicker, title)
    shape_rect(s, 46, 138, 418, 328, left_fill); shape_rect(s, 488, 138, 425, 328, right_fill)
    textbox(s, left, 67, 156, 374, 300, size, INK, False, space_after=8)
    textbox(s, right, 508, 156, 382, 300, size, INK, False, space_after=8)
    footer(s, label); set_notes(s, notes); return s


def statement_slide(prs, kicker, title, statement, body, notes, label, size=26, body_size=22):
    """A teal statement panel on the left and explanation on the right (like the existing slide 12 style)."""
    s = new_slide(prs); header(s, kicker, title)
    shape_rect(s, 46, 138, 352, 331, TEAL)
    textbox(s, statement, 68, 160, 308, 290, size, "FFFFFF", True)
    textbox(s, body, 428, 150, 470, 310, body_size, INK, False, space_after=9)
    footer(s, label); set_notes(s, notes); return s


def three_col_figure(prs, title, img, caption, notes, label):
    return figure_slide(prs, title, img, caption, notes, label)


def replace_two_panels(slide, left_head, left_body, right_head, right_body):
    bodies = [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.top > Pt(120) and sh.top < Pt(300)
              and sh.width > Pt(300) and not sh.text_frame.text.startswith("CML")]
    bodies = sorted(bodies, key=lambda sh: sh.left)[:2]
    assert len(bodies) == 2, f"expected two body panels, found {len(bodies)}"
    for sh, head, body in zip(bodies, (left_head, right_head), (left_body, right_body)):
        tf = sh.text_frame
        for r in list(tf.paragraphs[0].runs)[1:]: r.text = ""
        for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)
        p = tf.paragraphs[0]; ref = p.runs[0]
        ref.text = head; ref.font.bold = True; ref.font.size = Pt(19); ref.font.color.rgb = rgb(TEAL)
        p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = body
        r2.font.size = Pt(22); r2.font.name = "Calibri"; r2.font.bold = False; r2.font.color.rgb = rgb(INK); p2.space_before = Pt(8)


def replace_text(slide, old, new):
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for p in sh.text_frame.paragraphs:
            full = "".join(r.text for r in p.runs)
            if old in full and p.runs:
                p.runs[0].text = full.replace(old, new)
                for r in p.runs[1:]: r.text = ""
                n += 1
    return n


def order_deck(prs, wanted):
    """wanted: list of slide objects in final order. Everything else is removed."""
    lst = prs.slides._sldIdLst
    byslide = {}
    for e in list(lst): byslide[id(prs.part.related_part(e.rId))] = e
    keep = {id(s.part) for s in wanted}
    for e in list(lst):
        if id(prs.part.related_part(e.rId)) not in keep:
            prs.part.drop_rel(e.rId); lst.remove(e)
    els = [byslide[id(s.part)] for s in wanted]
    for e in list(lst): lst.remove(e)
    for e in els: lst.append(e)


def finalize(prs, code, week, dst):
    pat = re.compile(r"^" + re.escape(code) + r"\s+Week " + str(week) + r"\s+\d+$")
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame and pat.match(sh.text_frame.text.strip()):
                r = sh.text_frame.paragraphs[0].runs
                r[0].text = f"{code}   Week {week}   {i}"
                for x in r[1:]: x.text = ""
        s.part.partname = PackURI(f"/ppt/slides/slide{i}.xml")
        if s.has_notes_slide: s.notes_slide.part.partname = PackURI(f"/ppt/notesSlides/notesSlide{i}.xml")
    Path(dst).parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    names = zipfile.ZipFile(dst).namelist()
    assert len(names) == len(set(names)), "duplicate parts in saved package"


def spoken(*paras):
    """Join paragraphs of a spoken script."""
    return "\n\n".join(p.strip() for p in paras)


_code_slide_base = code_slide


def code_slide(prs, kicker, title, code, output, cues, caption, notes, label, code_size=17, **kw):
    """Same as the base builder, but refuses lines that would wrap inside the 540 pt code panel."""
    longest = max(len(l) for l in code.split("\n"))
    assert longest * code_size * 0.6 <= 540, f"code line too long for {title}: {longest} chars at {code_size} pt"
    out_size = kw.get("out_size", 16)
    assert max(len(l) for l in output.split("\n")) * out_size * 0.6 <= 540, f"output line too long for {title}"
    return _code_slide_base(prs, kicker, title, code, output, cues, caption, notes, label, code_size=code_size, **kw)


def edit_notes(slide, old, new):
    tf = slide.notes_slide.notes_text_frame
    assert old in tf.text, f"notes text not found: {old}"
    tf.text = tf.text.replace(old, new)


def set_body(slide, starts_with, paragraphs):
    """Replace the text of the body shape that starts with a given phrase, keeping its formatting."""
    import copy as _copy
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith(starts_with):
            tf = sh.text_frame; first = tf.paragraphs[0]._p
            for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
            for r in first.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")[1:]: first.remove(r)
            tf.paragraphs[0].runs[0].text = paragraphs[0]
            for txt in paragraphs[1:]:
                np_ = _copy.deepcopy(first); first.addnext(np_) if False else tf._txBody.append(np_)
                from pptx.text.text import _Paragraph
                _Paragraph(np_, tf).runs[0].text = txt
            return True
    raise AssertionError(f"no body shape starting with {starts_with!r}")
