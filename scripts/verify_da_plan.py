"""Catch schedule and section drift in the generated independent-track plan."""
from pathlib import Path
import re
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
def verify():
 doc=Document(ROOT/'DataAnalysis_Track/DataAnalysis_Track_Plan.docx')
 paragraphs=[p.text for p in doc.paragraphs]
 tables=[c.text for t in doc.tables for r in t.rows for c in r.cells]
 text='\n'.join(paragraphs+tables)
 for pattern in [r'FOUR WEEKS',r'the four weeks',r'twelve sessions',r'thirty-six hours',r'throughout week four',r'cannot give them the full Power Query']:
  if re.search(pattern,text,re.I):raise AssertionError('Stale plan wording: '+pattern)
 headings=[int(m.group(1)) for p in paragraphs if (m:=re.match(r'^(\d+) · ',p))]
 assert headings==list(range(1,11)),headings
 for week in range(1,6):assert any(f'Week {week} —' in t for t in tables),week
 sessions=[int(m.group(1)) for p in paragraphs if (m:=re.match(r'^Session (\d+) · ',p))]
 assert sessions==list(range(1,16)),sessions
 assert 'forty-five hours' in text and 'S13' in text and 'S15' in text
 print('DA plan: five weeks, fifteen sessions, 45 hours, ten unique sections; no stale duration/platform wording')
if __name__=='__main__':verify()
