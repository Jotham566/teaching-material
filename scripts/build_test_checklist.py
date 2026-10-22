from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for n in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Title'].font.size=Pt(22);d.styles['Heading 1'].font.size=Pt(14)
def p(t):d.add_paragraph(t)
def link(text,url):
 para=d.add_paragraph();h=OxmlElement('w:hyperlink');h.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');pr.append(color);r.append(pr);txt=OxmlElement('w:t');txt.text=text;r.append(txt);h.append(r);para._p.append(h)
d.add_paragraph('IUEA test grading checklist','Title')
p('Test 1 and Test 2 • CML3111, CML3112 and CML3113\nStudents submit in Colab; your Mac grades the batch afterwards.')
p('Use wirmbeesubs@gmail.com for the lecturer forms and Drive. Google AI Plus is not needed. Keep this account: it still owns the teaching files.')
p('For weekly A1–A5 and Quiz 1–5 grading, use the Continue grading guide linked below. Weekly work uses Export Grading ZIP in the assignment workbook. Tests use Export Test ZIP in the test workbook.')
d.add_heading('Prepare students',1)
p('A few days before: send the Week 7 or Week 14 revision PDF and its matching practice CSV. Students practise in a blank Colab notebook; no practice upload is required. The portal releases test files only during the enabled test window.')
d.add_heading('Before the test',1)
p('1. Open Test windows in the collection workbook. Set the correct course and test opening and closing times, including +03:00 for Kampala. Set enabled to YES. Use the format YYYY-MM-DDTHH:MM:SS+03:00.')
p('2. In Apps Script, open Colab_Test_Collection.gs. Choose syncColabExams in the function dropdown and click Run at the start. Automatic checks also run every five minutes; the exact submission deadline is enforced separately.')
p('3. Share the course/test entry link or QR from the Week 7 or Week 14 deck. Each learner downloads their notebook and CSV. Rehearse both downloads, the notebook button and form confirmation. Forms must collect Verified email; Student Registration Number stays required and freeform.')
d.add_heading('Tell the students',1)
p('Use your own Google account in both Colab and the form. Run the final notebook cell and approve the account prompt. Click Submit test, open the generated link, check your details, tick the email confirmation box and press Submit on the form. Wait for confirmation. Keep the snapshot in Drive until marks are released. Report any failure immediately.')
d.add_page_break()
d.add_heading('After the test',1)
p('4. Leave enabled as YES and keep the original closing time. The form closes by time; delayed collection can still recover on-time responses.')
p('5. In the test collection workbook, choose Export Test ZIP > Export latest submissions. Refresh once if the menu is missing; authorize it on first use. Use View progress / download and wait for the new export to say COMPLETE. While the request says PENDING, the old link may remain visible. Download the ZIP from Download / detail. It includes both tests across all three programmes; rehearsal entries are excluded.')
p('6. Start Docker Desktop. Double-click Grade_Tests.command in the test_operations folder. Drag the ZIP into the Terminal window and press Return. Keep the Mac on until grading finishes.')
p('7. Test_Marks.xlsx opens automatically. Check Current marks, All attempts and Collection issues. Match the submission count to the collection workbook. Resolve pending or conflicting rows before releasing marks; a blank is not zero.')
p('8. Keep the ZIP and results private. Results are in IUEA_Teaching_Plan/test_operations/grading_output (inside this repository). Re-running the same batch does not duplicate marks. Test 1 contributes 15 semester points; Test 2 contributes 20.')
p('If export stops: read Download / detail, fix the problem and choose Export latest submissions again. Existing submissions remain intact. Repeated clicks while queued or running do not start another export.')
p('9. Save and push the private snapshot before switching Macs, using step 6 of Continue grading. New test marks need production test-ledger integration into the cumulative workbook; the weekly grading command does not import them automatically.')
link('Weekly grading and private backup checklist','https://github.com/Jotham566/teaching-material/blob/main/IUEA_Teaching_Plan/assessment_operations/CONTINUE_GRADING.md')
link('Full test checklist and cumulative marks guidance','https://github.com/Jotham566/teaching-material/blob/main/IUEA_Teaching_Plan/test_operations/TEST_CHECKLIST.md')
link('Open test collection workbook','https://docs.google.com/spreadsheets/d/1BveqLGvv-6C-00TMbUF0qti6VoEnAcRX9K1s9eja8uM/edit')
link('Open Apps Script controls','https://script.google.com/home/projects/1XLdUAzi6wnSjQ2iFTV3hl5BQk9XWw-PIi1_vodC_l107B6S13yW_ZObw/edit')
for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for border in list(d.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
path=ROOT/'IUEA_Teaching_Plan/test_operations/Test_Grading_Checklist.docx';d.save(path);print(path)
