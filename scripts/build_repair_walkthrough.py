"""Render the private repair walkthrough as a compact Word reference."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
ROOT=Path(__file__).resolve().parents[1]
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(1.6);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 d.styles[name].font.name='Calibri';d.styles[name].font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(5);d.styles['Normal'].paragraph_format.line_spacing=1.0
d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(14)
for n in list(d.styles.element.iter()):
 if n.tag.endswith('}pBdr'):n.getparent().remove(n)
def p(s):d.add_paragraph(s)
def h(s):d.add_paragraph(s,'Heading 1')
def code(s):
 a=d.add_paragraph();a.paragraph_format.space_after=Pt(7)
 for i,line in enumerate(s.split('\n')):
  r=a.add_run(('\n' if i else '')+line);r.font.name='Consolas';r.font.size=Pt(10)
d.add_paragraph('Week 2 repair walkthrough','Title')
p('Lecturer only • Worked answers for notebooks 01 to 04')
p('Use during the main notebook practical. Model 01; students try 02 and 03 with support; guide 04. Allow about five minutes each. Ask students to predict, run, read the error/output, repair and explain. These are practice, not four graded submissions.')
h('01 Missing module')
code('import course_package_not_installed')
p('What happens: ModuleNotFoundError. Python cannot find this deliberately unavailable module. Importing does not install a package.')
p('Ask: “What does the final error line say? Do we need this module?” Remove the import. Demonstrate a valid standard-library import instead:')
code('import csv\nprint(csv.__name__)')
p('Expected: csv. No installation is needed. Do not install the fictional module. In real work, verify the intended package and environment before installing anything.')
h('02 Python version mismatch')
code('import sys\nassert sys.version_info.major == 2, "Notebook expects Python 2, but the course uses Python 3"')
p('What happens: AssertionError in Python 3. The deliberately wrong expectation is the problem; Colab is not broken. Ask which version is running and which version the assertion demands.')
p('Correction: keep the Python 3 runtime and change the test to 3.')
code('import sys\nprint(sys.version_info.major)\nassert sys.version_info.major == 3, "Use Python 3"')
p('Expected: 3, with no assertion error. An assertion checks a condition; it cannot change or install the Python version.')
h('Optional script demonstration')
p('hello_engineering.py is not required for submission. If useful, upload it to Colab and run !python hello_engineering.py in a cell. It prints a greeting and the dataset name; it does not read or verify the CSV.')
d.add_page_break();d.add_paragraph('File paths and stored values','Title')
h('03 Wrong file path')
p('Original CML3112 path: missing_folder/site_inspection_log.csv. It raises FileNotFoundError because that location does not exist. A file on the laptop is not automatically inside Colab.')
p('Ask: “Where is Python looking? Where did we upload the file?” Upload the cohort CSV through the Files sidebar, then use its actual name:')
code('with open("site_inspection_log.csv") as f:\n    print(f.readline().strip())')
p('CML3112: site_inspection_log.csv → reading_id,crack_width_mm\nCML3113: drill_shift_log.csv → reading_id,depth_m\nCML3111: motor_test_bench.csv → reading_id,current_a')
p('The text after each arrow is the expected first line. Substitute your cohort filename in the code. Explain that with closes the file afterwards. Check spelling and upload location; re-upload after a runtime reset if needed.')
h('04 Variable state and execution order')
code('readings = 12\nreadings = readings + 5\nreadings = 12\nprint(readings)')
p('It prints 12 without an exception. Trace the value: 12 → 17 → 12. The third line overwrites the addition. Ask: “Which line discards our result?” Remove the second initialization:')
code('readings = 12\nreadings = readings + 5\nprint(readings)')
p('Expected: 17. Rerunning this WHOLE cell still prints 17 because the first line resets the value every time. Successful execution does not guarantee the right answer.')
h('Return to the main notebook')
p('In proof_of_life.ipynb, initialization and increment are in SEPARATE cells. Run initialization once, then the increment three times: 17, 22, 27. Restart clears stored state; Run All initializes once and adds five once, giving 17.')
p('Ask students to explain that difference in their own words. Finish the main notebook and own README. Upload only those two files together, plus enter the repository URL. Keep 01–04 as practice copies.')
d.save(ROOT/'IUEA_Teaching_Plan/IUEA_ML_Week2_Repair_Walkthrough.docx')
