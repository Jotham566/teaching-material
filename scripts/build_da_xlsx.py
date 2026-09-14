# -*- coding: utf-8 -*-
"""The Data Analysis track workbook: session grid, assessment, datasets,
and a participant mark sheet the instructor fills in as the course runs.

Conventions follow build_xlsx.py, which is the approved house style for this
project: Arial, navy header band, banded rows, frozen panes. Totals and
weighted marks are real formulas, so the sheet recalculates when a mark is
typed in.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HERE = os.path.dirname(os.path.abspath(__file__))
_ROOT = os.path.dirname(_HERE)
sys.path.insert(0, _HERE)
import da_content as C                                   # noqa: E402

TRACK = os.path.join(_ROOT, "DataAnalysis_Track")
OUT = os.path.join(TRACK, "DataAnalysis_Track_Workbook.xlsx")
with open(os.path.join(TRACK, "data", "answer_key.json"),
          encoding="utf-8") as fh:
    K = json.load(fh)

FONT = "Arial"
NAVY, MID, LIGHT, AMBER, GREEN = "1F3864", "2E5C8A", "D9E2F3", "FFF2CC", "E2EFDA"
H_FILL = PatternFill("solid", fgColor=NAVY)
H_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
T_FONT = Font(name=FONT, size=17, bold=True, color=NAVY)
S_FONT = Font(name=FONT, size=10.5, italic=True, color=MID)
BODY = Font(name=FONT, size=9)
BODY_B = Font(name=FONT, size=9, bold=True)
INPUT = Font(name=FONT, size=9, color="0000FF")          # cells to fill in
TOP = Alignment(vertical="top", wrap_text=True)
TOPC = Alignment(vertical="top", horizontal="center", wrap_text=True)
CTR = Alignment(vertical="center", horizontal="center", wrap_text=True)


def title_block(ws, title, subtitle, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(1, 1, title); c.font = T_FONT; c.alignment = TOP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)
    c = ws.cell(2, 1, subtitle); c.font = S_FONT; c.alignment = TOP
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 16


def header_row(ws, row, headers, widths, height=30):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row, i, h)
        c.font = H_FONT; c.fill = H_FILL; c.alignment = CTR
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = height


def write_row(ws, row, values, fill=None, bold_first=False, height=None):
    for i, v in enumerate(values, 1):
        c = ws.cell(row, i, v)
        c.font = BODY_B if (bold_first and i == 1) else BODY
        c.alignment = TOP if i > 1 else TOP
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height


wb = Workbook()

# ============================================================ 1. session grid
ws = wb.active
ws.title = "Session grid"
title_block(ws, C.TRACK, "{} · {} hours · every session, in order".format(
    C.SUBTITLE, C.TOTAL_HOURS), 9)
header_row(ws, 4, ["#", "Week", "Tool", "Title", "New words",
                   "roadmap.sh topic", "Teach block 1", "Teach block 2",
                   "Lab, last 50 minutes"],
           [5, 7, 13, 30, 10, 26, 44, 44, 44])
r = 5
for s in C.SESSIONS:
    write_row(ws, r, [s["n"], s["week"], s["tool"], s["title"],
                      len(s.get("vocab", [])), s["roadmap"],
                      s["b1"], s["b2"], s["lab"]],
              fill=LIGHT if s["week"] % 2 else None, height=86)
    r += 1
ws.freeze_panes = "D5"

# ============================================================ 2. the 3h block
ws = wb.create_sheet("Three-hour block")
title_block(ws, "The three-hour block",
            "The same shape every session. Minutes sum to 180, checked by "
            "scripts/verify_da_track.py", 4)
header_row(ws, 4, ["Minutes", "Running", "What", "Why"], [10, 10, 30, 70])
r = 5
for i, (m, what, why) in enumerate(C.BLOCK):
    write_row(ws, r, [m, None, what, why], height=32)
    # running total as a formula, so an edited minute shows up immediately
    ws.cell(r, 2).value = ("=A5" if i == 0 else "=B{}+A{}".format(r - 1, r))
    ws.cell(r, 2).font = BODY
    r += 1
ws.cell(r, 1).value = "=SUM(A5:A{})".format(r - 1)
ws.cell(r, 1).font = BODY_B
ws.cell(r, 3).value = "Total, which must be 180"
ws.cell(r, 3).font = BODY_B
for col in (1, 2, 3, 4):
    ws.cell(r, col).fill = PatternFill("solid", fgColor=GREEN)
ws.freeze_panes = "A5"

# ============================================================ 3. assessment
ws = wb.create_sheet("Assessment")
title_block(ws, "Assessment", "Weights must total 100", 4)
header_row(ws, 4, ["What", "Weight", "When", "How it is marked"],
           [34, 10, 14, 70])
r = 5
for a, w, when, how in C.ASSESSMENT:
    write_row(ws, r, [a, w / 100.0, when, how], height=44)
    ws.cell(r, 2).number_format = "0%"
    r += 1
ws.cell(r, 1).value = "Total"
ws.cell(r, 1).font = BODY_B
ws.cell(r, 2).value = "=SUM(B5:B{})".format(r - 1)
ws.cell(r, 2).font = BODY_B
ws.cell(r, 2).number_format = "0%"
for col in range(1, 5):
    ws.cell(r, col).fill = PatternFill("solid", fgColor=GREEN)
ws.freeze_panes = "A5"

# ============================================================ 4. mark sheet
ws = wb.create_sheet("Mark sheet")
title_block(ws, "Participant mark sheet",
            "Type marks out of 100 in the blue cells. The weighted total is "
            "a formula and recalculates as you go.", 8)
labels = [a for a, _, _, _ in C.ASSESSMENT]
weights = [w for _, w, _, _ in C.ASSESSMENT]
header_row(ws, 4, ["Participant", "Email"] + labels + ["Weighted total",
                                                       "Standing"],
           [26, 26] + [16] * len(labels) + [15, 14], height=44)
# the weights live in their own labelled row, and the formulas reference it
ws.cell(5, 1, "Weight").font = BODY_B
for i, w in enumerate(weights):
    c = ws.cell(5, 3 + i, w / 100.0)
    c.font = BODY_B
    c.number_format = "0%"
    c.fill = PatternFill("solid", fgColor=AMBER)
ws.cell(5, 3 + len(labels), "=SUM(C5:{}5)".format(
    get_column_letter(2 + len(labels))))
ws.cell(5, 3 + len(labels)).number_format = "0%"
ws.cell(5, 3 + len(labels)).font = BODY_B

FIRST_DATA = 6
N_ROWS = 24
last_mark_col = get_column_letter(2 + len(labels))
tot_col = get_column_letter(3 + len(labels))
for i in range(N_ROWS):
    r = FIRST_DATA + i
    if i == 0:
        # one example row, in the format expected, per the workbook guidance
        write_row(ws, r, ["Example: Grace Nabirye",
                          "grace@example.com", 80, 72, 68, 75],
                  fill=AMBER)
        for cc in range(3, 3 + len(labels)):
            ws.cell(r, cc).font = INPUT
    else:
        write_row(ws, r, ["", ""] + [None] * len(labels))
        for cc in range(3, 3 + len(labels)):
            ws.cell(r, cc).font = INPUT
    # SUMPRODUCT against the weight row, guarded so a blank row stays blank.
    # The denominator is a second SUMPRODUCT rather than SUM(IF(...)):
    # SUM(IF()) over a range needs array entry, which openpyxl cannot write,
    # and it came back #VALUE! from the recalc.
    ws.cell(r, 3 + len(labels)).value = (
        '=IF(COUNT(C{r}:{lc}{r})=0,"",'
        'SUMPRODUCT(C{r}:{lc}{r},C$5:{lc}$5)'
        '/SUMPRODUCT((C{r}:{lc}{r}<>"")*(C$5:{lc}$5)))'.format(
            r=r, lc=last_mark_col))
    ws.cell(r, 3 + len(labels)).font = BODY_B
    ws.cell(r, 3 + len(labels)).number_format = "0.0"
    ws.cell(r, 4 + len(labels)).value = (
        '=IF({tc}{r}="","",IF({tc}{r}>=70,"Distinction",'
        'IF({tc}{r}>=50,"Pass","Not yet")))'.format(tc=tot_col, r=r))
    ws.cell(r, 4 + len(labels)).font = BODY

ws.cell(FIRST_DATA + N_ROWS + 1, 1,
        "The weighted total divides by the weight of the pieces ACTUALLY "
        "marked, so a participant halfway through the course shows their "
        "standing on what they have done rather than being penalised for "
        "work not yet set.").font = S_FONT
ws.cell(FIRST_DATA + N_ROWS + 2, 1,
        "Bands are a placeholder: 70 distinction, 50 pass. Replace them with "
        "the awarding body's if there is one.").font = S_FONT
ws.freeze_panes = "C6"

# ============================================================ 5. datasets
ws = wb.create_sheet("Datasets")
title_block(ws, "Datasets, with licences checked",
            "Nothing in this list is NonCommercial or internal-use, so the "
            "track is safe for commercial delivery", 6)
header_row(ws, 4, ["Dataset", "Where", "Licence", "Shape", "Account",
                   "What it is for"], [28, 44, 30, 40, 14, 60])
r = 5
for ds in C.DATASETS:
    write_row(ws, r, [ds["name"], ds["url"], ds["licence"], ds["shape"],
                      ds["account"], ds["use"]], height=62)
    r += 1
r += 1
c = ws.cell(r, 1, "Considered and rejected")
c.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
c.fill = PatternFill("solid", fgColor=MID)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
r += 1
for name, why in C.REJECTED:
    write_row(ws, r, [name, why], fill=AMBER, height=48)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    r += 1
ws.freeze_panes = "A5"

# ============================================================ 6. the answers
ws = wb.create_sheet("Answer key")
title_block(ws, "Answer key for the course dataset",
            "Computed by data/build_teaching_data.py from the same data the "
            "participants get. Re-run the generator and this sheet is "
            "rebuilt with it.", 3)
header_row(ws, 4, ["What", "Value", "Where it is used"], [46, 26, 54])
rows = [
    ("Orders", K["s2"]["orders_total"], "SQL lab question 1"),
    ("Orders, delivered", K["s2"]["orders_delivered"], "question 2"),
    ("Orders, cancelled", K["s2"]["orders_cancelled"], "question 2"),
    ("COUNT(*) on orders", K["s2"]["count_star_vs_customer_id"][0],
     "question 3, the NULL lesson"),
    ("COUNT(customer_id) on orders",
     K["s2"]["count_star_vs_customer_id"][1], "question 3"),
    ("Revenue, all lines, UGX", K["s3"]["revenue_all_ugx"], "question 6"),
    ("Revenue, delivered, UGX", K["s3"]["revenue_delivered_ugx"],
     "question 7, and the Python lab target"),
    ("Revenue lost to an INNER JOIN, UGX",
     K["s3"]["revenue_lost_to_inner_join_ugx"],
     "question 8, the walk-in trap"),
    ("Orders with no customer", K["s3"]["inner_join_drops_rows"],
     "question 8"),
    ("Value of everything returned, UGX", K["s3"]["returns_value_ugx"],
     "session 3 report"),
    ("Revenue from the raw CSVs, UGX",
     K["s3"]["revenue_delivered_from_raw_ugx"],
     "where a correct pandas clean actually lands"),
    ("Cent-rounding residual, UGX", K["s3"]["usd_rounding_residual_ugx"],
     "the precision lesson in session 5"),
    ("Mean order value, UGX", K["s7"]["order_value_mean_ugx"],
     "session 7, mean against median"),
    ("Median order value, UGX", K["s7"]["order_value_median_ugx"],
     "session 7"),
    ("Mean divided by median", K["s7"]["mean_over_median"],
     "session 7, how skewed it is"),
    ("90th percentile order value, UGX", K["s7"]["order_value_p90_ugx"],
     "session 7, percentiles"),
]
r = 5
for what, val, where in rows:
    write_row(ws, r, [what, val, where], height=16)
    if isinstance(val, float) and abs(val) > 1000:
        ws.cell(r, 2).number_format = "#,##0"
    r += 1
r += 1
ws.cell(r, 1, "Region and month breakdowns are in "
              "DataAnalysis_Track/data/answer_key.json, which is the "
              "machine-readable version of this sheet.").font = S_FONT
ws.freeze_panes = "A5"

# ============================================================ 7. glossary
ws = wb.create_sheet("Glossary")
title_block(ws, "Every word the track uses",
            "{} terms, generated from the sessions themselves so it cannot "
            "drift from what was taught".format(len(C.GLOSSARY)), 4)
header_row(ws, 4, ["Term", "What it means", "In a spreadsheet",
                   "Taught in session"], [24, 70, 34, 14])
def spreadsheet_safe(text):
    """Stop Excel error tokens in glossary text from becoming errors.

    "#VALUE!, with an explanation attached" is the honest way to describe a
    traceback to an Excel user, and LibreOffice parses a cell starting with
    it as the error value itself. recalc.py then reports three errors in a
    sheet that contains no formulas. The decks and the printed glossary keep
    the real tokens, because there the words are just words.
    """
    SAFE = {"#VALUE!": "the VALUE error", "#N/A": "the N/A error",
            "#REF!": "the REF error", "#NAME?": "the NAME error",
            "#DIV/0!": "the DIV/0 error"}
    for tok, plain in SAFE.items():
        text = text.replace(tok, plain)
    # a leading "the" reads oddly mid-sentence, so tidy the commonest case
    return text.replace("delete the the ", "delete the ").replace(
        "keep the the ", "keep the ")


r = 5
for term, definition, excel, sess in C.GLOSSARY:
    write_row(ws, r, [term, spreadsheet_safe(definition),
                      spreadsheet_safe(excel) if excel
                      else "no direct equivalent",
                      sess], height=30)
    if not excel:
        ws.cell(r, 3).font = Font(name=FONT, size=9, italic=True,
                                  color="9C4500")
    r += 1
r += 1
ws.cell(r, 1, "The {} terms with no spreadsheet equivalent are the "
              "genuinely new ideas in this course.".format(
                  sum(1 for _, _, x, _ in C.GLOSSARY if not x))).font = S_FONT
ws.freeze_panes = "A5"

# ============================================================ 8. what was cut
ws = wb.create_sheet("What was cut")
title_block(ws, "Cut to pay for the foundations",
            "Recorded so it does not creep back in, and so a fast cohort "
            "knows what to add", 2)
header_row(ws, 4, ["What came out", "Why, and what replaced it"], [30, 100])
r = 5
for what, why in C.CUT:
    write_row(ws, r, [what, why], height=58)
    r += 1

wb.save(OUT)
print("saved " + OUT)
print("  sheets: " + ", ".join(wb.sheetnames))
