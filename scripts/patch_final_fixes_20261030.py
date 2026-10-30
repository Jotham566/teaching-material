"""Last text repairs from the 30 October reviews. Idempotent.
  * Week 5 title slide (Friday cohorts): the nominal class time contradicted the 'no normal class' notice;
  * Week 12 optional video repeated the Week 10 video;
  * the exit slides repeated 'Save your work' twice;
  * code-slide captions that name a variable the code does not use.
Usage: python3 scripts/patch_final_fixes_20261030.py [--dir DIR]"""
import argparse, re, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation
from pptx.util import Pt
from deckkit_20261030 import *
import week_reference_numpy as R
from patch_weeks8b_12_20261030 import deck_name, first_title, DECKS

OLD_VIDEO = "https://www.youtube.com/watch?v=fSytzGwwBVw"
NEW_VIDEO = "https://www.youtube.com/watch?v=EuBBz3bI-aA"
GENERIC = "Run the cell and read the output below it."


def courier(slide):
    return "\n".join(p.text for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if any(r.font.name == "Courier New" for r in p.runs))


def fix_captions(prs):
    n = 0
    for sl in prs.slides:
        code = courier(sl)
        if not code: continue
        title = first_title(sl)
        for sh in sl.shapes:
            if not sh.has_text_frame or any(r.font.name == "Courier New" for p in sh.text_frame.paragraphs for r in p.runs): continue
            t = sh.text_frame.text.strip(); m = re.match(r"^(\w+): ", t)
            if m and "\n" not in t and not t.isupper() and t != title and not re.search(rf"\b{re.escape(m.group(1))}\b", code):
                n += replace_text(sl, t, GENERIC)
    return n


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--dir", default=None); a = ap.parse_args(); d = Path(a.dir) if a.dir else DECKS
    for code in R.COHORT:
        for week in range(1, 16):
            path = d / deck_name(code, week); prs = Presentation(str(path)); n = 0
            if week == 5 and code in ("CML3111", "CML3113"):
                n += replace_text(prs.slides[0], "Friday 11:00–14:00, Room 115", "No normal class on this date: replacement session as announced")
                n += replace_text(prs.slides[0], "Friday 08:00–11:00, Room 105", "No normal class on this date: replacement session as announced")
            if week == 12:
                v = find_slide(prs, "Optional revision video")
                for sh in v.shapes:
                    if not sh.has_text_frame: continue
                    for p in sh.text_frame.paragraphs:
                        for r in p.runs:
                            if r.hyperlink.address == OLD_VIDEO: r.hyperlink.address = NEW_VIDEO; n += 1
                n += replace_text(v, "Machine Learning Fundamentals Cross Validation", "Machine Learning Fundamentals: Bias and Variance")
                n += replace_text(v, "Listening question: Why evaluate using observations not used to fit that particular model?", "Listening question: What goes wrong with a model that is too simple, and with one that is too flexible?")
            for sl in prs.slides: n += replace_text(sl, "Save your work and check the required submission files.", "Check that the required files are in place.")
            if 8 <= week <= 12: n += fix_captions(prs)
            prs.save(str(path)); print(f"{code} W{week}: {n} repairs") if n else None


if __name__ == "__main__":
    main()
