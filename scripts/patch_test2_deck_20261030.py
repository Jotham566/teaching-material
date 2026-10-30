"""Week 14 decks: Practical Test 2 now covers Weeks 8 to 13 (idempotent).

Inserts a revision section (what the test covers, six words, one worked example for each new task) before the
'Test 2 revision checklist' slide, updates the marks and order-of-work slides, and sets spoken notes on everything it
touches. Run after apply_spoken_notes_20261030.py, because that script rewrites notes on the older Week 14 slides.
Usage: python3 scripts/patch_test2_deck_20261030.py [--dir DIR]"""
import argparse, glob, io, contextlib, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation
from deckkit_20261030 import *
from test2_deck_examples_20261030 import EXAMPLES as X
import patch_week8_20261030 as p8

DECKS = p8.ROOT / "IUEA_Teaching_Plan" / "decks"
FIRST = "The test now reaches back to Week 13"
ANCHOR = "Test 2 revision checklist"
INSERTED = [FIRST, "Words for the later weeks", "Counting warnings and their cost", "Cross-validation by equipment group", "Does an extra input help?",
            "Choosing a model inside two limits", "Ridge and nearest neighbours: build the models (1/2)", "Ridge and nearest neighbours: fit and score (2/2)",
            "Test 2 checklist for Weeks 9 to 13"]


def run(code):
    """Run a worked example against the NumPy stand-in or real scikit-learn and return what it prints."""
    import verify_week_slides_20261030 as V
    try:
        import sklearn  # noqa: F401
    except ImportError:
        V.install_shim()
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf): exec(compile(code, "example", "exec"), {})
    return buf.getvalue().strip()


def patch_deck(path):
    prs = Presentation(str(path)); code = Path(path).name[:7]; lab = lambda: f"{code}   Week 14   0"
    # --- existing slides: marks, order of work, number of tasks
    for old, new in [("Code tasks have a prompt, function name and marks.", "Eight code tasks, each with a prompt, function name and marks."),
                     ("Equipment-group split 20.", "Weeks 8 to 10: split 10, predictions 15, MAE and RMSE 5, group cross-validation 10."),
                     ("Baseline, linear and tree predictions 30.", "Weeks 9, 11 and 12: classification and cost 10, feature ablation 10, model selection 10."),
                     ("MAE and RMSE 20. Structured interpretation 30.", "Week 13: ridge and nearest neighbours 15. Six interpretation choices 15."),
                     ("Compare validation errors. Keep the specified final-test group untouched.", "Work down the eight tasks in order. Each is checked on its own, so move on if one is stuck.")]:
        for s in prs.slides: replace_text(s, old, new)
    notes = {
        "One notebook contains all answers": spoken(
            "Everything for the test lives in one notebook.",
            "There are eight code tasks. Each has a prompt, a function name and the marks it carries. Read the prompt, and write your code in the function with the name supplied.",
            "There are also six interpretation questions. For these you record your chosen option letters in the DECISIONS dictionary, as we practised.",
            "When you finish, you submit one notebook snapshot using the final button. A snapshot is a saved copy of your notebook at that moment."),
        "How the 100 marks are allocated": spoken(
            "The test is marked out of one hundred, and it reaches back over Weeks 8 to 13.",
            "From Weeks 8 to 10 we have the equipment-group split for ten marks, the baseline, linear and tree predictions for fifteen, MAE and RMSE for five, and cross-validation by equipment group for ten.",
            "From Weeks 9, 11 and 12 we have classification and cost for ten marks, the feature ablation for ten, and choosing a model under limits for ten.",
            "From Week 13 we have the ridge curve and nearest neighbours for fifteen marks. The six interpretation choices are worth fifteen marks together.",
            "Every one of these steps is something we have practised. The revision slides before this section cover each of them."),
        "Order of work": spoken(
            "The test has a sensible order, and it follows the practice closely.",
            "We begin by running the row filters that the notebook supplies. After that we split the data by equipment ID, so that each machine goes wholly to one group.",
            "Next we fit the imputer and the models on the training records only. The training rows teach the pipelines, and the other rows stay out of the learning.",
            "Then we work down the eight tasks in order. Each task is checked on its own, so an error in one does not stop the grader from checking the next. If a task will not work, leave it and move on, and come back if you have time.",
            "Throughout, the final-test group, E6, stays untouched until the very end."),
    }
    for s in prs.slides:
        t = slide_title_texts(s)
        for k, v in notes.items():
            if k in t[:3]: set_notes(s, v)
    # --- inserted revision section
    for t in INSERTED:                      # rebuild the section from scratch so the script can be re-run after edits
        for s in [x for x in prs.slides if t in slide_title_texts(x)[:3]]: delete_slide(prs, s)
    if find_slide(prs, FIRST) is None:
        anchor = find_slide(prs, ANCHOR); assert anchor is not None, "checklist slide not found"
        new = []
        new.append(idea_slide(prs, "WHAT TO DO", FIRST,
            [("EIGHT CODING TASKS", True), "Week 8: predictions and errors. Week 9: warnings and cost. Week 10: the split and cross-validation. Week 11: does an extra input help? Week 12: choose under limits. Week 13: ridge and neighbours."],
            [("SIX READING CHOICES", True), "Splits, fill values, cost, the sealed test rows, the ridge penalty and scaling. Record the letters in DECISIONS."],
            spoken("The test has grown since the first plan. It now reaches back over Weeks 8 to 13, so that it checks what we have really learned this term.",
                   "There are eight coding tasks and six reading choices. From Week 8 we have the predictions and the two error measures. From Week 9 we have counting warnings and working out their cost. From Week 10 we have the split by equipment and cross-validation. From Week 11 we have asking whether an extra input helps. From Week 12 we have choosing a model inside two limits. From Week 13 we have the ridge curve and the nearest neighbours.",
                   "The next few slides revise the later weeks with tiny examples, so that nothing in the test is a surprise. Everything we need is a small step from something we have already written."), lab(), size=20))
        new.append(table_slide(prs, "WORDS WE USE", "Words for the later weeks",
            ["Word", "Plain meaning", "Week"],
            [["Fold", "One round of cross-validation: one group held back, the rest used to fit.", "10"],
             ["Cross-validation", "Holding each group back in turn, then reporting the average and the spread.", "10"],
             ["Ablation", "Adding or removing one input and comparing the error on the same rows.", "11"],
             ["Ridge penalty", "A price on large coefficients that calms a wiggly curve. Alpha sets the price.", "13"],
             ["Nearest neighbours", "Predict from the average of the closest training rows.", "13"],
             ["Scaling", "Putting inputs on a common ruler so that one big number does not decide the distances.", "13"]],
            [200, 590, 80],
            spoken("Six words to keep in mind. A fold is one round of cross-validation: one group is held back and the rest are used to fit. Cross-validation means holding each group back in turn, and then reporting the average error and how far the errors spread.",
                   "Ablation means adding or removing one input and comparing the error on the same rows. A ridge penalty is a price on large coefficients, and it calms a wiggly curve. The number alpha sets the price.",
                   "Nearest neighbours predicts from the average of the closest training rows. Scaling puts the inputs on a common ruler, so that a column with big numbers does not decide all the distances.",
                   "Every one of these words came from a lesson in Weeks 10 to 13. Nothing here is new."), lab(), size=17))
        out = run(X["classify"])
        new.append(code_slide(prs, "WORKED EXAMPLE", "Counting warnings and their cost", X["classify"], out,
            ["tp: warnings given and needed. fn: warnings missed. fp: false alarms.", "zip pairs each actual label with its predicted label.", "A missed warning costs 5 and a false alarm costs 1."],
            "Label 1 means a warning. One warning was missed, so the cost is 5.",
            spoken("Task 4 asks us to count the outcomes of a warning model and work out a cost. Label one means a warning was needed, or a warning was given.",
                   "We pair each actual label with its predicted label using zip. Then we count three things. A true positive, tp, is a warning that was needed and given. A false negative, fn, is a warning that was needed and missed. A false positive, fp, is a warning given when none was needed, a false alarm.",
                   "The cost rule in the test is five for each missed warning and one for each false alarm. Here one warning was missed and there were no false alarms, so the cost is five. In the test we also work out precision and recall, and we use zero when a denominator would be zero."), lab(), code_size=14, out_size=16))
        out = run(X["groupcv"])
        new.append(code_slide(prs, "WORKED EXAMPLE", "Cross-validation by equipment group", X["groupcv"], out,
            ["Each pass holds one group back and fits on the others.", "A held-back group never appears in its own training rows.", "The test asks for E1 to E5, then the mean and the spread of the five errors."],
            "Three groups give three passes and three errors.",
            spoken("Task 5 repeats the check so that every group is held back once. This is cross-validation by equipment group.",
                   "The loop picks one group to hold out. The training rows are all the other groups, and the test rows are the held-out group. We fit a straight line on the training rows, predict the held-out rows, and print the error. Then the loop moves to the next group.",
                   "In the real test there are five groups, E1 to E5, and E6 is never used. The function returns the five errors in order, their average, and their spread, which is the sample standard deviation. A small spread means the model behaves much the same on every group."), lab(), code_size=12, out_size=16))
        out = run(X["ablation"])
        new.append(code_slide(prs, "WORKED EXAMPLE", "Does an extra input help?", X["ablation"], out,
            ["The new input is the load multiplied by the temperature.", "Both models use the same training rows and the same validation rows.", "The only change is one extra column."],
            "On this tiny file the extra input lowers the error. On the test file, report what the numbers show.",
            spoken("Task 6 asks whether an extra input earns its place. This is an ablation: one change, the same rows, the same error measure.",
                   "First we make the new input, load times temperature, in both tables. Then the loop fits two straight-line models. The first uses load and temperature. The second adds the new column. Each is fitted on the training rows and scored on the validation rows.",
                   "Compare the two printed errors. Here the second is smaller, so the extra input helped on this tiny file. The test asks for both errors and their difference, and it does not promise which way the answer goes."), lab(), code_size=12, out_size=16))
        out = run(X["select"])
        new.append(code_slide(prs, "WORKED EXAMPLE", "Choosing a model inside two limits", X["select"], out,
            ["kept holds the models that meet both limits.", "sorted orders them by error, then speed, then name.", "The first one in the sorted list is the choice."],
            "The tree is too slow and the mean is too inaccurate. Of the rest, knn has the smaller error.",
            spoken("Task 7 chooses a model under two limits, one for the error and one for the time a prediction takes. Each candidate is a dictionary with a name, a validation error and a time in milliseconds.",
                   "First we keep only the candidates that meet both limits. The tree is dropped because it takes thirty milliseconds, and the mean is dropped because its error is four. Two are left, the line and knn.",
                   "Then we sort. The smallest error comes first, then the shorter time, then the name in alphabetical order. The first one in the sorted list is knn. If nothing meets both limits, the test wants the answer None. We never use the test scores to choose."), lab(), code_size=12, out_size=16))
        out_b = "(nothing printed)"
        new.append(code_slide(prs, "WORKED EXAMPLE", "Ridge and nearest neighbours: build the models (1/2)", X["ridgeknn_build"], out_b,
            ["The ridge pipeline makes a cubic curve of load, scales it, then fits with a penalty.", "The neighbour pipeline scales both inputs first.", "Nothing is fitted yet."],
            "Two recipes. Each one will be fitted on the training rows only.",
            spoken("Task 8 compares two ideas from Week 13. First we build both models. Nothing is fitted yet, so nothing prints.",
                   "The ridge pipeline has three steps. PolynomialFeatures makes load, load squared and load cubed. StandardScaler puts those three columns on a common ruler. Ridge fits a curve with a penalty, and alpha equal to one sets the price of the penalty.",
                   "The neighbour pipeline has two steps. StandardScaler puts load and temperature on a common ruler, so that neither decides all the distances. Then KNeighborsRegressor predicts from the average of the three nearest training rows."), lab(), code_size=14, out_size=16))
        out = run(X["ridgeknn_build"] + "\n" + X["ridgeknn_score"])
        new.append(code_slide(prs, "WORKED EXAMPLE", "Ridge and nearest neighbours: fit and score (2/2)", X["ridgeknn_score"], out,
            ["Both models are fitted on the training rows only.", "The ridge model sees load alone. The neighbour model sees load and temperature.", "Both are scored on the same validation rows."],
            "Same rows, same measure. The smaller error is better.",
            spoken("Now we fit both models on the training rows and score them on the validation rows. The ridge model sees load alone. The neighbour model sees load and temperature, because that is how the task is written.",
                   "The loop predicts the validation rows with each model and prints the mean absolute error. Here the ridge error is a little over one, and the neighbour error is two. So ridge is better on this tiny file.",
                   "In the test, the function returns both errors without rounding, and the word ridge or knn for the smaller one. Every step is fitted on the training rows only. The validation rows are used only to score."), lab(), code_size=12, out_size=16))
        new.append(idea_slide(prs, "WHAT TO DO", "Test 2 checklist for Weeks 9 to 13",
            [("I CAN", True), "Count tp, fn and fp and work out a cost. Hold each group back in turn and report the mean and spread. Compare two models on the same rows. Choose the best model inside two limits."],
            [("I CAN EXPLAIN", True), "Why a very large ridge penalty can hurt. Why neighbours need scaled inputs. Why the test rows stay sealed. Why every fit uses training rows only."],
            spoken("Here is a short checklist for the later weeks. Before the test, we should be able to do four things. Count true positives, false negatives and false alarms and work out a cost. Hold each group back in turn and report the average error and the spread. Compare two models on the same rows. And choose the best model inside two limits.",
                   "We should also be able to explain four ideas in plain words. A very large ridge penalty can flatten the curve and make both errors worse. Nearest neighbours need scaled inputs, because distances depend on the size of the numbers. The test rows stay sealed until the choice is made. And every fit, including every fill value and every scaler, uses training rows only.",
                   "If you can do these and explain these, you are ready for the later part of the test."), lab(), size=20))
        for s in new: move_slide(prs, s, before_slide=anchor)
    finalize(prs, code, 14, path)
    return len(prs.slides)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--dir", default=None); a = ap.parse_args()
    d = Path(a.dir) if a.dir else DECKS
    for f in sorted(glob.glob(str(d / "CML31??_Week14_*.pptx"))):
        print(Path(f).name, patch_deck(f), "slides")


if __name__ == "__main__":
    main()
