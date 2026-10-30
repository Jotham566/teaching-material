"""Regenerate the lecturer-notes Markdown for all 45 decks, and the guided-practice notebooks (Weeks 8 to 13),
from the finished decks. Also refreshes guided_practice.ipynb inside the Week 8 to 12 student ZIPs.
Usage: python3 scripts/build_companions_20261030.py"""
import json, re, shutil, sys, zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation
import week_reference_numpy as R
from patch_weeks8b_12_20261030 import deck_name, first_title, DECKS, p8
from verify_week_slides_20261030 import courier_blocks

ROOT = p8.ROOT; BASE = ROOT / "IUEA_Teaching_Plan"
SKIP = ("quiz", "submission", "Optional revision video", "Revision and submission checklist", "Project preparation")
COURSE = {"CML3111": "Engineering", "CML3112": "Design and Infrastructure", "CML3113": "Mining and Petroleum"}


def plain_text(slide, title):
    out = []
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        if any(r.font.name == "Courier New" for p in sh.text_frame.paragraphs for r in p.runs): continue
        t = sh.text_frame.text.strip()
        if not t or t.isupper() or t == title or re.match(r"CML311\d\s+Week \d+\s+\d+$", t): continue
        out.append(t)
    return out


def slide_rows(path):
    prs = Presentation(str(path)); seen = {}; rows = []
    for i, s in enumerate(prs.slides, 1):
        t = first_title(s); n = seen.get(t, 0); seen[t] = n + 1
        rows.append(dict(i=i, title=t, text=plain_text(s, t), code=courier_blocks(s), notes=s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""))
    return rows


def write_md(code, week):
    rows = slide_rows(DECKS / deck_name(code, week)); out = [f"# {code} Week {week}", "", "Each slide shows its text, then the spoken script the lecturer reads aloud.", ""]
    for r in rows:
        out += [f"## Slide {r['i']}: {r['title']}", ""]
        for t in r["text"]: out += [t, ""]
        for cb in r["code"][:1]: out += ["```python", cb, "```", ""]
        if len(r["code"]) > 1: out += ["Output on the slide:", "", "```text", r["code"][1], "```", ""]
        out += ["**Spoken script**", "", r["notes"], ""]
    (DECKS / "lecturer_notes" / deck_name(code, week).replace(".pptx", ".md")).write_text("\n".join(out))


def notebook(code, week):
    rows = slide_rows(DECKS / deck_name(code, week)); c = R.COHORT[code]
    cells = []
    def md(txt): cells.append({"cell_type": "markdown", "metadata": {}, "source": txt.splitlines(True)})
    def cd(txt): cells.append({"cell_type": "code", "metadata": {}, "execution_count": None, "outputs": [], "source": txt.splitlines(True)})
    files = "engineering_data.csv" + (" and classification_case.csv" if week == 9 else "")
    md(f"# Week {week} guided practice\n\n{code}. Follow the lecturer’s PowerPoint in this order. This notebook contains worked examples, not your graded upload.\n\n"
       f"Extract the student ZIP and upload {files} into Colab with the Files panel before running the cells. The target in this notebook is {c['short']} in {c['unit']}.")
    buf = []
    for r in rows:
        if r["i"] == 1 or any(k in r["title"] for k in SKIP): continue
        title = re.sub(r"\s*\(\d+/\d+\)\s*$", "", r["title"])
        if r["code"]:
            m = re.search(r"\((\d+)/(\d+)\)\s*$", r["title"])
            if m:
                buf.append(r["code"][0])
                if int(m.group(1)) < int(m.group(2)): continue
                src = "\n".join(buf); buf = []
            else: src = r["code"][0]
            md(f"## {title}\n\n*PowerPoint slide {r['i']}*\n\n" + "\n\n".join(r["text"][:2]))
            cd(src)
            if len(r["code"]) > 1 and not r["code"][1].startswith("("): md("**Example output**\n\n```text\n" + r["code"][1] + "\n```")
        elif r["text"]:
            md(f"## {title}\n\n*PowerPoint slide {r['i']}*\n\n" + "\n\n".join(r["text"]))
    return {"cells": cells, "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}, "language_info": {"name": "python"}}, "nbformat": 4, "nbformat_minor": 5}


def main():
    zbackup = BASE / "student_packages" / "_backup_2026-10-30"; zbackup.mkdir(exist_ok=True)
    for code in R.COHORT:
        for week in range(1, 16): write_md(code, week)
        for week in range(8, 14):
            nb = notebook(code, week); d = BASE / "guided_practice" / code / f"W{week:02d}"; d.mkdir(parents=True, exist_ok=True)
            (d / "guided_practice.ipynb").write_text(json.dumps(nb, indent=1, ensure_ascii=False))
            z = BASE / "student_packages" / f"{code}_W{week:02d}_student.zip"
            if z.exists():
                if not (zbackup / z.name).exists(): shutil.copy2(z, zbackup / z.name)
                src = z; tmp = z.with_suffix(".tmp")
                with zipfile.ZipFile(src) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
                    for item in zi.infolist():
                        zo.writestr(item, json.dumps(nb, indent=1, ensure_ascii=False) if item.filename == "guided_practice.ipynb" else zi.read(item.filename))
                tmp.replace(z)
            print(f"{code} W{week}: {len(nb['cells'])} notebook cells" + ("; zip refreshed" if z.exists() else "; no zip for this week"))


if __name__ == "__main__":
    main()
