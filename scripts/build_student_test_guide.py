"""Build the shared beginner guide without including test answers or private links."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'IUEA_Teaching_Plan/test_operations'
ASSETS=OUT/'student_guide_images'; ASSETS.mkdir(exist_ok=True)
fontpath='/System/Library/Fonts/Supplemental/Arial.ttf'
def diagram(name, labels, sublabels):
 im=Image.new('RGB',(1500,210),'white'); dr=ImageDraw.Draw(im)
 font=ImageFont.truetype(fontpath,31); small=ImageFont.truetype(fontpath,25)
 width=440
 for i,(label,sub) in enumerate(zip(labels,sublabels)):
  x=10+i*510
  dr.rounded_rectangle((x,20,x+width,180),radius=18,fill='#eef4f7',outline='#b3c8d4',width=2)
  dr.text((x+width/2,65),label,font=font,fill='#12344a',anchor='mm')
  dr.text((x+width/2,117),sub,font=small,fill='#213945',anchor='mm')
  if i<2:
   dr.line((x+width+10,100,x+width+60,100),fill='#087e8b',width=5)
   dr.polygon([(x+width+60,100),(x+width+46,90),(x+width+46,110)],fill='#087e8b')
 path=ASSETS/name; im.save(path); return path
files=diagram('open_the_correct_files.png',['Open the test link','Notebook .ipynb','Data .csv'],['Enter your own details','Download and open in Colab','Download and upload in Files'])
submit=diagram('finish_both_submission_steps.png',['Submit test','Open the form link','Submit on the form'],['Creates a saved snapshot','Check your own account','Wait for confirmation'])
d=Document(); sec=d.sections[0]; sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.62);sec.left_margin=sec.right_margin=Inches(.7)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(11)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.05
d.styles['Title'].font.size=Pt(23);d.styles['Title'].paragraph_format.space_after=Pt(8)
d.styles['Heading 1'].font.size=Pt(14);d.styles['Heading 1'].paragraph_format.space_before=Pt(10);d.styles['Heading 1'].paragraph_format.space_after=Pt(5)
def p(text,bold=False):
 para=d.add_paragraph();para.add_run(text).bold=bold;return para
def h(text):d.add_heading(text,1)
def link(label,url):
 para=d.add_paragraph();el=OxmlElement('w:hyperlink');el.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');pr.append(c);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);el.append(r);para._p.append(el)
def pic(path,alt):
 shape=d.add_picture(str(path),width=Inches(6.85));shape._inline.docPr.set('descr',alt)
d.add_paragraph('Your Colab test guide','Title')
p('Test 1 and Test 2 • CML3111, CML3112 and CML3113')
p('Use this guide to open your test, load its data and submit your answers. Have your laptop, charger, internet connection and your own Google account ready. Use the same account throughout.')
h('1  Open your test page and get your files')
p('Open the test entry link your lecturer shares, or scan its QR code. Check the course and test number. Enter your full name, Google email and Student Registration Number, then click Get my test files. Download the notebook and CSV to your laptop. Returning with the same details gives you the same files.')
p('The .ipynb file is your notebook: it contains questions and spaces for code. engineering_data.csv is its data table. Keep these two files together. If you scan on a phone, open the same link on your laptop and enter the same details there. You do not need to choose a test version or extract a ZIP.')
pic(files,'Open the test page and enter your details. Download its notebook and CSV. Open the notebook in Colab and upload the CSV in Files.')
h('2  Open the notebook in Colab')
link('Open Google Colab — colab.research.google.com','https://colab.research.google.com/')
p('Sign into your own Google account. Choose File → Open notebook → Upload, then select TEST1_student.ipynb or TEST2_student.ipynb from your Downloads folder. If the opening window already shows Upload, select it directly.')
p('Choose File → Save a copy in Drive and work in that copy. Click Connect near the top right. This starts the computer session that runs your Python code; wait until it connects.')
h('3  Upload the CSV and enter your details')
p('Click the folder icon on the left to open Files. Click Upload (the upward-arrow icon), choose the engineering_data.csv you downloaded from your test page and wait until it appears. Keep its name unchanged. This puts the data into Colab; it does not submit your test.')
p('Find the first code cell with FULL_NAME, EMAIL and STUDENT_ID. Your details should already be filled in. Check them. Ask the invigilator if your registration number or email is wrong. STUDENT_ID means your Student Registration Number: enter it exactly as on your university record. Keep COHORT, TEST_ID and the final submission cell unchanged.')
h('4  Run cells and complete the questions')
p('A cell is one block of text or code. Click the triangle at the left of a code cell to run it, or click inside it and press Shift + Enter. Editing code alone does not run it. Work from the top down.')
p('Run the identity cell, then the supplied data-loading cell. You should see the first data rows and a row count. Read each question before replacing its placeholder with your code. Keep the required function names and answer formats. Put your interpretation choices in DECISIONS as instructed.')
d.add_page_break()
d.add_paragraph('Check and submit your test','Title')
p('Start these steps before the deadline announced by your lecturer. Do not leave submission until the final minute.')
h('5  Check your work from the beginning')
p('Save your notebook. Choose Runtime → Restart session and run all. This checks whether the notebook works from a fresh start. If the CSV is missing, upload the same engineering_data.csv again and rerun from the top.')
p('Fix the first error you can understand, then rerun the affected cells. Keep any unfinished tasks in the notebook. If you cannot finish a question, still submit your completed work before the deadline; partial credit is available.')
h('6  Prepare your submission in the notebook')
p('Run the final submission cell if it has not run. When Google asks, select the SAME account you used for this notebook and review the access request. The supplied code saves a separate notebook snapshot in your Drive and shares that snapshot privately with the lecturer. Ask the invigilator if you are unsure.')
p('Wait for the green Submit test button, then click it once. A snapshot is a saved copy of your answers at that moment. Wait for the link labelled “Finish submission: open the form and press Submit”.')
pic(submit,'Submission has three steps: click Submit test in the notebook, open its generated form link, then press Submit on the form and wait for confirmation.')
h('7  Finish submission on the Google Form')
p('Open the generated link. Check the course/test, your name and Student Registration Number. Check the signed-in Google email is the SAME account used in Colab. If your registration number or email is wrong, ask the invigilator to correct your allocation before submitting.')
p('Leave Test allocation code, Snapshot file ID and Snapshot SHA-256 unchanged; the button fills these automatically. Tick the email confirmation box, scroll to the bottom and click Submit. Wait for the page confirming that your response was recorded. Keep a screenshot of that confirmation for yourself.')
p('You are finished only after the form confirms submission. The notebook button alone is not enough.',bold=True)
p('No manual notebook download or upload is needed. Do not submit the ZIP, CSV, PDF or a screenshot of your answers. Keep the saved snapshot in Drive unchanged until marks are released. Marks are issued after the lecturer grades the batch.')
h('If something goes wrong')
p('CSV not found: check the Files sidebar and upload the correct file without renaming it. NameError: run the earlier cell that defines the missing name, especially your identity cell. No button: connect and run the final submission cell, then complete the Google prompt.')
p('Form closed, wrong account, failed submission or lost connection: tell the invigilator immediately, before the deadline. Do not assume your work was submitted. To replace an attempt before the deadline, prepare a new snapshot and submit its new form link.')
p('Work individually. Do not use AI-generated answers, messaging or copied work. Follow the lecturer’s rules for any other resources.',bold=True)
for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for border in list(d.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.core_properties.title='Your Colab test guide';d.core_properties.author='IUEA Teaching Resources'
path=OUT/'Student_Colab_Test_Guide.docx';d.save(path);print(path)
