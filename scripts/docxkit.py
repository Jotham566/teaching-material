# -*- coding: utf-8 -*-
"""Shared docx helpers for the IUEA teaching plan documents (python-docx)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x38, 0x64)
MID = RGBColor(0x2E, 0x5C, 0x8A)
GREY = RGBColor(0x59, 0x59, 0x59)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NAVY_H = "1F3864"
MID_H = "2E5C8A"
BAND_H = "F2F5FB"
LIGHT_H = "D9E2F3"
AMBER_H = "FFF2CC"
RED_H = "FCE4E4"
GREEN_H = "E2EFDA"


# ---------------------------------------------------------------- document
def new_doc(margin_cm=2.4):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.line_spacing = 1.15

    for name, size, colour, before, after in [
        ("Heading 1", 20, NAVY, 20, 9),
        ("Heading 2", 14, MID, 14, 7),
        ("Heading 3", 11.5, BLACK, 11, 5),
    ]:
        s = doc.styles[name]
        s.font.name = FONT
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = colour
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    sec = doc.sections[0]
    sec.top_margin = Cm(margin_cm)
    sec.bottom_margin = Cm(margin_cm)
    sec.left_margin = Cm(margin_cm)
    sec.right_margin = Cm(margin_cm)
    return doc


def usable_width(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


# ---------------------------------------------------------------- text
def para(doc, text, size=10.5, bold=False, italic=False, colour=None,
         align=None, after=8, before=0, space=1.15, indent=None, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if colour is not None:
        r.font.color.rgb = colour
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = space
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.right_indent = Cm(indent)
    return p


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def h3(doc, text):
    return doc.add_heading(text, level=3)


def quote(doc, text):
    p = para(doc, text, size=10.5, italic=True, colour=MID, indent=0.9, after=12, before=8)
    _border(p, "left", MID_H, 18)
    return p


def rule(doc, colour=LIGHT_H):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.space_before = Pt(2)
    _border(p, "bottom", colour, 8)
    return p


def spacer(doc, pts=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _border(p, edge, colour, size):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(size))
    e.set(qn("w:space"), "8")
    e.set(qn("w:color"), colour)
    bdr.append(e)


# ---------------------------------------------------------------- shading
def shade(cell, hexcolour):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolour)
    tcPr.append(shd)


def cell_margins(cell, top=60, bottom=60, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, v in (("top", top), ("start", left), ("bottom", bottom), ("end", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def table_borders(tbl, colour="B4C6E7", inner="D9E2F3"):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag, col, sz in (("top", colour, 6), ("left", colour, 6), ("bottom", colour, 6),
                         ("right", colour, 6), ("insideH", inner, 4), ("insideV", inner, 4)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), col)
        borders.append(e)
    tblPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


# ---------------------------------------------------------------- tables
def _fixed_grid(tbl, widths):
    """Pin the column widths.

    Setting autofit=False and a width on every cell is not enough: LibreOffice
    (and Word, on a table whose cells disagree) falls back to its own layout
    unless the table carries an explicit grid and a fixed layout algorithm.
    Without this a weight of 4 for a "#" column renders as wide as a weight of
    34, and the narrow columns then wrap into very tall rows.
    """
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    old = tbl._tbl.find(qn("w:tblGrid"))
    if old is not None:
        tbl._tbl.remove(old)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        col = OxmlElement("w:gridCol")
        # w:gridCol/@w:w is in dxa (twentieths of a point); the widths arrive
        # as EMU. 914400 EMU per inch / 1440 dxa per inch = 635.
        col.set(qn("w:w"), str(int(w / 635)))
        grid.append(col)
    tbl._tbl.insert(list(tbl._tbl).index(tblPr) + 1, grid)


def table(doc, headers, rows, weights, font=8.5, hdr_font=8.5,
          bold_first=False, shade_fn=None, band=True, hdr_fill=NAVY_H,
          hide_header=False):
    """weights: relative column widths. rows: list of lists of str (\\n = new para)."""
    total_w = usable_width(doc)
    tot = sum(weights)
    widths = [Emu(int(total_w * w / tot)) for w in weights]

    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed_grid(tbl, widths)
    table_borders(tbl)

    hdr = tbl.rows[0]
    repeat_header(hdr)
    for i, htxt in enumerate(headers):
        c = hdr.cells[i]
        c.width = widths[i]
        shade(c, hdr_fill)
        cell_margins(c)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(htxt)
        r.font.name = FONT
        r.font.size = Pt(hdr_font)
        r.bold = True
        r.font.color.rgb = WHITE

    for ri, rowvals in enumerate(rows):
        row = tbl.add_row()
        fill = None
        if shade_fn:
            fill = shade_fn(ri, rowvals)
        if fill is None and band and ri % 2 == 1:
            fill = BAND_H
        for i, val in enumerate(rowvals):
            c = row.cells[i]
            c.width = widths[i]
            cell_margins(c)
            if fill:
                shade(c, fill)
            lines = str(val).split("\n")
            for li, line in enumerate(lines):
                p = c.paragraphs[0] if li == 0 else c.add_paragraph()
                p.paragraph_format.space_after = Pt(0 if li == len(lines) - 1 else 3)
                p.paragraph_format.line_spacing = 1.05
                r = p.add_run(line)
                r.font.name = FONT
                r.font.size = Pt(font)
                if i == 0 and bold_first:
                    r.bold = True
    if hide_header:
        tbl._tbl.remove(tbl.rows[0]._tr)
    return tbl


def callout(doc, label, text, colour_hex=MID_H, fill=None):
    total_w = usable_width(doc)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag, sz in (("top", 6), ("bottom", 6), ("left", 24), ("right", 6)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), colour_hex)
        borders.append(e)
    tblPr.append(borders)

    c = tbl.rows[0].cells[0]
    c.width = Emu(int(total_w))
    cell_margins(c, top=110, bottom=110, left=170, right=140)
    if fill:
        shade(c, fill)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label)
    r.font.name = FONT
    r.font.size = Pt(9.5)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(colour_hex)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.12
    r2 = p2.add_run(text)
    r2.font.name = FONT
    r2.font.size = Pt(10)
    spacer(doc, 8)
    return tbl


# ---------------------------------------------------------------- fields
def toc(doc, levels="1-2"):
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f'TOC \\o "{levels}" \\h \\z \\u'
    r._r.append(instr)
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r._r.append(sep)
    t = OxmlElement("w:t")
    t.text = ("[To generate the contents list: open in Word, click here, then press F9 — "
              "or use References ▸ Update Table.]")
    r._r.append(t)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    r._r.append(end)
    return p


def _field(paragraph, code):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    r._r.append(f1)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
    r._r.append(it)
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f2)
    r.font.name = FONT
    r.font.size = Pt(8)
    r.font.color.rgb = GREY


def header_footer(doc, header_text):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(header_text)
    r.font.name = FONT
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    _border(hp, "bottom", "B4C6E7", 4)

    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page ")
    r.font.name = FONT; r.font.size = Pt(8); r.font.color.rgb = GREY
    _field(fp, "PAGE")
    r = fp.add_run(" of ")
    r.font.name = FONT; r.font.size = Pt(8); r.font.color.rgb = GREY
    _field(fp, "NUMPAGES")


def first_page_different(doc):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
