"""Check timelines and the actual weekly upload slides against issued banks."""
import json
import re
from pathlib import Path
from pptx import Presentation
from plan_content import SHARED

ROOT = Path(__file__).resolve().parents[1]


def verify():
    for week, content in SHARED.items():
        intervals = []
        for line in content['structure'].splitlines():
            match = re.search(r'(\d):(\d\d)-(\d):(\d\d)', line)
            if match:
                a, b, c, d = map(int, match.groups())
                intervals.append((a * 60 + b, c * 60 + d))
        assert intervals, f'Week {week}: no timeline'
        assert intervals[0][0] == 0 and intervals[-1][1] == 180, week
        assert all(a < b for a, b in intervals), week
        assert all(a[1] == b[0] for a, b in zip(intervals, intervals[1:])), week

    links = json.loads((ROOT / 'data/submission_links.json').read_text())
    checked = 0
    for code in ('CML3111', 'CML3112', 'CML3113'):
        for week in (6, 8, 9, 10, 11, 12):
            aid = f'W{week:02}'
            path = next((ROOT / 'IUEA_Teaching_Plan/decks').glob(f'{code}_Week{week:02}_*.pptx'))
            slide = Presentation(path).slides[-1]
            text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
            assert f'{aid}: your graded weekly submission' in text, path
            assert f'assessment_banks/{aid}/{code}/A/TASKS.md' in text, path
            assert 'Student Registration Number' in text, path
            urls = [r.hyperlink.address for s in slide.shapes if s.has_text_frame
                    for p in s.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
            assert urls == [links[code + '|' + aid]['url']], path
            checked += 1
    print(f'15 complete three-hour timelines; {checked} cohort-specific upload slides verified')


if __name__ == '__main__':
    verify()
