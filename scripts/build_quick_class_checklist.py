"""Build the concise two-page lecturer checklist with live control links."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(1.7);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(14)
d.styles['Heading 1'].paragraph_format.space_before=Pt(10)
for n in list(d.styles.element.iter()):
 if n.tag==qn('w:pBdr'):n.getparent().remove(n)
def p(s):return d.add_paragraph(s)
def h(s):d.add_paragraph(s,'Heading 1')
def link(label,url):
 a=p('');hyp=OxmlElement('w:hyperlink');hyp.set(qn('r:id'),a.part.relate_to(url,RT.HYPERLINK,is_external=True));run=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'0563C1');pr.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);hyp.append(run);a._p.append(hyp)
import json
L=json.loads((ROOT/'data/submission_links.json').read_text());Q=json.loads((ROOT/'data/quiz_links.json').read_text())
d.add_paragraph('Week 2 class checklist','Title')
p('Lecturer guide • Updated 21 September 2026 • Kampala time')
p('For the remaining CML3111 and CML3113 Week 2 classes. Teach from the matching PowerPoint. Page 2 has owner controls and student links.')
h('Before class')
p('1. Share the cohort W02 student ZIP. Students download, extract and read 00_START_HERE.txt. They start the notebooks with you in class.')
p('2. Open the revised PowerPoint, extracted ZIP and Week2 Repair Walkthrough. Test the projector, Colab, supplied CSV and a fresh run of proof_of_life.ipynb.')
p('3. Announce FOUR tasks: Week 1 Assignment A1, Week 1 Quiz 1, Week 2 Assignment A2 and Week 2 Quiz 2. A1 and A2 use separate upload forms.')
h('Teach in the slide order')
p('4. Define files, cells, numbers, values and variables before using assignment syntax or explaining notebook memory. Explain every new symbol, predict the result, then run it.')
p('5. Start proof_of_life.ipynb together. Demonstrate cells, CSV upload and repeated execution. Model repair 01; guide attempts at 02–03; trace 04 together. These are practice, not extra submissions.')
p('6. Return to the main notebook. Students write their explanation, restart and Run all, then download their completed notebook. Show the repository and README steps in the slides.')
p('hello_engineering.py is an optional script demonstration. Students do not submit it, the ZIP or the four repair notebooks.')
h('Explain the two assignments')
p('A1: three engineering problems; for each, possible inputs, target where appropriate, data source, limitation and ML suitability. Choose one problem/dataset and justify it. Upload a complete DOCX or selectable-text PDF; narrated PPTX or notebook evidence is also accepted.')
p('A2: upload the completed main notebook AND the student’s own README.md together. Enter the repository URL. The README states programme, files and run instructions. A replacement submission must contain both files again.')
p('Upload deadlines for BOTH A1 and A2: CML3113 — Friday 25 September, 08:00; CML3111 — Friday 25 September, 11:00. These are the existing published deadlines; confirm any timetable change before class.')
h('Before dismissal')
p('7. Activate Quiz 1 and Quiz 2 when you announce them. Each stays open for 24 hours from its own activation. Tell students the exact closing times and remind them to press Submit.')
p('8. Students enter their own Student Registration Number exactly as on their record. The required field now accepts any format. Refresh an already-open form if an old pattern warning remains. Keep local copies and each confirmation.')
d.add_page_break();d.add_paragraph('Controls and student links','Title')
p('Sign into owner controls only as wirmbeesubs@gmail.com. If a controls link fails with multiple Google accounts, use a private window signed into this account only. Share responder links with students, never these owner controls.')
h('Assignment uploads')
p('A1 catch-up and scheduled A2 uploads accept work until the cohort deadlines on page 1. The scheduler closes them automatically. If a due opening is missing, run syncSubmissions in Assignment controls, wait for completion and refresh the student link. Do not rebuild forms.')
link('Assignment controls — Apps Script','https://script.google.com/home/projects/1XLdUAzi6wnSjQ2iFTV3hl5BQk9XWw-PIi1_vodC_l107B6S13yW_ZObw/edit')
link('Assignment links index — Assignment links tab','https://docs.google.com/spreadsheets/d/1FxAVXwSvLEd7TEd3pJ2E0_ka7Mp6ZM6pW1yNsXKWdm4/edit')
h('Activate BOTH quizzes for the correct cohort')
p('In Quiz controls, select open_W1_CML3113 and Run once for CML3113 Quiz 1; select open_W2_CML3113 for Quiz 2. For CML3111, use open_W1_CML3111 and open_W2_CML3111.')
p('Run only when students are told to start. Each activation starts a separate 24-hour window. Wait for Execution completed, announce the logged deadline and check the student link. Re-running resets the window; use closeStatus to inspect it without reopening.')
link('Quiz controls — IUEA ML Quizzes','https://script.google.com/home/projects/1J7ePAMAAMbNzmAlPl-WtuiSDru5VXfZ1E2211YlxKDy0l9VOPogFhz7_/edit')
for cohort in ('CML3113','CML3111'):
 h(cohort+' student links')
 for label,url in [('A1 upload — Week 1',L[cohort+'|W01']['url']),('A2 upload — Week 2',L[cohort+'|W02']['url']),('Quiz 1 — Week 1',Q[cohort+'|1']),('Quiz 2 — Week 2',Q[cohort+'|2'])]:link(label,url)
p('Completed CML3112 class: its A1/A2 upload deadline remains 22 September at 11:00. Its earlier Quiz 1 window has ended; do not reopen it inadvertently. Existing Week 1 marks remain protected; new written grading awaits validation.')
footer=sec.footer.paragraphs[0];footer.text='IUEA • Lecturer use • ';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
d.save(ROOT/'IUEA_Teaching_Plan/IUEA_ML_Quick_Class_Checklist.docx')
