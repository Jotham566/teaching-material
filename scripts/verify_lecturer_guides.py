"""Check lecturer slide references against the current saved student decks."""
import json,re
from pathlib import Path
from docx import Document
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
cache={}
def slide_text(path,number):
 if path not in cache:cache[path]=Presentation(path)
 return [s.text for s in cache[path].slides[number-1].shapes if s.has_text_frame]

def main():
 notes=[r for p in (ROOT/'data').glob('lecturer_notes_*.json') for r in json.loads(p.read_text())]
 for r in notes:
  code,_,week=r['deck'].split()
  path=next((ROOT/'IUEA_Teaching_Plan/decks').glob(f'{code}_Week{int(week):02}_*.pptx'))
  assert r['title'] in slide_text(path,r['slide']),(r['deck'],r['slide'],r['title'])
 doc=Document(ROOT/'IUEA_Teaching_Plan/IUEA_ML_Lecturer_Teaching_Guide_W1-W15.docx')
 week=None;found=set()
 for p in doc.paragraphs:
  if p.style.name=='Heading 1' and p.text.startswith('Week '):week=int(p.text.split()[1])
  if p.style.name=='Heading 3' and p.text.startswith('Slide '):found.add((week,int(p.text.split()[1])))
 expected={(int(r['deck'].split()[-1]),r['slide']) for r in notes}
 assert found==expected,'IUEA guide missing or outdated slide references'
 doc=Document(ROOT/'DataAnalysis_Track/DataAnalysis_Lecturer_Guide.docx')
 tables=[t for t in doc.tables if t.cell(0,0).text=='Slide'];count=0
 assert len(tables)==15,'Expected all fifteen independent sessions'
 for session,table in enumerate(tables,1):
  path=next((ROOT/'DataAnalysis_Track/decks').glob(f'DA_S{session:02}_*.pptx'))
  for row in table.rows[1:]:
   number=int(row.cells[0].text);title=re.sub(r'^(Define|Figure|Live code|Worked): ','',row.cells[1].text)
   assert title=='—' or title.casefold() in [t.casefold() for t in slide_text(path,number)],(session,number,title)
   count+=1
 print(json.dumps(dict(iuea_directions=len(notes),iuea_merged_entries=len(found),da_directions=count,decks=len(cache),mismatches=0)))

if __name__=='__main__':main()
