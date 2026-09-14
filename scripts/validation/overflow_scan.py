# -*- coding: utf-8 -*-
"""Post-hoc overflow scan over every built deck.

Full-bleed decoration (title bands, footers) legitimately touches the slide
edge, so it is excluded by width. Everything else is text or a figure and must
sit inside the margins. Also reports the smallest body font actually emitted,
which is how an over-aggressive auto-shrink would show up.
"""
import glob
import re
from collections import Counter

from pptx import Presentation
import os as _os
_ROOT = _os.path.dirname(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))

W_IN, H_IN = 13.333, 7.5
EMU = 914400
BAND_W = 12.6 * EMU          # wider than this is decoration
BOTTOM_LIMIT = 7.44          # footer text baseline sits above this
TOP_LIMIT = -0.02
SIDE_SLACK = 0.06
# Anything smaller than this does not project to the back of a hundred-seat
# room. It is the rendered floor of pptxkit's type ramp minus a little slack,
# so a call site with a hardcoded small size shows up here instead of hiding
# in the histogram.
FLOOR_PT = 12.4        # nothing should EVER be smaller than this
TARGET_PT = 14.4       # and body text should normally clear this

over, tiny, rows = [], [], []
fs_hist = Counter()

# Both courses share pptxkit, so both get scanned. A change to the type
# ramp or a layout affects the IUEA decks and the Data Analysis decks
# together, and scanning only one of them would hide half the damage.
# The Data Analysis decks are taken from their MANIFEST rather than from a
# glob, because a renamed session leaves the old file on disk and scanning
# it would report overflows in a deck nobody is going to teach.
_da_dir = _ROOT + "/DataAnalysis_Track/decks"
_manifest = _os.path.join(_da_dir, "MANIFEST.txt")
if _os.path.exists(_manifest):
    with open(_manifest, encoding="utf-8") as _fh:
        _da = [_os.path.join(_da_dir, ln.strip()) for ln in _fh
               if ln.strip() and not ln.startswith("#")]
else:
    _da = sorted(glob.glob(_da_dir + "/DA_*.pptx"))

FILES = (sorted(glob.glob(_ROOT + "/IUEA_Teaching_Plan/decks/CML*.pptx"))
         + _da)

for f in FILES:
    base = f.split("/")[-1]
    if base.startswith("DA_"):
        code = "DA"
        wk = int(re.search(r"DA_S(\d+)", base).group(1))
    else:
        code = base.split("_")[0]
        wk = int(re.search(r"Week(\d+)", f).group(1))
    prs = Presentation(f)
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.left is None:
                continue
            w = (sh.width or 0)
            if w > BAND_W:                       # full-bleed decoration
                continue
            l, t = sh.left / EMU, sh.top / EMU
            b, r = t + (sh.height or 0) / EMU, l + w / EMU
            if b > BOTTOM_LIMIT or t < TOP_LIMIT or r > W_IN + SIDE_SLACK or l < -SIDE_SLACK:
                txt = sh.text_frame.text[:44].replace("\n", " ") if sh.has_text_frame else "[fig]"
                over.append((code, wk, i, round(l, 2), round(t, 2), round(r, 2),
                             round(b, 2), txt))
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for run in p.runs:
                        if run.font.size and run.text.strip():
                            pt = run.font.size.pt
                            fs_hist[round(pt, 1)] += 1
                            if pt < FLOOR_PT:
                                tiny.append((code, wk, i, pt,
                                             run.text[:40].replace("\n", " ")))
        rows.append((code, wk, i))

print(f"scanned {len(set((c, w) for c, w, _ in rows))} decks, {len(rows)} slides")
print(f"\nOUT OF BOUNDS: {len(over)}")
for o in over[:25]:
    print(f"  {o[0]} W{o[1]:02d} sl{o[2]:2d}  L{o[3]} T{o[4]} R{o[5]} B{o[6]}  {o[7]!r}")
if len(over) > 25:
    print(f"  ... and {len(over) - 25} more")

# Grouped by geometry, because 278 individual overflows are almost never 278
# problems. A left edge and a width identify the LAYOUT that produced them,
# which is the thing worth fixing; the worst overshoot tells you how much
# room it needs.
if over:
    geo = {}
    for c, w, i, l, t, r, b, txt in over:
        k = (l, round(r - l, 2))
        g = geo.setdefault(k, {"n": 0, "worst": 0.0, "eg": txt, "where": ""})
        g["n"] += 1
        if b - BOTTOM_LIMIT > g["worst"]:
            g["worst"] = round(b - BOTTOM_LIMIT, 2)
            g["eg"] = txt
            g["where"] = f"{c} W{w:02d} sl{i}"
    print("\nGROUPED BY LAYOUT (left, width): count, worst overshoot in inches")
    for k in sorted(geo, key=lambda k: -geo[k]["n"]):
        g = geo[k]
        print(f"  L{k[0]:<5} w{k[1]:<5} n={g['n']:<4} over by {g['worst']:<5} "
              f"{g['where']:<18} {g['eg']!r}")

below_target = sum(n for pt, n in fs_hist.items() if pt < TARGET_PT)
total_runs = sum(fs_hist.values())
print(f"\nRUNS BELOW TARGET {TARGET_PT}pt: {below_target} of {total_runs} "
      f"({100.0 * below_target / max(1, total_runs):.0f}%) "
      f"— dense slides that shrank; candidates for cutting words")
print(f"RUNS BELOW HARD FLOOR {FLOOR_PT}pt: {len(tiny)}")
for t in tiny[:20]:
    print(f"  {t[0]} W{t[1]:02d} sl{t[2]:2d}  {t[3]}pt  {t[4]!r}")
print("\nfont size histogram (pt: runs):")
for k in sorted(fs_hist):
    print(f"  {k:5.1f} : {fs_hist[k]}")
