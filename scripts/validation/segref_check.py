# -*- coding: utf-8 -*-
"""Check every "segment N" reference against the segment that actually exists.

s_segment numbers the dividers contiguously in the order they are added and
ignores any number passed to it, while the session-clock figure counts the
break as a block. Notes written against the clock therefore point one segment
too high for everything after the break. This prints each reference next to the
title of the segment it actually lands on, so a wrong one is obvious.
"""
import glob
import re

from pptx import Presentation
import os as _os
_ROOT = _os.path.dirname(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))))

REF = re.compile(r"segments?\s+(\d+)(?:\s*(?:and|to|,)\s*(\d+))?", re.I)

for f in sorted(glob.glob(_ROOT + "/"
                          "IUEA_Teaching_Plan/decks/CML3112_*.pptx")):
    wk = int(re.search(r"Week(\d+)", f).group(1))
    prs = Presentation(f)

    # First pass: what segments does this deck actually contain?
    segs = {}
    for i, s in enumerate(prs.slides, 1):
        num = title = None
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            m = re.fullmatch(r"SEGMENT (\d+)", t)
            if m:
                num = int(m.group(1))
            elif num is not None and title is None and 3 < len(t) < 80 \
                    and "MINUTE" not in t:
                title = t
        if num is not None:
            segs[num] = (title or "?", i)

    # Second pass: every reference, resolved against that map.
    hits = []
    for i, s in enumerate(prs.slides, 1):
        blocks = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if s.has_notes_slide:
            blocks.append(s.notes_slide.notes_text_frame.text)
        for b in blocks:
            for m in REF.finditer(b):
                for g in (m.group(1), m.group(2)):
                    if g is None:
                        continue
                    n = int(g)
                    tgt = segs.get(n)
                    ctx = b[max(0, m.start() - 46):m.end() + 40]
                    ctx = " ".join(ctx.split())
                    hits.append((i, n, tgt[0] if tgt else "*** NO SUCH SEGMENT ***",
                                 ctx))
    if not hits:
        continue
    print(f"\n===== Week {wk}  ({len(segs)} segments: "
          f"{', '.join(f'{k}={v[0][:22]}' for k, v in sorted(segs.items()))}) =====")
    for sl, n, tname, ctx in hits:
        print(f"  sl{sl:3d} -> segment {n} = {tname[:34]!r}")
        print(f"        ...{ctx}...")
